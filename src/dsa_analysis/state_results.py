"""Aggregate official state primary returns, preserving every reported candidate."""

import csv
import gzip
import hashlib
import io
import json
import re
import ssl
import subprocess
import sys
from collections import defaultdict
from datetime import datetime
from html.parser import HTMLParser
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.parse import quote, urlencode
from urllib.request import HTTPCookieProcessor, HTTPSHandler, Request, build_opener
from zipfile import ZipFile

from openpyxl import load_workbook
import certifi

from .fec_presidential import _download_workbook
from .document_corpus import _extract_pdf_text
from .io import read_csv, read_json, write_csv
from .paths import ANALYSIS_DATA_DIR, CONFIG_DIR, MANUAL_DIR, RAW_DIR

FIELDS = [
    "contest_id", "cycle", "election_date", "stage", "state_code", "chamber", "district", "term",
    "election_system", "primary_party", "candidate_name", "candidate_source_id", "party",
    "write_in", "votes", "source_id", "source_url", "source_sha256", "source_locators",
    "reporting_units",
]
PRIMARY_RESULT_STAGES = {
    "primary", "primary_runoff", "special_primary", "special_primary_runoff",
    "nonpartisan_primary",
}


def _votes(value: object) -> int:
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError(f"Missing/invalid vote total: {value!r}")
    number = int(value)
    if number < 0 or number != float(value):
        raise ValueError(f"Invalid vote total: {value!r}")
    return number


def _aggregate(source: dict, records: list[dict], path: Path) -> list[dict]:
    groups = defaultdict(list)
    seen = set()
    for row in records:
        key = (row["contest"], row["candidate_id"], row["unit"])
        if key in seen:
            raise ValueError(f'{source["source_id"]}: duplicate reporting unit {key}')
        seen.add(key)
        groups[(row["contest"], row["candidate_id"])].append(row)
    output = []
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    for (contest, candidate), rows in sorted(groups.items()):
        identities = {
            (row["name"], row["party"], row["chamber"], row["district"], row["term"])
            for row in rows
        }
        if len(identities) != 1:
            raise ValueError(f"Inconsistent identity within candidate {candidate}: {identities}")
        name, party, chamber, district, term = next(iter(identities))
        ballot_parties = {row.get("primary_party", row["party"]) for row in rows}
        if len(ballot_parties) != 1:
            raise ValueError("Candidate reporting units disagree about the primary ballot category")
        primary_party = next(iter(ballot_parties)) if source["election_system"].startswith("partisan_") else ""
        contest_id = "-".join((
            source["state"].lower(), source["date"], chamber.lower(), district.lower(),
            term, primary_party.lower() or "all",
        ))
        output.append({
            "contest_id": contest_id,
            "cycle": source["date"][:4],
            "election_date": source["date"],
            "stage": (
                "superseded_primary"
                if chamber == "House" and district in source.get("superseded_house_districts", [])
                else source.get("stage", "primary")
            ),
            "state_code": source["state"],
            "chamber": chamber,
            "district": district,
            "term": term,
            "election_system": source["election_system"],
            "primary_party": primary_party,
            "candidate_name": name,
            "candidate_source_id": candidate,
            "party": party,
            "write_in": rows[0]["write_in"],
            "votes": sum(row["votes"] for row in rows),
            "source_id": source["source_id"],
            "source_url": source["url"],
            "source_sha256": source_hash,
            "source_locators": _compress_locators([row["locator"] for row in rows]),
            "reporting_units": len(rows),
        })
    if not output:
        raise ValueError(f'{source["source_id"]}: no congressional candidates parsed')
    return output


def _compress_locators(locators: list[str]) -> str:
    files = defaultdict(set)
    literal = set()
    for locator in locators:
        if ":" in locator and locator.rsplit(":", 1)[1].isdigit():
            filename, number = locator.rsplit(":", 1)
            files[filename].add(int(number))
        else:
            literal.add(locator)
    output = []
    for filename, numbers in sorted(files.items()):
        ranges = []
        start = end = None
        for number in sorted(numbers):
            if start is None:
                start = end = number
            elif number == end + 1:
                end = number
            else:
                ranges.append(str(start) if start == end else f"{start}-{end}")
                start = end = number
        ranges.append(str(start) if start == end else f"{start}-{end}")
        output.append(f"{filename}:{','.join(ranges)}")
    return " | ".join(output + sorted(literal))


def parse_california(path: Path, source: dict) -> list[dict]:
    records = []
    counties = set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["SOV Candidates Export"]
        values = sheet.iter_rows(values_only=True)
        header = next(values)
        required = {"Election Date", "County Id", "Contest ID", "Contest Name",
                    "Candidate ID", "Candidate Name", "Write-in Flag", "Party Name", "Vote Total"}
        if not required <= set(header):
            raise ValueError("California statement-of-vote schema changed")
        for number, values in enumerate(values, 2):
            row = dict(zip(header, values, strict=True))
            contest = str(row["Contest Name"] or "")
            if not contest:
                continue
            counties.add(row["County Id"])
            house = re.fullmatch(r"United States Representative District (\d+)", contest)
            senate = contest.startswith("US Senate -")
            if not house and not senate:
                continue
            election_date = datetime.strptime(str(row["Election Date"]), "%m/%d/%Y").date()
            if election_date.isoformat() != source["date"]:
                raise ValueError("California workbook contains a different election date")
            records.append({
                "contest": str(row["Contest ID"]), "candidate_id": str(row["Candidate ID"]),
                "unit": str(row["County Id"]), "name": str(row["Candidate Name"]).strip(),
                "party": str(row["Party Name"]).strip(),
                "chamber": "House" if house else "Senate",
                "district": house[1].zfill(2) if house else "S",
                "term": "special" if "Partial/Unexpired" in contest else "regular",
                "write_in": str(row["Write-in Flag"]) == "Y",
                "votes": _votes(row["Vote Total"]), "locator": f"{sheet.title}:{number}",
            })
    finally:
        workbook.close()
    output = _aggregate(source, records, path)
    house_districts = {row["district"] for row in output if row["chamber"] == "House"}
    senate_contests = {row["contest_id"] for row in output if row["chamber"] == "Senate"}
    if house_districts != {f"{i:02}" for i in range(1, source["expected_house_districts"] + 1)}:
        raise ValueError("California House district coverage is incomplete")
    if len(senate_contests) != source["expected_senate_contests"]:
        raise ValueError("California Senate regular/special contest coverage is incomplete")
    if len(counties) != source["expected_counties"]:
        raise ValueError("California county coverage is incomplete")
    for row in output:
        if row["chamber"] == "Senate" and row["reporting_units"] != source["expected_counties"]:
            raise ValueError("California Senate candidate is missing county returns")
    return output


def parse_florida(path: Path, source: dict) -> list[dict]:
    contests = {}
    counties = set()
    with ZipFile(path) as archive:
        names = sorted(
            (name for name in archive.namelist() if name.lower().endswith(".txt")),
            key=lambda name: ("recount" in name.lower(), name),
        )
        for filename in names:
            file_contests = defaultdict(list)
            text = archive.read(filename).decode("utf-8-sig")
            for number, row in enumerate(csv.reader(io.StringIO(text), delimiter="\t"), 1):
                if len(row) != 19:
                    raise ValueError(f"Florida precinct schema changed: {filename}:{number}")
                counties.add(row[0])
                if row[11] not in {"United States Senator", "Representative in Congress"}:
                    continue
                if row[14].strip().casefold() in {"overvotes", "undervotes"}:
                    continue
                if datetime.strptime(row[3], "%m/%d/%Y").date().isoformat() != source["date"]:
                    raise ValueError("Florida precinct file contains a different election date")
                chamber = "Senate" if row[11] == "United States Senator" else "House"
                district_match = re.fullmatch(r"(?:District\s+)?(\d+)", row[12].strip(), re.I)
                if chamber == "House" and district_match is None:
                    raise ValueError(f"Unknown Florida congressional district: {row[12]!r}")
                district = "S" if chamber == "Senate" else district_match[1].zfill(2)
                key = (row[0], row[11], district, row[15])
                file_contests[key].append({
                    "contest": f"{chamber}-{district}-{row[15]}",
                    "candidate_id": row[17], "unit": f"{row[0]}:{row[5]}",
                    "name": row[14].strip(), "party": row[15],
                    "chamber": chamber, "district": district, "term": "regular",
                    "write_in": row[15] == "WRI", "votes": _votes(row[18]),
                    "locator": f"{filename}:{number}",
                })
            for key, rows in file_contests.items():
                if key in contests and "recount" not in filename.lower():
                    raise ValueError(f"Duplicate Florida county contest in {filename}")
                # Recounts replace the entire affected county/party contest, not add to it.
                contests[key] = rows
    if len(counties) != source["expected_counties"]:
        raise ValueError("Florida statewide archive is missing county files")
    return _aggregate(source, [row for group in contests.values() for row in group], path)


def parse_north_carolina(path: Path, source: dict) -> list[dict]:
    records = []
    counties = set()
    with ZipFile(path) as archive:
        for filename in archive.namelist():
            if not filename.endswith(".txt"):
                continue
            text = archive.read(filename).decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text), delimiter="\t")
            required = {"County", "Election Date", "Precinct", "Contest Name", "Choice",
                        "Choice Party", "Total Votes"}
            if not required <= set(reader.fieldnames or []):
                raise ValueError("North Carolina precinct schema changed")
            for number, row in enumerate(reader, 2):
                counties.add(row["County"])
                name = row["Contest Name"]
                house = re.fullmatch(
                    r"US HOUSE OF REPRESENTATIVES DISTRICT (\d+)(?: \([A-Z]+\))?", name,
                )
                senate = re.fullmatch(r"US SENATE(?: \([A-Z]+\))?", name)
                if not house and not senate:
                    continue
                if datetime.strptime(row["Election Date"], "%m/%d/%Y").date().isoformat() != source["date"]:
                    raise ValueError("North Carolina returns contain a different election date")
                components = ("Election Day", "Early Voting", "Absentee by Mail", "Provisional")
                total = _votes(row["Total Votes"])
                if all(field in row for field in components):
                    if sum(_votes(row[field]) for field in components) != total:
                        raise ValueError(f"North Carolina vote components differ: {filename}:{number}")
                records.append({
                    "contest": name, "candidate_id": row["Choice"].strip(),
                    "unit": f'{row["County"]}:{row["Precinct"]}',
                    "name": row["Choice"].strip(), "party": row["Choice Party"].strip(),
                    "chamber": "House" if house else "Senate",
                    "district": house[1].zfill(2) if house else "S",
                    "term": source.get("term", "regular"),
                    "write_in": "WRITE-IN" in row["Choice"].upper(),
                    "votes": total, "locator": f"{filename}:{number}",
                })
    if "expected_counties" in source and len(counties) != source["expected_counties"]:
        raise ValueError("North Carolina export is missing counties")
    return _aggregate(source, records, path)


def parse_virginia(path: Path, source: dict) -> list[dict]:
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"CandidateId", "CandidateName", "TOTAL_VOTES", "Party", "WriteInVote",
                    "LocalityCode", "PrecinctId", "DistrictName", "DistrictType",
                    "OfficeTitle", "ElectionDate", "ElectionType"}
        if not required <= set(reader.fieldnames or []):
            raise ValueError("Virginia official results schema changed")
        for number, row in enumerate(reader, 2):
            house = row["DistrictType"].casefold() == "congressional"
            senate = bool(re.search(r"(?:United States|U\.S\.)\s+Senat", row["OfficeTitle"], re.I))
            if not house and not senate:
                continue
            if row["ElectionDate"] != source["date"] or row["ElectionType"] != "Primary":
                raise ValueError("Virginia source is not the requested primary")
            if not row["CandidateId"] or not row["CandidateName"]:
                raise ValueError("Virginia congressional result lacks candidate identity")
            district = str(int(row["DistrictName"])).zfill(2) if house else "S"
            records.append({
                "contest": f'{row["DistrictType"]}-{district}-{row["Party"]}',
                "candidate_id": row["CandidateId"],
                "unit": f'{row["LocalityCode"]}:{row["PrecinctId"]}',
                "name": row["CandidateName"], "party": row["Party"],
                "chamber": "House" if house else "Senate", "district": district,
                "term": source.get("term", "regular"), "write_in": row["WriteInVote"] == "1",
                "votes": _votes(row["TOTAL_VOTES"]),
                "locator": f"{source.get('filename', path.name)}:{number}",
            })
    return _aggregate(source, records, path)


def _virginia_report_source(source: dict, raw_dir: Path, download: bool) -> dict:
    metadata_path = (raw_dir / source["filename"]).with_suffix(".json")
    if download:
        payload = _download_workbook(source["url"])
        metadata = json.loads(payload)
    else:
        metadata = read_json(metadata_path)
    if metadata["electionDate"] != source["date"] or metadata["isOfficialResults"] is not True:
        raise ValueError("Virginia publication is not certified for the requested date")
    reports = [
        report for category in metadata["publicReportCategories"]
        for report in category["reports"] if report["reportName"] == "Election Results"
    ]
    if len(reports) != 1:
        raise ValueError("Virginia publication lacks a unique complete results report")
    report_url = (
        "https://enr.elections.virginia.gov/cdn/results/"
        + quote(metadata["jurisdictionId"], safe="") + "/"
        + quote(reports[0]["blobName"], safe="")
    )
    if download:
        metadata_path.write_bytes(payload)
    return {**source, "url": report_url}


def parse_reviewed_totals(path: Path, source: dict) -> list[dict]:
    rows = [
        row for row in read_csv(MANUAL_DIR / "congressional_source_extractions.csv")
        if row["source_id"] == source["source_id"]
    ]
    if len(rows) != source["expected_ballot_options"]:
        raise ValueError("Reviewed state result extraction has an unexpected row count")
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    output = []
    for row in rows:
        if row["source_sha256"] != source_hash or row["election_date"] != source["date"]:
            raise ValueError("Reviewed totals do not match the current source document")
        primary_party = row["party"] if source["election_system"].startswith("partisan_") else ""
        output.append({
            "contest_id": "-".join((
                source["state"].lower(), source["date"], row["chamber"].lower(),
                row["district"].lower(), row["term"], primary_party.lower() or "all",
            )),
            "cycle": source["date"][:4], "election_date": source["date"],
            "stage": source.get("stage", "primary"), "state_code": source["state"],
            "chamber": row["chamber"], "district": row["district"], "term": row["term"],
            "election_system": source["election_system"],
            "primary_party": primary_party,
            "candidate_name": row["candidate_name"], "candidate_source_id": row["candidate_name"],
            "party": row["party"], "write_in": row["write_in"] == "true",
            "votes": _votes(row["votes"]), "source_id": source["source_id"],
            "source_url": source["url"], "source_sha256": source_hash,
            "source_locators": row["locator"], "reporting_units": 1,
        })
    return output


def parse_reviewed_contests(path: Path, source: dict) -> list[dict]:
    extraction = read_json(MANUAL_DIR / source["extraction_file"])[source["source_id"]]
    if hashlib.sha256(path.read_bytes()).hexdigest() != extraction["source_sha256"]:
        raise ValueError("Reviewed contest extraction does not match the current source")
    records = []
    if "expected_reporting_units" in source:
        if {row.get("reporting_unit") for row in extraction["contests"]} != set(source["expected_reporting_units"]):
            raise ValueError("Reviewed county extraction is missing a required reporting unit")
    for contest in extraction["contests"]:
        if "expected_total" in contest and sum(contest["votes"]) != contest["expected_total"]:
            raise ValueError("Reviewed candidate totals differ from the official contest total")
        for name, votes in zip(contest["names"], contest["votes"], strict=True):
            write_in = name in contest.get("write_in_names", [])
            records.append({
                "contest": f'{contest["chamber"]}-{contest["district"]}-{contest["party"]}',
                "candidate_id": name,
                "unit": contest.get("reporting_unit", "statewide"), "name": name,
                "party": "" if write_in else contest["party"],
                "primary_party": contest["party"], "chamber": contest["chamber"],
                "district": contest["district"], "term": contest.get("term", "regular"),
                "write_in": write_in, "votes": _votes(votes), "locator": contest["locator"],
            })
    output = _aggregate(source, records, path)
    if len(output) != source["expected_ballot_options"]:
        raise ValueError("Reviewed contest extraction has an unexpected candidate count")
    if "expected_reporting_units" in source and any(
        row["reporting_units"] != len(source["expected_reporting_units"]) for row in output
    ):
        raise ValueError("Reviewed candidate is missing a county total")
    return output


def parse_maine_totals(path: Path, source: dict) -> list[dict]:
    if source.get("expected_sha256") and hashlib.sha256(path.read_bytes()).hexdigest() != source["expected_sha256"]:
        raise ValueError("Maine's non-year-specific download changed; review before updating the pinned snapshot")
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        sheet = workbook[source["sheet"]] if "sheet" in source else workbook.worksheets[0]
        rows = list(sheet.values)
        header = rows[0]
        first = source["first_candidate_column"] - 1
        blank = next(i for i, value in enumerate(header) if str(value).strip().upper() == "BLANK")
        tbc = next(
            (i for i, value in enumerate(header) if str(value).strip().upper() == "TBC"),
            len(header) - 1,
        )
        if source.get("totals_row"):
            total_number = source["totals_row"]
            total = rows[total_number - 1]
        else:
            totals = [
                (number, row) for number, row in enumerate(rows, 1)
                if any(
                    str(value).strip().casefold() == source["totals_label"].casefold()
                    if "totals_label" in source
                    else bool(re.fullmatch(r"state totals?", str(value).strip(), re.I))
                    for value in row
                )
            ]
            if len(totals) != 1:
                raise ValueError("Maine workbook lacks a unique statewide total")
            total_number, total = totals[0]
        if sum(_votes(total[i]) for i in range(first, blank + 1)) != _votes(total[tbc]):
            raise ValueError("Maine candidate/blank totals differ from total ballots cast")
        for column in range(first, blank):
            if not header[column] or not str(header[column]).strip():
                raise ValueError("Maine candidate column is missing its name")
            raw_name = str(header[column]).strip()
            name = re.sub(r"\s*\(Declared Write-In\)$", "", raw_name, flags=re.I)
            records.append({
                "contest": f'{source["chamber"]}-{source["district"]}-{source["party"]}',
                "candidate_id": name, "unit": "statewide",
                "name": name, "party": source["party"], "chamber": source["chamber"],
                "district": source["district"], "term": "regular",
                "write_in": "write-in" in raw_name.casefold()
                or any("write-in" in str(row[column]).casefold() for row in rows[1:3]),
                "votes": _votes(total[column]),
                "locator": f"{sheet.title} column {column + 1}:{total_number}",
            })
    finally:
        workbook.close()
    return _aggregate(source, records, path)


def parse_rhode_island(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        sheet = workbook["Candidate_Breakout"]
        values = sheet.iter_rows(values_only=True)
        header = next(values)
        if tuple(header) != (
            "Party", "Contest Title", "Candidate", "Precinct", "Total",
            "Election Day", "Mail", "Early Voting",
        ):
            raise ValueError("Rhode Island candidate export schema changed")
        for number, row in enumerate(values, 2):
            contest = str(row[1] or "")
            match = re.fullmatch(
                r"(?:DEM|REP) (Senator in Congress|Representative in Congress District (\d+))",
                contest,
            )
            if not match:
                continue
            total = _votes(row[4])
            if sum(_votes(value) for value in row[5:8]) != total:
                raise ValueError("Rhode Island voting-mode totals do not reconcile")
            records.append({
                "contest": contest, "candidate_id": str(row[2]), "unit": str(row[3]),
                "name": str(row[2]), "party": row[0],
                "chamber": "House" if match[2] else "Senate",
                "district": match[2].zfill(2) if match[2] else "S",
                "term": "regular", "write_in": "write-in" in str(row[2]).casefold(),
                "votes": total, "locator": f"{sheet.title}:{number}",
            })
    finally:
        workbook.close()
    return _aggregate(source, records, path)


def parse_civera(path: Path, source: dict) -> list[dict]:
    records = []
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {
            "contest_id", "election_date", "election_type", "primary_party", "office_name",
            "office_modifier", "district_name", "candidate_id", "candidate_name",
            "division_id", "division_type", "vote_channel", "candidate_party_name", "votes",
        }
        if not required <= set(reader.fieldnames or []):
            raise ValueError("Civera public export schema changed")
        for number, row in enumerate(reader, 2):
            if (
                row["election_date"][:10] != source["date"]
                or row["election_type"] != "Primary"
                or row["office_name"] != source["office_name"]
                or row["office_modifier"] not in source.get("office_modifiers", [""])
                or row["division_type"] != "County" or row["vote_channel"]
            ):
                continue
            if row["candidate_name"].casefold() in {
                "total votes cast", "total ballots cast", "blank votes", "overvotes", "undervotes",
            }:
                continue
            if not row["primary_party"]:
                raise ValueError("Partisan Civera primary lacks its ballot-party category")
            records.append({
                "contest": row["contest_id"], "candidate_id": row["candidate_id"],
                "unit": row["division_id"], "name": row["candidate_name"],
                "party": row["candidate_party_name"], "primary_party": row["primary_party"],
                "chamber": source["chamber"], "district": str(int(row["district_name"])).zfill(2),
                "term": source.get("term", "regular"),
                "write_in": row["candidate_name"].casefold() in {"write-in", "scattering"},
                "votes": _votes(row["votes"]),
                "locator": f"{source.get('filename', path.name)}:{number}",
            })
    output = _aggregate(source, records, path)
    if len({row["contest_id"] for row in output}) != source["expected_contests"]:
        raise ValueError("Civera primary contest coverage differs from the reviewed inventory")
    return output


def parse_washington(path: Path, source: dict) -> list[dict]:
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not {"Race", "Candidate", "Party", "Votes"} <= set(reader.fieldnames or []):
            raise ValueError("Washington congressional export schema changed")
        for number, row in enumerate(reader, 2):
            house = re.fullmatch(r"Congressional District (\d+) - U\.S\. Representative", row["Race"], re.I)
            senate = row["Race"] == "U.S. Senator"
            if not house and not senate:
                raise ValueError(f"Unexpected race in congressional-only export: {row['Race']}")
            records.append({
                "contest": row["Race"].casefold(), "candidate_id": row["Candidate"],
                "unit": "statewide", "name": row["Candidate"], "party": row["Party"].strip(),
                "chamber": "House" if house else "Senate",
                "district": house[1].zfill(2) if house else "S",
                "term": "regular", "write_in": row["Candidate"] == "WRITE-IN",
                "votes": _votes(row["Votes"]), "locator": f"{source['filename']}:{number}",
            })
    output = _aggregate(source, records, path)
    if len({row["contest_id"] for row in output}) != source["expected_contests"]:
        raise ValueError("Washington federal contest coverage is incomplete")
    return output


def parse_alabama(path: Path, source: dict) -> list[dict]:
    try:
        import xlrd
    except ImportError as error:
        raise RuntimeError("Legacy canvasses require xlrd: use `uv run --with xlrd ...`") from error
    records = []
    files = []
    with ZipFile(path) as archive:
        for filename in archive.namelist():
            if not filename.lower().endswith(".xls"):
                continue
            files.append(filename)
            workbook = xlrd.open_workbook(file_contents=archive.read(filename))
            sheet = workbook.sheet_by_index(0)
            header = sheet.row_values(0)
            if header[:3] != ["Contest Title", "Party", "Candidate"]:
                raise ValueError("Alabama precinct spreadsheet schema changed")
            groups = defaultdict(list)
            for number in range(1, sheet.nrows):
                row = sheet.row_values(number)
                title, party, name = (str(value).strip() for value in row[:3])
                house = re.fullmatch(
                    r"UNITED STATES REPRESENTATIVE, (\d+)(?:ST|ND|RD|TH) CONGRESSIONAL DISTRICT"
                    r"(?: \((?:DEM|REP)\))?", title,
                )
                senate = re.fullmatch(r"UNITED STATES SENATOR(?: \((?:DEM|REP)\))?", title)
                if not house and not senate:
                    continue
                if name.casefold() in {"over votes", "under votes", "overvotes", "undervotes"}:
                    continue
                key = ("House" if house else "Senate", house[1].zfill(2) if house else "S", party)
                groups[key].append((number + 1, name, row))
            for (chamber, district, party), candidates in groups.items():
                for column in range(3, sheet.ncols):
                    values = [row[column] for _, _, row in candidates]
                    if all(value == "" for value in values):
                        continue
                    if any(value == "" for value in values):
                        raise ValueError(
                            f"{filename}:{header[column]}: partially missing candidate vote cells"
                        )
                    for number, name, row in candidates:
                        records.append({
                            "contest": f"{chamber}-{district}-{party}", "candidate_id": name,
                            "unit": f"{filename}:column-{column + 1}",
                            "name": name, "party": party, "chamber": chamber,
                            "district": district, "term": "regular",
                            "write_in": "write-in" in name.casefold(),
                            "votes": _votes(row[column]),
                            "locator": f"{filename} column {column + 1}:{number}",
                        })
    if len(files) != source["expected_counties"]:
        raise ValueError("Alabama precinct archive has an unexpected county-file count")
    return _aggregate(source, records, path)


def parse_alabama_replacement_summary(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    county_records = []
    state_totals = {}
    try:
        for sheet in workbook:
            rows = list(sheet.values)
            for number, row in enumerate(rows):
                for column, value in enumerate(row):
                    match = re.fullmatch(r"U\.S\. Congress . Congressional District (\d+)", str(value))
                    if not match:
                        continue
                    district = match[1].zfill(2)
                    if district not in source["verified_districts"]:
                        continue
                    candidate_rows = []
                    reported_total = None
                    for index in range(number + 3, len(rows)):
                        values = rows[index]
                        name = values[column]
                        if not name:
                            break
                        votes = _votes(values[column + 1])
                        if name == "Total Votes":
                            reported_total = votes
                            break
                        candidate_rows.append((index + 1, str(name).strip(), votes))
                    if not candidate_rows or reported_total is None:
                        raise ValueError("Alabama replacement summary is missing a candidate block or total")
                    if sum(count for _, _, count in candidate_rows) != reported_total:
                        raise ValueError("Alabama replacement candidate votes do not match the block total")
                    for index, name, votes in candidate_rows:
                        if sheet.title == "Summary":
                            state_totals[(district, name)] = votes
                        else:
                            county_records.append({
                                "contest": district, "candidate_id": name,
                                "unit": sheet.title, "name": name, "party": "Republican",
                                "chamber": "House", "district": district, "term": "regular",
                                "write_in": False, "votes": votes,
                                "locator": f"{sheet.title} column {column + 1}:{index}",
                            })
    finally:
        workbook.close()
    output = _aggregate(source, county_records, path)
    aggregated = {(row["district"], row["candidate_name"]): row["votes"] for row in output}
    if aggregated != state_totals:
        raise ValueError("Alabama replacement county totals do not match the official statewide summary")
    return output


def parse_tally(path: Path, source: dict) -> list[dict]:
    data = read_json(path)
    info = read_json(path.parent / source["info_filename"])
    metadata = read_json(path.parent / source["contests_filename"])
    election = data.get("ElectionData", data.get("ElectionInfo"))
    if not election or election.get("IsOfficial") is not True:
        raise ValueError("Tally result file is not marked official")
    if info.get("isOfficial") is not True or metadata.get("isOfficial") is not True:
        raise ValueError("Tally metadata is not marked official")
    if info["response"]["date"][:10] != source["date"]:
        raise ValueError("Tally metadata has the wrong election date")
    if str(election["ElectionID"]) != str(info["response"]["id"]):
        raise ValueError("Tally result/metadata election IDs differ")
    if info["versionID"] != metadata["versionID"]:
        raise ValueError("Tally name metadata was acquired from different result versions")
    current = "ElectionData" in data
    result_map = {
        str(row["ContestID"] if current else row["ContestName"]): (index, row)
        for index, row in enumerate(data["ContestData"])
    }
    parties = info["response"].get("parties") or info["response"].get("partyInfo")
    if parties is not None and not isinstance(parties, dict):
        raise ValueError("Unsupported Tally party dictionary schema")
    records = []
    for contest_id, definition in metadata["response"]["contests"].items():
        title = definition["contestName"]
        override = source.get("contest_overrides", {}).get(title)
        match = re.fullmatch(
            r"(?:(DEM|REP|LIB|GRN) )?U\.S\. (Senate|Congress District (\d+))"
            r"(?: - (DEM|REP|LIB|GRN))?", title,
        )
        if not match and not override:
            continue
        primary_party = override["primary_party"] if override else match[1] or match[4]
        chamber = override["chamber"] if override else ("House" if match[3] else "Senate")
        district = override["district"] if override else (match[3].zfill(2) if match[3] else "S")
        if not primary_party:
            raise ValueError(f"Tally partisan primary lacks ballot category: {title}")
        key = str(contest_id) if current else title
        if key not in result_map:
            raise ValueError(f"Congressional contest metadata has no result: {title}")
        index, result = result_map[key]
        if result["PrecinctsReporting"] != result["TotalPrecincts"]:
            raise ValueError(f"Tally federal contest reporting is incomplete: {title}")
        options = result["Choices"] + result.get("WriteinChoices", []) if current else result["Candidates"]
        if sum(_votes(option["TotalVotes"]) for option in options) != _votes(result["TotalVotes"]):
            raise ValueError(f"Tally candidate totals do not reconcile: {title}")
        names = {choice["name"]: choice for choice in definition["choices"].values()}
        seen = set()
        for option_index, option in enumerate(options):
            choice = definition["choices"].get(str(option["ChoiceID"])) if current else names.get(option["Name"])
            if choice is None:
                raise ValueError(f"Unidentified Tally candidate option in {title}")
            seen.add(str(choice["id"]))
            name = " ".join(choice["name"].split())
            write_in = choice.get("isWriteIn") is True or name.casefold() in {
                "write-in", "write in", "scattering", "scattered",
            }
            party_id = str(choice.get("partyID") or "")
            party = parties.get(party_id, {}).get("partyName", party_id) if parties else party_id
            records.append({
                "contest": str(contest_id), "candidate_id": str(choice["id"]),
                "unit": "statewide", "name": name,
                "party": "" if write_in else party, "primary_party": primary_party,
                "chamber": chamber, "district": district,
                "term": source.get("term", "regular"), "write_in": write_in,
                "votes": _votes(option["TotalVotes"]),
                "locator": f"ContestData[{index}].{'Choices' if current else 'Candidates'}[{option_index}]",
            })
        if seen != {str(choice["id"]) for choice in definition["choices"].values()}:
            raise ValueError(f"Tally result omits a listed congressional candidate: {title}")
    return _aggregate(source, records, path)


def parse_montana(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    counties = set()
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if header[0] != "County ID" or header[7] != "Votes":
            raise ValueError("Montana precinct export schema changed")
        for number, row in enumerate(rows, 2):
            if row[0]:
                counties.add(str(row[0]))
            office = str(row[3] or "")
            senate = office == "UNITED STATES SENATOR"
            house = office == "UNITED STATES REPRESENTATIVE" or bool(
                re.fullmatch(r"US REPRESENTATIVE DIST \d+", office)
            )
            if not house and not senate:
                continue
            district_match = re.fullmatch(r"(\d+)(?:ST|ND|RD|TH) CONGRESSIONAL", str(row[4]))
            if house and not district_match:
                raise ValueError(f"Unrecognized Montana congressional district: {row[4]}")
            district = district_match[1].zfill(2) if house else "S"
            name = str(row[8]).strip()
            if not name or name == "None":
                raise ValueError("Montana federal result lacks a ballot name")
            records.append({
                "contest": f"{office}-{district}-{row[6]}", "candidate_id": name,
                "unit": f"{row[0]}:{row[2]}", "name": name, "party": str(row[6]).strip(),
                "chamber": "House" if house else "Senate", "district": district,
                "term": "regular", "write_in": "WRITE-IN" in name.upper(),
                "votes": _votes(row[7]), "locator": f"{sheet.title}:{number}",
            })
    finally:
        workbook.close()
    if len(counties) != 56:
        raise ValueError("Montana precinct export must cover all 56 counties")
    output = _aggregate(source, records, path)
    if {row["district"] for row in output if row["chamber"] == "House"} != {"01", "02"}:
        raise ValueError("Montana must cover both congressional districts")
    if not any(row["chamber"] == "Senate" for row in output):
        raise ValueError("Montana source is missing its scheduled Senate primary")
    return output


def _download_webform(source: dict) -> bytes:
    class Form(HTMLParser):
        def __init__(self):
            super().__init__()
            self.fields = {}

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if tag == "input" and attrs.get("type") == "hidden" and attrs.get("name"):
                self.fields[attrs["name"]] = attrs.get("value", "")

    user_agent = "dsa-analysis/0.1 (+source-first academic research)"
    opener = build_opener(
        HTTPCookieProcessor(CookieJar()),
        HTTPSHandler(context=ssl.create_default_context(cafile=certifi.where())),
    )
    request = Request(source["url"], headers={"User-Agent": user_agent})
    with opener.open(request, timeout=60) as response:
        parser = Form()
        parser.feed(response.read().decode("utf-8"))
    if "__VIEWSTATE" not in parser.fields:
        raise ValueError("Public export form no longer has its expected state fields")
    parser.fields.update(source["form_fields"])
    request = Request(
        source["url"], data=urlencode(parser.fields).encode(),
        headers={"Content-Type": "application/x-www-form-urlencoded",
                 "Referer": source["url"], "User-Agent": user_agent},
    )
    with opener.open(request, timeout=120) as response:
        if "attachment" not in response.headers.get("Content-Disposition", "").lower():
            raise ValueError("Public export form did not return a downloadable result file")
        return response.read()


def parse_statewide_export(path: Path, source: dict) -> list[dict]:
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"ContestName", "PartyCode", "CandidateName", "CandidateVotes", "PrecinctsReporting"}
        if not required <= set(reader.fieldnames or []):
            raise ValueError("Statewide export schema changed")
        for number, row in enumerate(reader, 2):
            office = source["offices"].get(row["ContestName"])
            if office is None:
                continue
            reporting = re.fullmatch(r"(\d+)/(\d+)", row["PrecinctsReporting"].strip())
            if not reporting or int(reporting[1]) != int(reporting[2]):
                raise ValueError("Statewide congressional export is not fully reported")
            name = " ".join(row["CandidateName"].split())
            write_in = name.casefold() in {"write-in", "write in", "scattered"}
            records.append({
                "contest": f'{row["ContestName"]}-{row["PartyCode"]}',
                "candidate_id": name, "unit": "statewide", "name": name,
                "party": "" if write_in else row["PartyCode"],
                "primary_party": source["party_names"][row["PartyCode"]],
                "chamber": office["chamber"], "district": office["district"],
                "term": source.get("term", "regular"), "write_in": write_in,
                "votes": _votes(row["CandidateVotes"]),
                "locator": f"{source['filename']}:{number}",
            })
    output = _aggregate(source, records, path)
    if len({row["contest_id"] for row in output}) != source["expected_contests"]:
        raise ValueError("Statewide export does not cover the expected federal primary contests")
    return output


def parse_hawaii(path: Path, source: dict) -> list[dict]:
    payload = path.read_bytes()
    encoding = "utf-16" if payload.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    lines = payload.decode(encoding).splitlines()
    if len(lines) < 3 or not lines[1].startswith("#Contest ID"):
        raise ValueError("Hawaii summary is missing its format/header declaration")
    lines[1] = lines[1].lstrip("#")
    delimiter = "\t" if "\t" in lines[1] else ","
    reader = csv.DictReader(io.StringIO("\n".join(lines[1:])), delimiter=delimiter)
    required = {
        "Contest ID", "Contest Title", "Contest Party", "Candidate ID", "Candidate Name",
        "Candidate Party", "Mail Votes", "In-Person Votes", "Total Votes",
        "Total Precincts", "Counted Precincts",
    }
    if not required <= set(reader.fieldnames or []):
        raise ValueError("Hawaii primary summary schema changed")
    records = []
    for number, row in enumerate(reader, 3):
        title = row["Contest Title"]
        house = re.fullmatch(r"U\.S\. Representative, Dist (I|II)", title)
        senate = title == "U.S. Senator"
        if not house and not senate:
            continue
        if _votes(row["Counted Precincts"]) != _votes(row["Total Precincts"]):
            raise ValueError("Hawaii federal contest is not fully reported")
        total = _votes(row["Total Votes"])
        if _votes(row["Mail Votes"]) + _votes(row["In-Person Votes"]) != total:
            raise ValueError("Hawaii federal voting-mode totals do not reconcile")
        records.append({
            "contest": row["Contest ID"], "candidate_id": row["Candidate ID"],
            "unit": "statewide", "name": row["Candidate Name"],
            "party": row["Candidate Party"].strip(),
            "primary_party": row["Contest Party"],
            "chamber": "House" if house else "Senate",
            "district": {"I": "01", "II": "02"}[house[1]] if house else "S",
            "term": "regular", "write_in": False, "votes": total,
            "locator": f"{source['filename']}:{number}",
        })
    output = _aggregate(source, records, path)
    if len({row["contest_id"] for row in output}) != source["expected_contests"]:
        raise ValueError("Hawaii federal primary contest coverage is incomplete")
    return output


def parse_dc(path: Path, source: dict) -> list[dict]:
    info = read_json(path.parent / source["info_filename"])
    if info["ElectionResultsTypeName"] != "Certified Results" or info["ElectionDate"][:10] != source["date"]:
        raise ValueError("DC metadata does not certify the requested primary")
    if info["PrecinctsCounted"] != info["TotalPrecincts"]:
        raise ValueError("DC certified election metadata does not report every precinct")
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"ElectionDate", "ContestNumber", "ContestName", "PrecinctNumber",
                    "WardNumber", "Candidate", "Party", "Votes"}
        if not required <= set(reader.fieldnames or []):
            raise ValueError("DC certified CSV schema changed")
        for number, row in enumerate(reader, 2):
            if not re.search(r"DELEGATE TO THE (?:U\.S\. )?HOUSE OF REPRESENTATIVES", row["ContestName"]):
                continue
            if row["Candidate"].strip().upper() in {"OVER VOTES", "UNDER VOTES"}:
                continue
            date = datetime.strptime(row["ElectionDate"].split()[0], "%m/%d/%Y").date()
            if date.isoformat() != source["date"]:
                raise ValueError("DC CSV has a different election date")
            name = row["Candidate"].strip()
            if not name:
                raise ValueError("DC delegate result lacks its ballot-option name")
            write_in = name.casefold() == "write-in"
            records.append({
                "contest": row["ContestNumber"], "candidate_id": name,
                "unit": f'{row["WardNumber"]}:{row["PrecinctNumber"]}',
                "name": name, "party": "" if write_in else row["Party"],
                "primary_party": row["Party"], "chamber": "House", "district": "00",
                "term": "regular", "write_in": write_in, "votes": _votes(row["Votes"]),
                "locator": f"{source['filename']}:{number}",
            })
    return _aggregate(source, records, path)


def parse_delaware(path: Path, source: dict) -> list[dict]:
    records = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"office", "candidatename", "partyname", "electiondate",
                    "machinevotessum", "absenteevotessum", "earlyvotessum", "totalvotessum"}
        if not required <= set(reader.fieldnames or []):
            raise ValueError("Delaware result CSV schema changed")
        for number, row in enumerate(reader, 2):
            house = row["office"] == "U.S. Representative in Congress"
            senate = row["office"] == "U.S. Senator"
            if not house and not senate:
                continue
            if row["electiondate"][:10] != source["date"]:
                raise ValueError("Delaware result has the wrong election date")
            total = _votes(row["totalvotessum"])
            if sum(_votes(row[key]) for key in (
                "machinevotessum", "absenteevotessum", "earlyvotessum",
            )) != total:
                raise ValueError("Delaware voting-mode totals do not reconcile")
            records.append({
                "contest": f'{row["office"]}-{row["partyname"]}',
                "candidate_id": row["candidatename"], "unit": "statewide",
                "name": row["candidatename"], "party": row["partyname"],
                "chamber": "House" if house else "Senate",
                "district": "00" if house else "S", "term": "regular",
                "write_in": False, "votes": total,
                "locator": f"{source['filename']}:{number}",
            })
    return _aggregate(source, records, path)


def parse_wyoming(path: Path, source: dict) -> list[dict]:
    with ZipFile(path) as archive:
        members = [name for name in archive.namelist() if "Results Summaries" in name and name.endswith(".xlsx")]
        if len(members) != 1:
            raise ValueError("Wyoming archive lacks a unique official statewide summary")
        workbook = load_workbook(io.BytesIO(archive.read(members[0])), read_only=True, data_only=True)
    records = []
    try:
        sheet = workbook["Statewide Candidates"]
        rows = list(sheet.values)
        totals = [(number, row) for number, row in enumerate(rows, 1) if row[0] == "Total"]
        if len(totals) != 1:
            raise ValueError("Wyoming summary lacks a unique statewide total row")
        total_number, total = totals[0]
        county_rows = rows[5:total_number - 1]
        if len(county_rows) != 23 or len({row[0] for row in county_rows}) != 23:
            raise ValueError("Wyoming summary must reconcile all 23 counties")
        office = party = ""
        for column in range(1, sheet.max_column):
            if rows[2][column]:
                office = str(rows[2][column]).split(", Continued")[0]
            if rows[3][column]:
                party = str(rows[3][column])
            name = rows[4][column]
            if office not in {"United States Senator", "United States Representative"}:
                continue
            if not name or str(name).strip().casefold() in {"overvotes", "undervotes"}:
                continue
            name = " ".join(str(name).split())
            count = _votes(total[column])
            if sum(_votes(row[column]) for row in county_rows) != count:
                raise ValueError(f"Wyoming county votes do not reconcile for {name}")
            write_in = name.casefold() == "write-ins"
            records.append({
                "contest": f"{office}-{party}", "candidate_id": name, "unit": "statewide",
                "name": name, "party": "" if write_in else party, "primary_party": party,
                "chamber": "Senate" if office == "United States Senator" else "House",
                "district": "S" if office == "United States Senator" else "00",
                "term": "regular", "write_in": write_in, "votes": count,
                "locator": f"{members[0]} / {sheet.title} column {column + 1}:{total_number}",
            })
    finally:
        workbook.close()
    output = _aggregate(source, records, path)
    if len({row["contest_id"] for row in output}) != 4:
        raise ValueError("Wyoming must include both federal offices and both major-party primaries")
    return output


def parse_iowa_text(text: str, path: Path, source: dict) -> list[dict]:
    heading = re.compile(
        r"United States (Senator|Representative District (\d+)) - (Rep|Dem|Lib)\."
    )
    section_boundary = re.compile(
        r"IOWA SECRETARY OF STATE\s+" + re.escape(source["date"][:4])
        + r" Primary Election CANVASS SUMMARY"
    )
    triples = re.compile(
        r"Election Day ([\d, ]+)\s*\nAbsentee ([\d, ]+)\s*\nTotal ([\d, ]+)"
    )
    records = []
    for match in heading.finditer(text):
        boundary = section_boundary.search(text, match.end())
        end = boundary.start() if boundary else len(text)
        section = text[match.end():end]
        if "Write-in Under Votes Over Votes Total" not in section:
            raise ValueError("Iowa federal section lacks the expected candidate/total header")
        header = section.split("Write-in Under Votes Over Votes Total", 1)[0]
        names = [
            " ".join(name.split()).strip().rstrip(",")
            for name in re.findall(r"(.*?),\s*\n(?:REP|DEM|LIB)\b", header)
        ]
        rows = list(triples.finditer(section))
        if not rows:
            raise ValueError("Iowa federal contest has no reconciled statewide total triple")
        last = rows[-1]
        values = [
            [int(value.replace(",", "")) for value in line.split()]
            for line in last.groups()
        ]
        expected_columns = len(names) + 4
        if any(len(row) != expected_columns for row in values):
            raise ValueError("Iowa candidate count and total-column count differ")
        if any(day + absentee != total for day, absentee, total in zip(*values, strict=True)):
            raise ValueError("Iowa statewide election-day/absentee totals do not reconcile")
        if any(sum(row[:-1]) != row[-1] for row in values):
            raise ValueError("Iowa candidate, blank and overvote totals do not reconcile")
        party = {"Rep": "Republican", "Dem": "Democratic", "Lib": "Libertarian"}[match[3]]
        chamber = "House" if match[2] else "Senate"
        district = match[2].zfill(2) if match[2] else "S"
        before_start = list(re.finditer(r"\[PDF page (\d+)\]", text[:match.start()]))
        before_end = list(re.finditer(r"\[PDF page (\d+)\]", text[:match.end() + last.end()]))
        if not before_start or not before_end:
            raise ValueError("Iowa extraction lacks PDF page locators")
        locator = f"PDF pages {before_start[-1][1]}-{before_end[-1][1]}; statewide TOTAL"
        for index, name in enumerate([*names, "Write-in"]):
            write_in = index == len(names)
            records.append({
                "contest": match.group(0), "candidate_id": name, "unit": "statewide",
                "name": name, "party": "" if write_in else party, "primary_party": party,
                "chamber": chamber, "district": district, "term": "regular",
                "write_in": write_in, "votes": values[2][index], "locator": locator,
            })
    output = _aggregate(source, records, path)
    if len({row["contest_id"] for row in output}) != source["expected_contests"]:
        raise ValueError("Iowa federal primary contest coverage is incomplete")
    return output


def parse_iowa(path: Path, source: dict) -> list[dict]:
    _, text = _extract_pdf_text(path.read_bytes())
    if not text.strip():
        raise ValueError("Iowa canvass contains no machine-readable text")
    return parse_iowa_text(text, path, source)


def parse_missouri(path: Path, source: dict) -> list[dict]:
    _, text = _extract_pdf_text(path.read_bytes())
    if "Official Election Returns" not in text:
        raise ValueError("Missouri source is not an official canvass")
    headings = re.compile(
        r"U\.S\. (Senator|Representative - District (\d+)) "
        r"\((\d+) of (\d+) Precincts Reported\)"
    )
    any_heading = re.compile(r"^[^\n]+\(\d+ of \d+ Precincts Reported\)", re.M)
    candidate_line = re.compile(
        r"^(.+?) (Republican|Democratic|Libertarian|Green|Constitution) "
        r"([\d,]+) [\d.]+%$", re.M,
    )
    records = []
    for heading in headings.finditer(text):
        if heading[3] != heading[4]:
            raise ValueError("Missouri federal contest is not fully reported")
        boundary = any_heading.search(text, heading.end())
        section = text[heading.end():boundary.start() if boundary else len(text)]
        candidates = list(candidate_line.finditer(section))
        total = re.search(r"Total Votes ([\d,]+)", section)
        if not candidates or not total:
            raise ValueError("Missouri federal section lacks candidate rows or total votes")
        if sum(int(row[3].replace(",", "")) for row in candidates) != int(total[1].replace(",", "")):
            raise ValueError("Missouri candidate votes do not reconcile to the official total")
        party_totals = re.findall(r"Party Total ([\d,]+)", section)
        parties_in_order = list(dict.fromkeys(row[2] for row in candidates))
        if len(party_totals) != len(parties_in_order):
            raise ValueError("Missouri party subtotal count differs from candidate parties")
        for party, subtotal in zip(parties_in_order, party_totals, strict=True):
            if sum(int(row[3].replace(",", "")) for row in candidates if row[2] == party) != int(subtotal.replace(",", "")):
                raise ValueError("Missouri party candidate totals do not reconcile")
        for row in candidates:
            pages = list(re.finditer(r"\[PDF page (\d+)\]", text[:heading.end() + row.end()]))
            if not pages:
                raise ValueError("Missouri candidate row lacks a PDF page locator")
            name = re.sub(r"^\[PDF page \d+\]\s*", "", row[1]).strip()
            records.append({
                "contest": heading.group(0) + "-" + row[2],
                "candidate_id": name, "unit": "statewide", "name": name, "party": row[2],
                "chamber": "House" if heading[2] else "Senate",
                "district": heading[2].zfill(2) if heading[2] else "S",
                "term": "regular", "write_in": name.casefold() in {"write-in", "write-ins"},
                "votes": int(row[3].replace(",", "")), "locator": f"PDF page {pages[-1][1]}",
            })
    output = _aggregate(source, records, path)
    if {row["district"] for row in output if row["chamber"] == "House"} != {
        f"{number:02}" for number in range(1, source["expected_house_districts"] + 1)
    }:
        raise ValueError("Missouri source is missing a congressional district")
    return output


def parse_kansas(path: Path, source: dict) -> list[dict]:
    _, text = _extract_pdf_text(path.read_bytes())
    if f'{source["date"][:4]} Primary Election' not in text or "Official Vote Totals" not in text:
        raise ValueError("Kansas source is not the expected official primary canvass")
    marker = source["federal_section_end"]
    if marker not in text:
        raise ValueError("Kansas federal-section boundary changed")
    federal = text.split(marker, 1)[0]
    heading = re.compile(r"^United States (Senate|House of Representatives (\d+))$", re.M)
    sections = list(heading.finditer(federal))
    records = []
    for index, match in enumerate(sections):
        end = sections[index + 1].start() if index + 1 < len(sections) else len(federal)
        body = federal[match.end():end]
        candidates = list(re.finditer(r"^([DR])-(.+?) ([\d,]+) [\d.]+%$", body, re.M))
        if not candidates:
            raise ValueError("Kansas federal contest has no candidate rows")
        for candidate in candidates:
            party = {"D": "Democratic", "R": "Republican"}[candidate[1]]
            pages = list(re.finditer(r"\[PDF page (\d+)\]", text[:match.end() + candidate.start()]))
            records.append({
                "contest": match.group(0) + "-" + party,
                "candidate_id": candidate[2], "unit": "statewide",
                "name": candidate[2], "party": party,
                "chamber": "House" if match[2] else "Senate",
                "district": match[2].zfill(2) if match[2] else "S",
                "term": "regular", "write_in": False,
                "votes": int(candidate[3].replace(",", "")),
                "locator": f"PDF page {pages[-1][1]}",
            })
    output = _aggregate(source, records, path)
    if len(output) != source["expected_ballot_options"]:
        raise ValueError("Kansas federal candidate count differs from the reviewed canvass")
    if len({row["contest_id"] for row in output}) != source["expected_contests"]:
        raise ValueError("Kansas federal district/party coverage is incomplete")
    return output


def export_ranked_choice_rounds(
    sources: list[dict], results: list[dict], raw_dir: Path, output_dir: Path,
) -> None:
    path = MANUAL_DIR / "congressional_rcv_rounds.csv"
    if not path.exists():
        return
    rows = read_csv(path)
    sources_by_id = {source["source_id"]: source for source in sources}
    grouped = defaultdict(lambda: defaultdict(list))
    for row in rows:
        source = sources_by_id[row["source_id"]]
        expected_hash = hashlib.sha256((raw_dir / source["filename"]).read_bytes()).hexdigest()
        if row["source_sha256"] != expected_hash or row["election_date"] != source["date"]:
            raise ValueError("Ranked-choice rounds do not match their registered source")
        _votes(row["votes"])
        grouped[row["source_id"]][int(row["round"])].append(row)
    for source_id, rounds in grouped.items():
        if set(rounds) != set(range(1, len(rounds) + 1)):
            raise ValueError("Ranked-choice rounds are not contiguous")
        totals = {sum(_votes(row["votes"]) for row in group) for group in rounds.values()}
        option_sets = [set(row["ballot_option"] for row in group) for group in rounds.values()]
        if len(totals) != 1 or any(options != option_sets[0] for options in option_sets):
            raise ValueError("Ranked-choice round totals or ballot options do not reconcile")
        if any(len(options) != len(group) for options, group in zip(option_sets, rounds.values())):
            raise ValueError("Duplicate ranked-choice ballot option")
        primary = {
            row["candidate_name"]: _votes(row["votes"])
            for row in results if row["source_id"] == source_id
        }
        first = {
            row["ballot_option"]: _votes(row["votes"])
            for row in rounds[1] if row["ballot_option"] != "Exhausted Ballots"
        }
        if primary != first:
            raise ValueError("Primary first preferences differ from the final RCV report")
        final = rounds[max(rounds)]
        winners = [row for row in final if row["is_elected"] == "true"]
        exhausted = next(_votes(row["votes"]) for row in final if row["ballot_option"] == "Exhausted Ballots")
        if len(winners) != 1 or _votes(winners[0]["votes"]) <= (next(iter(totals)) - exhausted) / 2:
            raise ValueError("Declared ranked-choice winner lacks a final-round majority")
    if rows:
        write_csv(output_dir / "ranked_choice_rounds.csv", rows, list(rows[0]))


def collect_state_primaries(*, download: bool = False) -> dict:
    config = read_json(CONFIG_DIR / "congressional_state_sources.json")
    sources = config["sources"]
    raw_dir = RAW_DIR / "state_results"
    output_dir = ANALYSIS_DATA_DIR / "congressional"
    raw_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    manifest = []
    adapters = {
        "california": parse_california, "florida": parse_florida,
        "north_carolina": parse_north_carolina,
        "virginia": parse_virginia,
        "reviewed_totals": parse_reviewed_totals,
        "reviewed_contests": parse_reviewed_contests,
        "maine_totals": parse_maine_totals,
        "rhode_island": parse_rhode_island,
        "civera": parse_civera, "alabama": parse_alabama,
        "alabama_replacement_summary": parse_alabama_replacement_summary,
        "washington": parse_washington,
        "tally": parse_tally,
        "montana": parse_montana,
        "statewide_export": parse_statewide_export,
        "hawaii": parse_hawaii,
        "dc": parse_dc,
        "delaware": parse_delaware, "wyoming": parse_wyoming,
        "iowa": parse_iowa,
        "missouri": parse_missouri,
        "kansas": parse_kansas,
    }
    for source in sources:
        if "search" in source:
            source = {
                **source, "url": source["url"] + "?" + urlencode({
                    "search": json.dumps(source["search"], separators=(",", ":")),
                }),
            }
        if source["adapter"] == "virginia":
            source = _virginia_report_source(source, raw_dir, download)
        if source["adapter"] in {"tally", "dc"} and download:
            for field in (("info", "contests") if source["adapter"] == "tally" else ("info",)):
                payload = _download_workbook(source[f"{field}_url"])
                json.loads(payload)
                (raw_dir / source[f"{field}_filename"]).write_bytes(payload)
        path = raw_dir / source["filename"]
        if download:
            temporary = path.with_name(path.stem + ".download" + path.suffix)
            payload = (
                _download_webform(source) if "form_fields" in source
                else _download_workbook(source["url"])
            )
            if source.get("storage_compression") == "gzip":
                payload = gzip.compress(payload, mtime=0)
            temporary.write_bytes(payload)
            try:
                rows = adapters[source["adapter"]](temporary, source)
                temporary.replace(path)
            finally:
                temporary.unlink(missing_ok=True)
        else:
            rows = adapters[source["adapter"]](path, source)
        results.extend(rows)
        manifest.append({
            "source_id": source["source_id"], "state_code": source["state"],
            "election_date": source["date"], "source_url": source["url"],
            "publication_url": source["publication"],
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "candidate_rows": len(rows),
            "contests": len({row["contest_id"] for row in rows}),
            "coverage": "reported_primary_returns_only; uncontested nominations separate",
        })
    providers = []
    for provider in config.get("supplemental_outputs", []):
        if download:
            subprocess.run([sys.executable, "-m", provider["regenerate_module"]], check=True)
        path = output_dir / provider["results"]
        supplemental = read_csv(path)
        for row in supplemental:
            if set(row) != set(FIELDS):
                raise ValueError(f"{path.name}: supplemental schema differs from canonical results")
            if row["stage"] not in PRIMARY_RESULT_STAGES:
                raise ValueError(
                    f"{path.name}: stage {row['stage']!r} is not a canonical primary result stage; "
                    "rosters/general-election nominees cannot fill primary result gaps"
                )
            _votes(row["votes"])
        manifest_path = output_dir / provider["source_manifest"]
        if not manifest_path.exists():
            raise ValueError(f"Missing supplemental source manifest: {manifest_path}")
        results.extend(supplemental)
        providers.append({
            "provider": provider["provider"], "results_file": provider["results"],
            "results_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "source_manifest": provider["source_manifest"],
            "manifest_sha256": hashlib.sha256(manifest_path.read_bytes()).hexdigest(),
            "candidate_rows": len(supplemental),
            "contests": len({row["contest_id"] for row in supplemental}),
        })
    cutoff = datetime.strptime(
        read_json(CONFIG_DIR / "sources.json")["research_cutoff"], "%Y-%m-%d",
    ).date()
    for row in results:
        if datetime.strptime(row["election_date"], "%Y-%m-%d").date() > cutoff:
            raise ValueError(f"{row['source_id']}: primary results postdate research cutoff {cutoff}")
    keys = [
        (row["source_id"], row["contest_id"], row["candidate_source_id"])
        for row in results
    ]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate candidate/result source identity across primary providers")
    export_ranked_choice_rounds(sources, results, raw_dir, output_dir)
    write_csv(output_dir / "state_primary_results.csv", results, FIELDS)
    write_csv(output_dir / "state_primary_sources.csv", manifest, list(manifest[0]))
    write_csv(
        output_dir / "state_primary_providers.csv", providers,
        ["provider", "results_file", "results_sha256", "source_manifest",
         "manifest_sha256", "candidate_rows", "contests"],
    )
    return {
        "candidate_rows": len(results), "source_count": len(sources),
        "contests": len({row["contest_id"] for row in results}),
    }
