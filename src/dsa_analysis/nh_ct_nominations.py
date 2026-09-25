"""Recover New Hampshire results and Connecticut nomination accounting."""

import argparse
import csv
import hashlib
import re
from pathlib import Path

from openpyxl import load_workbook

from .io import write_csv
from .paths import ANALYSIS_DATA_DIR, RAW_DIR
from .state_results import FIELDS

RAW_SUBDIR = "nh_ct_nominations"
RESULT_FILENAME = "nh_ct_primary_results.csv"

PARTY_CONTEST_FIELDS = [
    "cycle", "election_date", "state_code", "chamber", "district", "term",
    "stage", "primary_party", "status", "candidate_names", "source_urls",
    "locators", "notes",
]

NH_2024_SOURCES = (
    {
        "source_id": "nh-2024-cd01",
        "district": "01",
        "filename": (
            "new_hampshire/"
            "2024-sp-congressional-district-1-democratic_4.xlsx"
        ),
        "url": (
            "https://www.sos.nh.gov/sites/g/files/ehbemt561/files/"
            "inline-documents/sonh/"
            "2024-sp-congressional-district-1-democratic_4.xlsx"
        ),
        "archive_url": (
            "https://web.archive.org/web/20250613142834id_/"
            "https://www.sos.nh.gov/sites/g/files/ehbemt561/files/"
            "inline-documents/sonh/"
            "2024-sp-congressional-district-1-democratic_4.xlsx"
        ),
    },
    {
        "source_id": "nh-2024-cd02",
        "district": "02",
        "filename": (
            "new_hampshire/"
            "2024-sp-congressional-district-2-democratic_4.xlsx"
        ),
        "url": (
            "https://www.sos.nh.gov/sites/g/files/ehbemt561/files/"
            "inline-documents/sonh/"
            "2024-sp-congressional-district-2-democratic_4.xlsx"
        ),
        "archive_url": (
            "https://web.archive.org/web/20251219043444id_/"
            "https://www.sos.nh.gov/sites/g/files/ehbemt561/files/"
            "inline-documents/sonh/"
            "2024-sp-congressional-district-2-democratic_4.xlsx"
        ),
    },
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _votes(value: object) -> int:
    if value is None:
        return 0
    number = int(value)
    if number < 0:
        raise ValueError(f"Invalid vote total: {value!r}")
    return number


def parse_new_hampshire(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = list(next(sheet.iter_rows(min_row=3, max_row=3, values_only=True)))
        total_row = None
        total_number = None
        for number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            if row[0] == "Totals":
                total_row = list(row)
                total_number = number
                break
    finally:
        workbook.close()
    if total_row is None:
        raise ValueError(f"New Hampshire workbook lacks Totals row: {path.name}")
    output = []
    source_hash = _sha256(path)
    for column in range(1, len(headers)):
        label = str(headers[column] or "").strip()
        match = re.fullmatch(r"(.+?),\s*([dr])", label, re.I)
        if match is None:
            continue
        name = match[1].strip()
        party = "Democratic" if match[2].lower() == "d" else "Republican"
        output.append({
            "contest_id": (
                f"nh-2024-09-10-house-{source['district']}-regular-"
                f"{party.lower()}"
            ),
            "cycle": "2024",
            "election_date": "2024-09-10",
            "stage": "primary",
            "state_code": "NH",
            "chamber": "House",
            "district": source["district"],
            "term": "regular",
            "election_system": "partisan_primary",
            "primary_party": party,
            "candidate_name": name,
            "candidate_source_id": f"{source['source_id']}-column-{column + 1}",
            "party": party,
            "write_in": False,
            "votes": _votes(total_row[column]),
            "source_id": source["source_id"],
            "source_url": source["archive_url"],
            "source_sha256": source_hash,
            "source_locators": f"{sheet.title}:R3C{column + 1},R{total_number}C{column + 1}",
            "reporting_units": 1,
        })
    if {row["primary_party"] for row in output} != {"Democratic", "Republican"}:
        raise ValueError("New Hampshire workbook does not contain both party primaries")
    return output


def _ct_2024_nominations(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _accounting(
    nh_results: list[dict],
    ct_nominations: list[dict],
    northeast_results: list[dict],
) -> list[dict]:
    rows = []
    groups = {}
    for result in nh_results:
        key = (result["district"], result["primary_party"])
        groups.setdefault(key, []).append(result)
    for (district, party), group in groups.items():
        rows.append({
            "cycle": "2024", "election_date": "2024-09-10",
            "state_code": "NH", "chamber": "House", "district": district,
            "term": "regular", "stage": "primary", "primary_party": party,
            "status": "results_complete",
            "candidate_names": " | ".join(sorted(row["candidate_name"] for row in group)),
            "source_urls": group[0]["source_url"],
            "locators": " | ".join(row["source_locators"] for row in group),
            "notes": "Archived copy of the official Secretary of State workbook.",
        })
    for party in ("Democratic", "Republican"):
        for chamber, district in (
            ("Senate", "S"), ("House", "01"), ("House", "02"),
        ):
            rows.append({
                "cycle": "2026", "election_date": "2026-09-08",
                "state_code": "NH", "chamber": chamber, "district": district,
                "term": "regular", "stage": "primary", "primary_party": party,
                "status": "unresolved", "candidate_names": "",
                "source_urls": (
                    f"https://www.sos.nh.gov/2026-{party.lower()}-state-primary"
                ),
                "locators": "",
                "notes": (
                    "Official page was archived before federal result files were linked; "
                    "live page access returned a blocker and no CAPTCHA was solved."
                ),
            })
    ct_results = {
        (
            row["cycle"], row["chamber"], row["district"], row["primary_party"],
        ): row
        for row in northeast_results if row["state_code"] == "CT"
    }
    ct_result_groups = {}
    for row in northeast_results:
        if row["state_code"] != "CT":
            continue
        key = (
            row["cycle"], row["chamber"], row["district"], row["primary_party"],
        )
        ct_result_groups.setdefault(key, []).append(row)
    for key, group in ct_result_groups.items():
        first = group[0]
        rows.append({
            "cycle": first["cycle"], "election_date": first["election_date"],
            "state_code": "CT", "chamber": first["chamber"],
            "district": first["district"], "term": first["term"],
            "stage": first["stage"], "primary_party": first["primary_party"],
            "status": "results_complete",
            "candidate_names": " | ".join(
                sorted(row["candidate_name"] for row in group)
            ),
            "source_urls": first["source_url"],
            "locators": " | ".join(row["source_locators"] for row in group),
            "notes": "Numeric result is provided by the integrated Connecticut source.",
        })
    for row in ct_nominations:
        key = (
            row["cycle"], row["chamber"], row["district"], row["primary_party"],
        )
        if key in ct_results:
            result = ct_results[key]
            rows.append({
                "cycle": row["cycle"], "election_date": row["election_date"],
                "state_code": "CT", "chamber": row["chamber"],
                "district": row["district"], "term": "regular", "stage": "primary",
                "primary_party": row["primary_party"], "status": "results_complete",
                "candidate_names": "", "source_urls": result["source_url"],
                "locators": result["source_locators"],
                "notes": "Numeric result is provided by the integrated Connecticut source.",
            })
        else:
            rows.append({
                "cycle": row["cycle"], "election_date": row["election_date"],
                "state_code": "CT", "chamber": row["chamber"],
                "district": row["district"], "term": "regular", "stage": "primary",
                "primary_party": row["primary_party"],
                "status": "unopposed_nomination_verified",
                "candidate_names": row["candidate_name"],
                "source_urls": row["source_url"],
                "locators": row["source_locator"],
                "notes": (
                    "Official general-election nominee plus complete primary export "
                    "with no contest row; no numeric primary vote is fabricated."
                ),
            })
    for district in range(1, 6):
        for party in ("Democratic", "Republican"):
            rows.append({
                "cycle": "2026", "election_date": "2026-08-11",
                "state_code": "CT", "chamber": "House",
                "district": f"{district:02}", "term": "regular", "stage": "primary",
                "primary_party": party, "status": "unresolved",
                "candidate_names": "",
                "source_urls": (
                    "https://portal.ct.gov/sots/election-services/"
                    "certificate-of-endorsement/2026-certificate-of-endorsements"
                ),
                "locators": f"ct-2026-cd{district}-endorsements.pdf",
                "notes": (
                    "Endorsement/15-percent forms are captured, but final primary-ballot "
                    "qualification cannot be inferred from the combined forms alone."
                ),
            })
    rows.append({
        "cycle": "2026", "election_date": "2026-08-11",
        "state_code": "CT", "chamber": "Senate", "district": "S",
        "term": "regular", "stage": "primary", "primary_party": "",
        "status": "not_scheduled", "candidate_names": "",
        "source_urls": "https://portal.ct.gov/sots/election-services/",
        "locators": "", "notes": "No U.S. Senate seat was scheduled in Connecticut in 2026.",
    })
    unique = {}
    for row in rows:
        key = tuple(row[field] for field in PARTY_CONTEST_FIELDS[:8])
        unique[key] = row
    return sorted(unique.values(), key=lambda row: tuple(
        row[field] for field in PARTY_CONTEST_FIELDS[:8]
    ))


def collect_nh_ct(
    raw_dir: Path | None = None,
    output_dir: Path | None = None,
) -> dict:
    raw_dir = raw_dir or RAW_DIR / RAW_SUBDIR
    output_dir = output_dir or ANALYSIS_DATA_DIR / "congressional"
    ledger = output_dir / "nh_ct"
    ledger.mkdir(parents=True, exist_ok=True)
    results = []
    for source in NH_2024_SOURCES:
        results.extend(parse_new_hampshire(raw_dir / source["filename"], source))
    results.sort(key=lambda row: (
        row["district"], row["primary_party"], row["candidate_name"],
    ))
    write_csv(output_dir / RESULT_FILENAME, results, FIELDS)
    ct_nominations = _ct_2024_nominations(
        raw_dir / "connecticut/ct-2024-verified-nominations.csv"
    )
    northeast_results_path = output_dir / "northeast_primary_results.csv"
    with northeast_results_path.open(newline="", encoding="utf-8") as handle:
        northeast_results = list(csv.DictReader(handle))
    accounting = _accounting(results, ct_nominations, northeast_results)
    write_csv(
        ledger / "party_contest_accounting.csv",
        accounting,
        PARTY_CONTEST_FIELDS,
    )
    manifest = []
    for source in NH_2024_SOURCES:
        path = raw_dir / source["filename"]
        subset = [row for row in results if row["source_id"] == source["source_id"]]
        manifest.append({
            "source_id": source["source_id"], "state_code": "NH",
            "election_date": "2024-09-10", "source_url": source["archive_url"],
            "publication_url": source["url"], "raw_filename": source["filename"],
            "sha256": _sha256(path), "candidate_rows": len(subset),
            "contests": len({row["contest_id"] for row in subset}),
            "status": "archived_official_workbook_recovered",
            "notes": "Workbook contains both Democratic and Republican candidate columns.",
        })
    manifest.append({
        "source_id": "ct-2024-nomination-evidence", "state_code": "CT",
        "election_date": "2024-08-13",
        "source_url": (
            "https://portal.ct.gov/-/media/sots/electionservices/2024/"
            "voter_guide/voter_guide_2024_general_election.pdf"
        ),
        "publication_url": "https://portal.ct.gov/sots/election-services/",
        "raw_filename": "connecticut/ct-2024-verified-nominations.csv",
        "sha256": _sha256(
            raw_dir / "connecticut/ct-2024-verified-nominations.csv"
        ),
        "candidate_rows": len(ct_nominations), "contests": len(ct_nominations),
        "status": "official_nomination_evidence",
        "notes": "Used only when the complete Civera primary export has no contest row.",
    })
    manifest.append({
        "source_id": "ct-2026-endorsement-evidence", "state_code": "CT",
        "election_date": "2026-08-11",
        "source_url": (
            "https://portal.ct.gov/sots/election-services/"
            "certificate-of-endorsement/2026-certificate-of-endorsements"
        ),
        "publication_url": (
            "https://portal.ct.gov/sots/election-services/"
            "certificate-of-endorsement/2026-certificate-of-endorsements"
        ),
        "raw_filename": "connecticut/ct-2026-endorsements.html",
        "sha256": _sha256(raw_dir / "connecticut/ct-2026-endorsements.html"),
        "candidate_rows": 0, "contests": 10,
        "status": "captured_but_final_qualification_unresolved",
        "notes": "Combined endorsement/15-percent forms are not treated as final nominations.",
    })
    for party in ("democratic", "republican"):
        path = raw_dir / f"new_hampshire/2026-{party}-state-primary.html"
        manifest.append({
            "source_id": f"nh-2026-{party}-page", "state_code": "NH",
            "election_date": "2026-09-08",
            "source_url": f"https://www.sos.nh.gov/2026-{party}-state-primary",
            "publication_url": f"https://www.sos.nh.gov/2026-{party}-state-primary",
            "raw_filename": f"new_hampshire/2026-{party}-state-primary.html",
            "sha256": _sha256(path), "candidate_rows": 0, "contests": 0,
            "status": "archived_before_federal_result_links",
            "notes": "Not used as vote data; live access blocker was not solved.",
        })
    write_csv(ledger / "source_manifest.csv", manifest, list(manifest[0]))
    return {
        "result_rows": len(results),
        "contests": len({row["contest_id"] for row in results}),
        "accounting_rows": len(accounting),
        "sources": len(manifest),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", type=Path)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    summary = collect_nh_ct(args.raw_dir, args.output_dir)
    print(" ".join(f"{key}={value}" for key, value in summary.items()))


if __name__ == "__main__":
    main()
