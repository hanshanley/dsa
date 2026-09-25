"""Recover Louisiana, Mississippi, and Oklahoma congressional primaries."""

import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .io import write_csv
from .paths import ANALYSIS_DATA_DIR, RAW_DIR
from .state_results import FIELDS

RAW_SUBDIR = "gulf_primaries"
RESULT_FILENAME = "gulf_primary_results.csv"
MANIFEST_SUBDIR = "gulf"
DATA_CUTOFF = "2026-09-22"

SOURCE_MANIFEST_FIELDS = [
    "source_id",
    "state_code",
    "election_date",
    "stage",
    "artifact_kind",
    "authority",
    "publication_url",
    "source_url",
    "raw_filename",
    "source_date",
    "sha256",
    "acquisition_status",
    "verification_status",
    "certification_status",
    "provisional_status",
    "candidate_rows",
    "contests",
    "coverage",
    "source_locators",
    "notes",
]

ATTEMPT_FIELDS = [
    "attempt_id",
    "attempted_on",
    "state_code",
    "cycle",
    "source_url",
    "method",
    "outcome",
    "raw_evidence",
    "notes",
]

GAP_FIELDS = [
    "gap_id",
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "primary_party",
    "election_system",
    "coverage_status",
    "missing_scope",
    "blocker",
    "next_authoritative_step",
    "source_urls",
]

CONTEST_FIELDS = [
    "contest_id",
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "term",
    "election_system",
    "primary_party",
    "candidate_count",
    "unopposed",
    "actual_zero_rows",
    "has_write_in",
    "result_status",
    "certification_status",
    "source_id",
]

ROSTER_FIELDS = [
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "term",
    "election_system",
    "primary_party",
    "candidate_name",
    "candidate_source_id",
    "source_id",
    "source_url",
    "source_sha256",
    "source_locators",
    "verification_status",
    "notes",
]

COVERAGE_FIELDS = [
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "term",
    "election_system",
    "primary_party",
    "coverage_status",
    "result_rows",
    "roster_rows",
    "missing_reason",
    "source_ids",
]

OBSERVATION_FIELDS = [
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "primary_party",
    "election_system",
    "observation_status",
    "canonical_requirement",
    "source_id",
    "source_url",
    "notes",
]

LA_PUBLICATION = "https://www.sos.la.gov/elections-voting/election-results-statistics"
LA_WORKBOOK_BASE = (
    "https://s3-us-west-2.amazonaws.com/mediaresults.sos.la.gov/"
    "HumanReadableElectionResults"
)
MS_PUBLICATION = "https://www.sos.ms.gov/elections-voting/election-results"
OK_PUBLICATION = "https://oklahoma.gov/elections/elections-results.html"
OK_RESULTS = "https://results.okelections.gov/OKER/"

LA_SOURCES = (
    {
        "source_id": "la-2024-open-congressional-primary",
        "date": "2024-11-05",
        "stage": "nonpartisan_primary",
        "system": "nonpartisan_open_primary",
        "filename": "la/2024/la-2024-11-05-open-primary-results.xlsx",
        "url": (
            f"{LA_WORKBOOK_BASE}/20241105/"
            "Election+Results+(11-05-2024).xlsx"
        ),
    },
    {
        "source_id": "la-2026-senate-party-primary",
        "date": "2026-05-16",
        "stage": "primary",
        "system": "partisan_primary",
        "filename": "la/2026/la-2026-05-16-party-primary-results.xlsx",
        "url": (
            f"{LA_WORKBOOK_BASE}/20260516/"
            "Election+Results+(05-16-2026).xlsx"
        ),
    },
    {
        "source_id": "la-2026-senate-second-party-primary",
        "date": "2026-06-27",
        "stage": "primary_runoff",
        "system": "partisan_primary",
        "filename": "la/2026/la-2026-06-27-second-party-primary-results.xlsx",
        "url": (
            f"{LA_WORKBOOK_BASE}/20260627/"
            "Election+Results+(06-27-2026).xlsx"
        ),
    },
)

MS_SOURCES = (
    {
        "source_id": "ms-2024-democratic-primary-recap",
        "date": "2024-03-12",
        "stage": "primary",
        "party": "Democratic",
        "filename": "ms/2024/ms-2024-democratic-primary-recap.pdf",
    },
    {
        "source_id": "ms-2024-republican-primary-recap",
        "date": "2024-03-12",
        "stage": "primary",
        "party": "Republican",
        "filename": "ms/2024/ms-2024-republican-primary-recap.pdf",
    },
    {
        "source_id": "ms-2024-republican-primary-runoff-recap",
        "date": "2024-04-02",
        "stage": "primary_runoff",
        "party": "Republican",
        "filename": "ms/2024/ms-2024-republican-primary-runoff-recap.pdf",
    },
    {
        "source_id": "ms-2026-democratic-primary-recap",
        "date": "2026-03-10",
        "stage": "primary",
        "party": "Democratic",
        "filename": "ms/2026/ms-2026-democratic-primary-recap.pdf",
    },
    {
        "source_id": "ms-2026-republican-primary-recap",
        "date": "2026-03-10",
        "stage": "primary",
        "party": "Republican",
        "filename": "ms/2026/ms-2026-republican-primary-recap.pdf",
    },
)

OK_SOURCES = (
    {
        "source_id": "ok-2024-june-primary-results",
        "date": "2024-06-18",
        "stage": "primary",
        "filename": "ok/2024/ok-2024-06-18-primary.json",
    },
    {
        "source_id": "ok-2024-august-runoff-results",
        "date": "2024-08-27",
        "stage": "primary_runoff",
        "filename": "ok/2024/ok-2024-08-27-runoff.json",
    },
    {
        "source_id": "ok-2026-june-primary-results",
        "date": "2026-06-16",
        "stage": "primary",
        "filename": "ok/2026/ok-2026-06-16-primary.json",
    },
    {
        "source_id": "ok-2026-august-runoff-results",
        "date": "2026-08-25",
        "stage": "primary_runoff",
        "filename": "ok/2026/ok-2026-08-25-runoff.json",
    },
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _integer(value: object) -> int:
    text = str(value).strip().replace(",", "")
    if re.fullmatch(r"\d+", text) is None:
        raise ValueError(f"Invalid vote total: {value!r}")
    return int(text)


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")


def _contest_id(
    state: str,
    date: str,
    chamber: str,
    district: str,
    term: str,
    party: str,
) -> str:
    return "-".join((
        state.casefold(),
        date,
        chamber.casefold(),
        district.casefold(),
        term,
        _slug(party or "all"),
    ))


def _la_candidate(value: str) -> tuple[str, str]:
    text = " ".join(value.split())
    mappings = (
        (" Democratic (DEM)", "Democratic"),
        (" Republican (REP)", "Republican"),
        (" No Party (NOPTY)", "No Party"),
        (" Libertarian (LBT)", "Libertarian"),
        (" Green (GRN)", "Green"),
    )
    for suffix, party in mappings:
        if text.endswith(suffix):
            return text.removesuffix(suffix), party
    match = re.fullmatch(r"(.+?) \(([^)]+)\)", text)
    if match:
        return match[1], {
            "DEM": "Democratic",
            "REP": "Republican",
            "NOPTY": "No Party",
            "LBT": "Libertarian",
            "GRN": "Green",
        }.get(match[2], match[2])
    raise ValueError(f"Unknown Louisiana candidate label: {value!r}")


def _la_heading(value: str) -> tuple[str, str, str] | None:
    senate = re.fullmatch(r"U\. S\. Senator -- (Democratic|Republican) Party", value)
    if senate:
        return "Senate", "S", senate[1]
    house = re.fullmatch(
        r"U\. S\. Representative -- (\d+)(?:st|nd|rd|th) Congressional District",
        value,
    )
    if house:
        return "House", house[1].zfill(2), ""
    return None


def parse_louisiana(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Multi-Parish(Parish)"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if len(rows) < 8 or str(rows[2][1]) != f"Election Date: {source['date']}":
        raise ValueError("Louisiana workbook contains a different election date")
    source_hash = _sha256(path)
    output = []
    for index, row in enumerate(rows):
        heading = str(row[0] or "") if row else ""
        parsed = _la_heading(heading)
        if parsed is None:
            continue
        chamber, district, heading_party = parsed
        if index + 2 >= len(rows):
            raise ValueError("Louisiana contest table is truncated")
        candidates = rows[index + 1][1:]
        totals = rows[index + 2][1:]
        if rows[index + 2][0] != "Total Votes":
            raise ValueError(f"Louisiana total row is missing: {heading}")
        for candidate_number, (label, votes) in enumerate(
            zip(candidates, totals, strict=False),
            1,
        ):
            if label is None:
                continue
            name, candidate_party = _la_candidate(str(label))
            primary_party = heading_party
            if primary_party and candidate_party != primary_party:
                raise ValueError(f"Louisiana primary party differs: {name}")
            write_in = "WRITE-IN" in name.upper()
            output.append({
                "contest_id": _contest_id(
                    "LA", source["date"], chamber, district, "regular", primary_party
                ),
                "cycle": source["date"][:4],
                "election_date": source["date"],
                "stage": source["stage"],
                "state_code": "LA",
                "chamber": chamber,
                "district": district,
                "term": "regular",
                "election_system": source["system"],
                "primary_party": primary_party,
                "candidate_name": name,
                "candidate_source_id": (
                    f"la-{source['date']}-{chamber.casefold()}-"
                    f"{district.casefold()}-{candidate_number}"
                ),
                "party": "" if write_in else candidate_party,
                "write_in": write_in,
                "votes": _integer(votes),
                "source_id": source["source_id"],
                "source_url": source["url"],
                "source_sha256": source_hash,
                "source_locators": (
                    f"{path.name}:{sheet.title}:row-{index + 2}:"
                    f"column-{candidate_number + 1}"
                ),
                "reporting_units": 1,
            })
    if not output:
        raise ValueError(f'{source["source_id"]}: no congressional candidates parsed')
    return output


def parse_mississippi(pdf_path: Path, source: dict) -> list[dict]:
    text_path = pdf_path.with_suffix(".txt")
    records = {}
    pages = defaultdict(set)
    current = None
    pending_continuation = None
    heading_candidate_count = 0
    page = 0
    for line_number, raw_line in enumerate(
        text_path.read_text(encoding="utf-8").splitlines(),
        1,
    ):
        line = " ".join(raw_line.split())
        page_match = re.fullmatch(r"PAGE (\d+)", line)
        if page_match:
            page = int(page_match[1])
            current = pending_continuation
            heading_candidate_count = 0
            continue
        if line == "United States-President":
            current = None
            heading_candidate_count = 0
            continue
        if line == "United States-Senate":
            current = ("Senate", "S")
            heading_candidate_count = 0
            continue
        house = re.fullmatch(r"US House Of Rep (\d+)-.*", line)
        if house:
            current = ("House", house[1].zfill(2))
            heading_candidate_count = 0
            continue
        if line == "Official Recapitulation":
            if current is not None and heading_candidate_count == 0:
                pending_continuation = current
            continue
        if current is None:
            continue
        candidate = re.match(r"(.+?) (Democrat|Republican) (.+)$", line)
        if candidate is None:
            continue
        party = "Democratic" if candidate[2] == "Democrat" else "Republican"
        if party != source["party"]:
            continue
        heading_candidate_count += 1
        numbers = [
            _integer(value)
            for value in re.findall(r"\b[\d,]+\b", candidate[3])
        ]
        if not numbers:
            continue
        key = (*current, candidate[1])
        records[key] = {
            "votes": numbers[-1],
            "page": page,
            "line": line_number,
        }
        pages[key].add(page)
    source_hash = _sha256(pdf_path)
    output = []
    contest_counts = defaultdict(int)
    for (chamber, district, name), record in records.items():
        contest = (chamber, district)
        contest_counts[contest] += 1
        output.append({
            "contest_id": _contest_id(
                "MS", source["date"], chamber, district, "regular", source["party"]
            ),
            "cycle": source["date"][:4],
            "election_date": source["date"],
            "stage": source["stage"],
            "state_code": "MS",
            "chamber": chamber,
            "district": district,
            "term": "regular",
            "election_system": "partisan_primary",
            "primary_party": source["party"],
            "candidate_name": name,
            "candidate_source_id": (
                f"ms-{source['date']}-{chamber.casefold()}-"
                f"{district.casefold()}-{source['party'].casefold()}-"
                f"{contest_counts[contest]}"
            ),
            "party": source["party"],
            "write_in": False,
            "votes": record["votes"],
            "source_id": source["source_id"],
            "source_url": (
                "https://www.sos.ms.gov/elections-voting/election-results"
            ),
            "source_sha256": source_hash,
            "source_locators": (
                f"{pdf_path.name}:page-{record['page']} | "
                f"{text_path.name}:line-{record['line']}"
            ),
            "reporting_units": 1,
        })
    if not output:
        raise ValueError(f'{source["source_id"]}: no federal candidates parsed')
    return output


def _ok_contest(race: dict) -> tuple[str, str, str] | None:
    title = race["raceTitle"]
    if title == "FOR UNITED STATES SENATOR":
        chamber, district = "Senate", "S"
    else:
        match = re.fullmatch(
            r"FOR UNITED STATES REPRESENTATIVE DISTRICT (\d+)",
            title,
        )
        if match is None:
            return None
        chamber, district = "House", match[1].zfill(2)
    party = {
        "Democrat": "Democratic",
        "Republican": "Republican",
        "Libertarian": "Libertarian",
    }.get(race["raceTitle2"])
    if party is None:
        return None
    return chamber, district, party


def parse_oklahoma(path: Path, source: dict) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("isOfficial") is not True:
        raise ValueError(f'{source["source_id"]}: Oklahoma results are not official')
    if payload["electionDate"][:10] != source["date"]:
        raise ValueError("Oklahoma result date differs")
    results = {result["raceID"]: result for result in payload["results"]}
    source_hash = _sha256(path)
    output = []
    for race_index, race in enumerate(payload["races"]):
        parsed = _ok_contest(race)
        if parsed is None:
            continue
        chamber, district, party = parsed
        result = results[race["raceID"]]
        candidates = race["raceCandidates"]
        candidate_results = result["candResults"]
        if len(candidates) != len(candidate_results):
            raise ValueError("Oklahoma candidate and result arrays differ")
        if result["reportingPrecincts"] != result["totalPrecincts"]:
            raise ValueError(f"Oklahoma contest is not fully reported: {race['raceID']}")
        for candidate_index, (candidate, votes) in enumerate(
            zip(candidates, candidate_results, strict=True),
            1,
        ):
            name = " ".join(candidate["candName"].split())
            write_in = "WRITE-IN" in name.upper()
            output.append({
                "contest_id": _contest_id(
                    "OK", source["date"], chamber, district, "regular", party
                ),
                "cycle": source["date"][:4],
                "election_date": source["date"],
                "stage": source["stage"],
                "state_code": "OK",
                "chamber": chamber,
                "district": district,
                "term": "regular",
                "election_system": "partisan_primary",
                "primary_party": party,
                "candidate_name": name,
                "candidate_source_id": f'{race["raceID"]}:{candidate_index}',
                "party": "" if write_in else party,
                "write_in": write_in,
                "votes": _integer(votes["totalVotes"]),
                "source_id": source["source_id"],
                "source_url": (
                    f"{OK_RESULTS}?elecDate={source['date'].replace('-', '')}"
                ),
                "source_sha256": source_hash,
                "source_locators": (
                    f"{path.name}:$.races[{race_index}].raceCandidates"
                    f"[{candidate_index - 1}] + $.results[raceID={race['raceID']}]"
                    f".candResults[{candidate_index - 1}]"
                ),
                "reporting_units": _integer(result["totalPrecincts"]),
            })
    return output


def _roster_row(
    *,
    state: str,
    date: str,
    chamber: str,
    district: str,
    party: str,
    name: str,
    candidate_id: str,
    source_id: str,
    source_url: str,
    source_hash: str,
    locator: str,
    notes: str,
) -> dict:
    return {
        "state_code": state,
        "cycle": date[:4],
        "election_date": date,
        "stage": "primary",
        "chamber": chamber,
        "district": district,
        "term": "regular",
        "election_system": "partisan_primary",
        "primary_party": party,
        "candidate_name": name,
        "candidate_source_id": candidate_id,
        "source_id": source_id,
        "source_url": source_url,
        "source_sha256": source_hash,
        "source_locators": locator,
        "verification_status": "official_roster_no_primary_result",
        "notes": notes,
    }


def parse_oklahoma_2024_roster(
    path: Path,
    actual_units: set[tuple[str, str, str]],
) -> tuple[list[dict], dict[tuple[str, str, str], str]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    groups = defaultdict(list)
    for row in data:
        if not row["OfficeDesc"].startswith("UNITED STATES REPRESENTATIVE"):
            continue
        match = re.search(r"DISTRICT (\d+)", row["OfficeDesc"])
        party = {
            "Democrat": "Democratic",
            "Republican": "Republican",
        }.get(row["CandidateAffiliation"])
        if match and party and not row["IsRemoved"]:
            groups[("House", match[1].zfill(2), party)].append(row)
    expected = {
        ("House", f"{district:02}", party)
        for district in range(1, 6)
        for party in ("Democratic", "Republican")
    }
    roster = []
    missing = {}
    source_hash = _sha256(path)
    for unit in sorted(expected - actual_units):
        candidates = groups.get(unit, [])
        if len(candidates) == 1:
            candidate = candidates[0]
            roster.append(_roster_row(
                state="OK",
                date="2024-06-18",
                chamber=unit[0],
                district=unit[1],
                party=unit[2],
                name=candidate["CandidateCombinedName"],
                candidate_id=str(candidate["CandidateID"]),
                source_id="ok-2024-official-candidate-roster",
                source_url=(
                    "https://filings.okelections.gov/ViewCandidates/"
                    "2024040320240405/99/all"
                ),
                source_hash=source_hash,
                locator=f"{path.name}:CandidateID={candidate['CandidateID']}",
                notes="Single filed, non-removed party candidate; no result contest.",
            ))
            missing[unit] = "unopposed_roster_only"
        elif not candidates:
            missing[unit] = "no_filed_party_candidate"
        else:
            missing[unit] = "candidate_status_requires_review"
    return roster, missing


def parse_oklahoma_2026_roster(path: Path, pdf_path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    source_hash = _sha256(pdf_path)
    output = []
    for row in rows:
        output.append(_roster_row(
            state="OK",
            date="2026-06-16",
            chamber=row["chamber"],
            district=row["district"],
            party=row["primary_party"],
            name=row["candidate_name"],
            candidate_id=row["candidate_source_id"],
            source_id="ok-2026-official-candidate-book",
            source_url=(
                "https://oklahoma.gov/content/dam/ok/en/elections/"
                "candidate-filing-archives/2026-candidate-filing-archives/"
                "2026-candidate-list-book.pdf"
            ),
            source_hash=source_hash,
            locator=f"{row['source_pdf']}:page-{row['page']}",
            notes=row["notes"],
        ))
    return output


def _unit(row: dict) -> tuple[str, str, str, str, str, str]:
    return (
        row["state_code"],
        row["election_date"],
        row["stage"],
        row["chamber"],
        row["district"],
        row["primary_party"],
    )


def _coverage_rows(
    results: list[dict],
    rosters: list[dict],
    ok_2024_missing: dict[tuple[str, str, str], str],
) -> list[dict]:
    result_groups = defaultdict(list)
    roster_groups = defaultdict(list)
    for row in results:
        result_groups[_unit(row)].append(row)
    for row in rosters:
        roster_groups[_unit(row)].append(row)
    expected = []
    expected.extend(
        (
            "LA", "2024-11-05", "nonpartisan_primary", "House",
            f"{district:02}", "regular", "nonpartisan_open_primary", "",
        )
        for district in range(1, 7)
    )
    expected.extend(
        (
            "LA", date, stage, "Senate", "S", "regular",
            "partisan_primary", party,
        )
        for date, stage in (
            ("2026-05-16", "primary"),
            ("2026-06-27", "primary_runoff"),
        )
        for party in ("Democratic", "Republican")
    )
    expected.extend(
        (
            "LA", "2026-11-03", "nonpartisan_primary", "House",
            f"{district:02}", "regular", "nonpartisan_open_primary", "",
        )
        for district in range(1, 7)
    )
    for year, date in (("2024", "2024-03-12"), ("2026", "2026-03-10")):
        expected.extend(
            (
                "MS", date, "primary", "Senate", "S", "regular",
                "partisan_primary", party,
            )
            for party in ("Democratic", "Republican")
        )
        expected.extend(
            (
                "MS", date, "primary", "House", f"{district:02}", "regular",
                "partisan_primary", party,
            )
            for district in range(1, 5)
            for party in ("Democratic", "Republican")
        )
    expected.append((
        "MS", "2024-04-02", "primary_runoff", "House", "02", "regular",
        "partisan_primary", "Republican",
    ))
    expected.extend(
        (
            "OK", "2024-06-18", "primary", "House", f"{district:02}",
            "regular", "partisan_primary", party,
        )
        for district in range(1, 6)
        for party in ("Democratic", "Republican")
    )
    expected.extend(
        (
            "OK", "2026-06-16", "primary", "House", f"{district:02}",
            "regular", "partisan_primary", party,
        )
        for district in range(1, 6)
        for party in ("Democratic", "Republican")
    )
    expected.extend((
        "OK", "2026-06-16", "primary", "Senate", "S", "regular",
        "partisan_primary", party,
    ) for party in ("Democratic", "Republican", "Libertarian"))
    expected.append((
        "OK", "2026-08-25", "primary_runoff", "Senate", "S", "regular",
        "partisan_primary", "Democratic",
    ))
    expected.append((
        "OK", "2026-08-25", "primary_runoff", "House", "01", "regular",
        "partisan_primary", "Republican",
    ))
    output = []
    for state, date, stage, chamber, district, term, system, party in expected:
        key = (state, date, stage, chamber, district, party)
        result_rows = result_groups.get(key, [])
        roster_rows = roster_groups.get(key, [])
        if result_rows:
            status, reason = "actual_results", ""
        elif roster_rows:
            status, reason = (
                "unopposed_roster_only",
                "official candidate roster exists but no primary result contest",
            )
        elif state == "LA" and date == "2026-11-03":
            status, reason = (
                "future_scheduled",
                "election is scheduled after the September 22, 2026 data cutoff",
            )
        elif (
            state == "OK"
            and date == "2026-08-25"
            and chamber == "House"
            and district == "01"
        ):
            status, reason = (
                "runoff_cancelled_after_withdrawal",
                "second-place candidate withdrew after the June primary",
            )
        elif state == "OK" and date == "2024-06-18":
            status = ok_2024_missing.get(
                (chamber, district, party),
                "candidate_status_requires_review",
            )
            reason = status.replace("_", " ")
        else:
            status, reason = (
                "no_published_primary_result",
                "no official result contest or contemporaneous roster recovered",
            )
        source_ids = " | ".join(sorted({
            row["source_id"] for row in result_rows + roster_rows
        }))
        if state == "LA" and date == "2026-11-03":
            source_ids = (
                "la-2026-house-primary-postponement-authority | "
                "fec-2026-congressional-primary-calendar-louisiana"
            )
        output.append({
            "state_code": state,
            "cycle": date[:4],
            "election_date": date,
            "stage": stage,
            "chamber": chamber,
            "district": district,
            "term": term,
            "election_system": system,
            "primary_party": party,
            "coverage_status": status,
            "result_rows": len(result_rows),
            "roster_rows": len(roster_rows),
            "missing_reason": reason,
            "source_ids": source_ids,
        })
    return output


def _assert_future_scheduled_coverage(
    coverage: list[dict],
    results: list[dict],
) -> None:
    future = [
        row for row in coverage
        if row["state_code"] == "LA"
        and row["cycle"] == "2026"
        and row["chamber"] == "House"
        and row["coverage_status"] == "future_scheduled"
    ]
    if len(future) != 6:
        raise ValueError("Louisiana future House coverage must contain six districts")
    if any(row["election_date"] <= DATA_CUTOFF for row in future):
        raise ValueError("future_scheduled rows must occur after the data cutoff")
    future_units = {
        (
            row["state_code"],
            row["election_date"],
            row["stage"],
            row["chamber"],
            row["district"],
            row["primary_party"],
        )
        for row in future
    }
    result_units = {_unit(row) for row in results}
    if future_units & result_units:
        raise ValueError("future_scheduled units cannot have numeric vote records")


def _update_party_accounting_future_status(
    analysis_root: Path,
    results: list[dict],
) -> None:
    path = analysis_root / MANIFEST_SUBDIR / "party_contest_accounting.csv"
    if not path.exists():
        return
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    future_rows = [
        row for row in rows
        if row["state_code"] == "LA"
        and row["cycle"] == "2026"
        and row["election_date"] == "2026-11-03"
        and row["chamber"] == "House"
        and row["stage"] == "nonpartisan_primary"
    ]
    if len(future_rows) != 6:
        raise ValueError("Party accounting must contain six Louisiana House rows")
    result_units = {_unit(row) for row in results}
    for row in future_rows:
        if row["election_date"] <= DATA_CUTOFF:
            raise ValueError("future_scheduled accounting row is not future-dated")
        unit = (
            row["state_code"],
            row["election_date"],
            row["stage"],
            row["chamber"],
            row["district"],
            row["primary_party"],
        )
        if unit in result_units:
            raise ValueError("future_scheduled accounting row has numeric votes")
        row["status"] = "future_scheduled"
        row["notes"] = row["notes"].replace(
            "status enum represented as not_scheduled. ",
            "",
        )
    provenance_updates = (
        {
            "criteria": {
                "state_code": "OK",
                "cycle": "2024",
                "election_date": "2024-06-18",
                "stage": "primary",
                "chamber": "House",
                "district": "03",
                "primary_party": "Democratic",
            },
            "status": "no_candidate_verified",
            "source_urls": (
                "https://filings.okelections.gov/ViewCandidates/"
                "2024040320240405/99/all"
            ),
            "locators": (
                "ok-2024-candidate-roster.json:complete 606-row official filing "
                "export; OfficeDesc='UNITED STATES REPRESENTATIVE - DISTRICT 03'; "
                "Republican CandidateIDs 44514,44773,44932; no Democrat party row"
            ),
            "notes": (
                "Verified against the complete official candidate filing export; "
                "District 3 contains three Republican filings and no Democratic filing."
            ),
        },
        {
            "criteria": {
                "state_code": "NJ",
                "cycle": "2026",
                "election_date": "2026-06-02",
                "stage": "primary",
                "chamber": "House",
                "district": "08",
                "primary_party": "Republican",
            },
            "status": "no_candidate_verified",
            "source_urls": (
                "https://www.nj.gov/state/elections/assets/pdf/election-results/"
                "2026/2026-official-primary-candidates-us-house.pdf"
            ),
            "locators": (
                "2026-official-primary-candidates-us-house.pdf:page-10, "
                "Eighth Congressional District section; only Democratic candidates "
                "Rob Menendez and Mussab Ali are listed"
            ),
            "notes": (
                "Official primary candidate PDF explicitly lists the District 8 "
                "candidate section without a Republican candidate."
            ),
        },
        {
            "criteria": {
                "state_code": "OK",
                "cycle": "2026",
                "election_date": "2026-08-25",
                "stage": "primary_runoff",
                "chamber": "House",
                "district": "01",
                "primary_party": "Republican",
            },
            "status": "no_primary_verified",
            "source_urls": (
                "https://oklahoma.gov/elections/candidates/"
                "2026-candidate-filing-information/2026-candidate-withdrawals.html"
                " | https://results.okelections.gov/OKER/?elecDate=20260825"
            ),
            "locators": (
                "ok-2026-candidate-withdrawals.html:withdrawal table row Number 522, "
                "Jackson Lahmeyer, U.S. Representative District 1, 6/19/2026 | "
                "ok-2026-08-25-runoff.json:$.races; no U.S. Representative "
                "District 1 contest is present"
            ),
            "notes": (
                "Official withdrawal record and official August runoff payload "
                "jointly verify that no House District 1 Republican runoff occurred."
            ),
        },
    )
    for update in provenance_updates:
        matches = [
            row for row in rows
            if all(row.get(field) == value for field, value in update["criteria"].items())
        ]
        if matches:
            if len(matches) != 1:
                raise ValueError("Canonical provenance update matched multiple rows")
            row = matches[0]
            row["status"] = update["status"]
            row["source_urls"] = update["source_urls"]
            row["locators"] = update["locators"]
            row["notes"] = update["notes"]
    resolved_statuses = {
        "results_complete",
        "unopposed_nomination_verified",
        "no_primary_verified",
        "no_candidate_verified",
        "future_scheduled",
    }
    missing_provenance = [
        row for row in rows
        if row["status"] in resolved_statuses
        and (not row.get("source_urls") or not row.get("locators"))
    ]
    if missing_provenance:
        raise ValueError(
            f"Resolved accounting rows lack provenance: {missing_provenance}"
        )
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=rows[0].keys(),
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def _contest_rows(rows: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in rows:
        groups[row["contest_id"]].append(row)
    output = []
    for contest_id, candidates in sorted(groups.items()):
        first = candidates[0]
        output.append({
            "contest_id": contest_id,
            "state_code": first["state_code"],
            "cycle": first["cycle"],
            "election_date": first["election_date"],
            "stage": first["stage"],
            "chamber": first["chamber"],
            "district": first["district"],
            "term": first["term"],
            "election_system": first["election_system"],
            "primary_party": first["primary_party"],
            "candidate_count": len(candidates),
            "unopposed": len(candidates) == 1,
            "actual_zero_rows": sum(candidate["votes"] == 0 for candidate in candidates),
            "has_write_in": any(candidate["write_in"] for candidate in candidates),
            "result_status": "actual_numeric",
            "certification_status": "official",
            "source_id": first["source_id"],
        })
    return output


def _gap_rows(coverage: list[dict]) -> list[dict]:
    output = []
    for row in coverage:
        if row["coverage_status"] in {
            "actual_results",
            "unopposed_roster_only",
            "future_scheduled",
            "runoff_cancelled_after_withdrawal",
        }:
            continue
        output.append({
            "gap_id": "-".join((
                row["state_code"].casefold(),
                row["election_date"],
                row["stage"],
                row["chamber"].casefold(),
                row["district"].casefold(),
                _slug(row["primary_party"] or "all"),
            )),
            "state_code": row["state_code"],
            "cycle": row["cycle"],
            "election_date": row["election_date"],
            "stage": row["stage"],
            "chamber": row["chamber"],
            "district": row["district"],
            "primary_party": row["primary_party"],
            "election_system": row["election_system"],
            "coverage_status": row["coverage_status"],
            "missing_scope": "numeric candidate totals and verified primary roster",
            "blocker": row["missing_reason"],
            "next_authoritative_step": (
                "obtain a contemporaneous certified primary candidate roster"
            ),
            "source_urls": (
                MS_PUBLICATION if row["state_code"] == "MS" else OK_PUBLICATION
            ),
        })
    return output


def _observations() -> list[dict]:
    return [
        {
            "state_code": "LA",
            "cycle": "2024",
            "election_date": "2024-11-05",
            "stage": "nonpartisan_primary",
            "chamber": "House",
            "district": "01-06",
            "primary_party": "",
            "election_system": "nonpartisan_open_primary",
            "observation_status": "canonical_stage_required",
            "canonical_requirement": (
                "Do not label these November contests as Democratic or Republican "
                "closed primaries; all candidate parties share one open-primary contest."
            ),
            "source_id": "la-2024-open-congressional-primary",
            "source_url": LA_PUBLICATION,
            "notes": "The December 7 workbook contains no congressional runoff.",
        },
        {
            "state_code": "LA",
            "cycle": "2026",
            "election_date": "2026-11-03",
            "stage": "nonpartisan_primary",
            "chamber": "House",
            "district": "01-06",
            "primary_party": "",
            "election_system": "nonpartisan_open_primary",
            "observation_status": "future_scheduled",
            "canonical_requirement": (
                "No House result rows may be integrated before November 3, 2026."
            ),
            "source_id": "la-2026-house-primary-postponement-authority",
            "source_url": (
                "https://gov.louisiana.gov/assets/2026-Executive-Orders/"
                "JML-Exective-Order-26-038.pdf"
            ),
            "notes": (
                "Executive Order JML 26-038 suspended the May 16 and June 27 "
                "closed U.S. House party primaries after Louisiana v. Callais. "
                "The FEC calendar dated May 18, 2026 confirms House primary "
                "November 3 and runoff December 12. The current official "
                "Candidate Inquiry roster is retained separately without vote "
                "or outcome inference."
            ),
        },
        {
            "state_code": "OK",
            "cycle": "2026",
            "election_date": "2026-08-25",
            "stage": "primary_runoff",
            "chamber": "House",
            "district": "01",
            "primary_party": "Republican",
            "election_system": "partisan_primary",
            "observation_status": "runoff_cancelled_after_withdrawal",
            "canonical_requirement": (
                "Do not infer a runoff winner or final nominee from the June plurality."
            ),
            "source_id": "ok-2026-candidate-withdrawals",
            "source_url": (
                "https://oklahoma.gov/elections/candidates/"
                "2026-candidate-filing-information/2026-candidate-withdrawals.html"
            ),
            "notes": "Second-place candidate Jackson Lahmeyer withdrew June 19.",
        },
    ]


def _source_counts(rows: list[dict], source_id: str) -> tuple[int, int]:
    selected = [row for row in rows if row["source_id"] == source_id]
    return len(selected), len({row["contest_id"] for row in selected})


def _manifest_row(
    raw_root: Path,
    path: Path,
    *,
    source_id: str,
    state_code: str,
    election_date: str,
    stage: str,
    artifact_kind: str,
    authority: str,
    publication_url: str,
    source_url: str,
    source_date: str,
    verification_status: str,
    certification_status: str,
    provisional_status: str,
    candidate_rows: int,
    contests: int,
    coverage: str,
    source_locators: str,
    notes: str,
) -> dict:
    return {
        "source_id": source_id,
        "state_code": state_code,
        "election_date": election_date,
        "stage": stage,
        "artifact_kind": artifact_kind,
        "authority": authority,
        "publication_url": publication_url,
        "source_url": source_url,
        "raw_filename": path.relative_to(raw_root).as_posix(),
        "source_date": source_date,
        "sha256": _sha256(path),
        "acquisition_status": "captured",
        "verification_status": verification_status,
        "certification_status": certification_status,
        "provisional_status": provisional_status,
        "candidate_rows": candidate_rows,
        "contests": contests,
        "coverage": coverage,
        "source_locators": source_locators,
        "notes": notes,
    }


def _source_manifest(
    raw_root: Path,
    results: list[dict],
    rosters: list[dict],
) -> list[dict]:
    manifest = []
    for source in LA_SOURCES:
        path = raw_root / source["filename"]
        candidate_rows, contests = _source_counts(results, source["source_id"])
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source["source_id"],
            state_code="LA",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="official_human_readable_results_workbook",
            authority="Louisiana Secretary of State",
            publication_url=LA_PUBLICATION,
            source_url=source["url"],
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="official",
            provisional_status="final_official_totals",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="all federal contests published for this election",
            source_locators="workbook sheet, row, and candidate column",
            notes=(
                "2024 House rows are a shared nonpartisan open-primary ballot. "
                "2026 Senate rows are closed party-primary contests."
            ),
        ))
    election_dates = raw_root / "la/election-dates.json"
    manifest.append(_manifest_row(
        raw_root,
        election_dates,
        source_id="la-election-dates-and-official-status",
        state_code="LA",
        election_date="",
        stage="inventory",
        artifact_kind="official_election_inventory_json",
        authority="Louisiana Secretary of State",
        publication_url=LA_PUBLICATION,
        source_url=(
            "https://voterportal.sos.la.gov/ElectionResults/"
            "ElectionResults/Data?blob=ElectionDates.htm"
        ),
        source_date="2026-09-22",
        verification_status="verified_official",
        certification_status="ResultsOfficial flags",
        provisional_status="inventory_only",
        candidate_rows=0,
        contests=0,
        coverage="2024 and 2026 election dates and official flags",
        source_locators="$.Dates.Date",
        notes="November 3, 2026 remains future and ResultsOfficial=0.",
    ))
    for path, source_id, source_url, notes in (
        (
            raw_root / "la/2026/elections-calendar-2026.pdf",
            "la-2026-revised-election-calendar",
            "https://www.sos.la.gov/media/p31j4zkc/elections-calendar-2026.pdf",
            "Revised July 2026 calendar places U.S. House open primaries on November 3.",
        ),
        (
            raw_root / "la/2026/051426-fall-house-races.pdf",
            "la-2026-act-7-house-election-notice",
            "https://www.sos.la.gov/media/dcvl5ojl/051426-fall-house-races.pdf",
            "Official notice explains removal of U.S. House from spring party primaries.",
        ),
        (
            raw_root / "la/2024/la-2024-12-07-general-results.xlsx",
            "la-2024-december-no-congressional-runoff-check",
            (
                f"{LA_WORKBOOK_BASE}/20241207/"
                "Election+Results+(12-07-2024).xlsx"
            ),
            "The December workbook contains no congressional contest.",
        ),
    ):
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source_id,
            state_code="LA",
            election_date="",
            stage="calendar_or_runoff_inventory",
            artifact_kind="official_calendar_notice_or_results_workbook",
            authority="Louisiana Secretary of State",
            publication_url=LA_PUBLICATION,
            source_url=source_url,
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="supporting_artifact",
            provisional_status="not_candidate_vote_rows",
            candidate_rows=0,
            contests=0,
            coverage="system transition, future calendar, or runoff absence",
            source_locators="document title and relevant election table",
            notes=notes,
        ))
    executive_order = raw_root / "la/2026/JML-Executive-Order-26-038.pdf"
    manifest.append(_manifest_row(
        raw_root,
        executive_order,
        source_id="la-2026-house-primary-postponement-authority",
        state_code="LA",
        election_date="2026-11-03",
        stage="nonpartisan_primary",
        artifact_kind="governor_executive_order",
        authority="Office of the Governor of Louisiana",
        publication_url=LA_PUBLICATION,
        source_url=(
            "https://gov.louisiana.gov/assets/2026-Executive-Orders/"
            "JML-Exective-Order-26-038.pdf"
        ),
        source_date="2026-04-30",
        verification_status="verified_official",
        certification_status="controlling_calendar_authority",
        provisional_status="future_election",
        candidate_rows=0,
        contests=6,
        coverage="suspension of May and June U.S. House closed party primaries",
        source_locators="pages 1-3; operative Sections 1-3",
        notes=(
            "Suspends House closed party primaries after Louisiana v. Callais "
            "without altering the federal November election date."
        ),
    ))
    executive_order_text = executive_order.with_name(
        "JML-Executive-Order-26-038-ocr.txt"
    )
    manifest.append(_manifest_row(
        raw_root,
        executive_order_text,
        source_id="la-2026-house-primary-postponement-authority-ocr",
        state_code="LA",
        election_date="2026-11-03",
        stage="nonpartisan_primary",
        artifact_kind="derived_pdf_ocr",
        authority="Derived from Louisiana Governor executive order",
        publication_url=LA_PUBLICATION,
        source_url=(
            "https://gov.louisiana.gov/assets/2026-Executive-Orders/"
            "JML-Exective-Order-26-038.pdf"
        ),
        source_date="2026-09-22",
        verification_status="derived_parser_audit",
        certification_status="inherits_executive_order",
        provisional_status="future_election",
        candidate_rows=0,
        contests=6,
        coverage="OCR of all three executive-order pages",
        source_locators="PAGE markers correspond to physical PDF pages",
        notes="Extracted locally with the macOS Vision framework.",
    ))
    fec_calendar = raw_root / "la/2026/2026pdates-fec.pdf"
    manifest.append(_manifest_row(
        raw_root,
        fec_calendar,
        source_id="fec-2026-congressional-primary-calendar-louisiana",
        state_code="LA",
        election_date="2026-11-03",
        stage="nonpartisan_primary",
        artifact_kind="federal_congressional_primary_calendar",
        authority="Federal Election Commission",
        publication_url=LA_PUBLICATION,
        source_url="https://www.fec.gov/documents/5910/2026pdates.pdf",
        source_date="2026-05-18",
        verification_status="verified_official_supporting",
        certification_status="calendar_reference",
        provisional_status="future_election",
        candidate_rows=0,
        contests=6,
        coverage="Louisiana Senate and House primary/runoff dates",
        source_locators="Louisiana row and footnotes 4/2",
        notes=(
            "Lists Senate May 16/June 27 and House November 3/December 12; "
            "cites Executive Order JML 26-038."
        ),
    ))
    future_roster = (
        raw_root / "la/2026/la-2026-11-03-house-candidate-inquiry.html"
    )
    manifest.append(_manifest_row(
        raw_root,
        future_roster,
        source_id="la-2026-house-future-qualifying-roster",
        state_code="LA",
        election_date="2026-11-03",
        stage="nonpartisan_primary",
        artifact_kind="official_candidate_inquiry_roster",
        authority="Louisiana Secretary of State",
        publication_url=LA_PUBLICATION,
        source_url=(
            "https://voterportal.sos.la.gov/"
            "CandidateInquiry?electionDate=20261103"
        ),
        source_date="2026-09-22",
        verification_status="verified_official_future_roster",
        certification_status="qualifying_roster_not_results",
        provisional_status="future_election",
        candidate_rows=39,
        contests=6,
        coverage="qualified U.S. House candidates in districts 1-6",
        source_locators="office IDs 70195-70200 and candidate sections",
        notes=(
            "Retained separately in future_qualifying_roster.csv. "
            "No votes, winner status, or nomination outcome is inferred."
        ),
    ))
    for source in MS_SOURCES:
        path = raw_root / source["filename"]
        candidate_rows, contests = _source_counts(results, source["source_id"])
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source["source_id"],
            state_code="MS",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="official_statewide_recap_pdf",
            authority="Mississippi Secretary of State",
            publication_url=MS_PUBLICATION,
            source_url=MS_PUBLICATION,
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="official_results",
            provisional_status="final_county_recap",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="published statewide federal party-primary totals",
            source_locators="PDF page and extracted-text line",
            notes="X cells mean county not in contest, not zero votes.",
        ))
        text_path = path.with_suffix(".txt")
        manifest.append(_manifest_row(
            raw_root,
            text_path,
            source_id=f"{source['source_id']}-text",
            state_code="MS",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="derived_pdf_text",
            authority="Derived from Mississippi Secretary of State PDF",
            publication_url=MS_PUBLICATION,
            source_url=MS_PUBLICATION,
            source_date="2026-09-22",
            verification_status="derived_parser_input",
            certification_status="inherits_pdf",
            provisional_status="inherits_pdf",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="all PDF pages with physical PAGE markers",
            source_locators="text line locators recorded per result row",
            notes="Extracted with pypdf 6.16.0; result rows hash the PDF.",
        ))
    for source in OK_SOURCES:
        path = raw_root / source["filename"]
        candidate_rows, contests = _source_counts(results, source["source_id"])
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source["source_id"],
            state_code="OK",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="official_public_results_application_json",
            authority="Oklahoma State Election Board",
            publication_url=OK_PUBLICATION,
            source_url=f"{OK_RESULTS}?elecDate={source['date'].replace('-', '')}",
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="isOfficial=true",
            provisional_status="all_precincts_reporting",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="all published statewide federal primary contests",
            source_locators="race and result JSON array indexes",
            notes="Captured from the public application's authenticated network response.",
        ))
    for path, source_id, notes in (
        (
            raw_root / "ok/2024/ok-2024-candidate-roster.json",
            "ok-2024-official-candidate-roster",
            "Used only for missing unopposed/no-candidate primary units.",
        ),
        (
            raw_root / "ok/2026/ok-2026-candidate-list-book.pdf",
            "ok-2026-official-candidate-book",
            "Official ballot-proof candidate book; roster-only rows carry no votes.",
        ),
        (
            raw_root / "ok/2026/ok-2026-candidate-withdrawals.html",
            "ok-2026-candidate-withdrawals",
            "Documents the June 19 House District 1 Republican withdrawal.",
        ),
        (
            raw_root / "ok/2026/ok-2026-roster-only.csv",
            "ok-2026-derived-roster-only-extraction",
            "Structured extraction of three no-result primary units from the candidate book.",
        ),
    ):
        extracted_roster_rows = (
            len(rosters) if path.name == "ok-2026-roster-only.csv" else 0
        )
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source_id,
            state_code="OK",
            election_date="",
            stage="roster_or_observation",
            artifact_kind="official_candidate_roster_or_withdrawal",
            authority="Oklahoma State Election Board",
            publication_url=OK_PUBLICATION,
            source_url=OK_PUBLICATION,
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="not_vote_results",
            provisional_status="roster_only",
            candidate_rows=(
                extracted_roster_rows
                or sum(row["source_id"] == source_id for row in rosters)
            ),
            contests=extracted_roster_rows,
            coverage="missing primary units or nomination-resolution observation",
            source_locators="candidate ID, PDF page, or withdrawal table",
            notes=notes,
        ))
    return manifest


def _attempt_rows() -> list[dict]:
    date = "2026-09-22"
    return [
        {
            "attempt_id": "la-01-official-result-api",
            "attempted_on": date,
            "state_code": "LA",
            "cycle": "2024-2026",
            "source_url": LA_PUBLICATION,
            "method": "discover public data blobs and human-readable workbooks",
            "outcome": "success_complete_past_events",
            "raw_evidence": "la/election-dates.json",
            "notes": "Official flags confirm all captured past events are final.",
        },
        {
            "attempt_id": "la-02-2024-open-primary",
            "attempted_on": date,
            "state_code": "LA",
            "cycle": "2024",
            "source_url": LA_SOURCES[0]["url"],
            "method": "download official November statewide workbook",
            "outcome": "success_complete_open_primary",
            "raw_evidence": LA_SOURCES[0]["filename"],
            "notes": "Six House contests; no closed party-primary labels.",
        },
        {
            "attempt_id": "la-03-2026-senate-primary",
            "attempted_on": date,
            "state_code": "LA",
            "cycle": "2026",
            "source_url": LA_SOURCES[1]["url"],
            "method": "download party-primary and second-primary workbooks",
            "outcome": "success_complete_senate_primary_and_runoff",
            "raw_evidence": LA_SOURCES[1]["filename"],
            "notes": "Both Democratic and Republican Senate contests recovered.",
        },
        {
            "attempt_id": "la-04-revised-house-calendar",
            "attempted_on": date,
            "state_code": "LA",
            "cycle": "2026",
            "source_url": (
                "https://www.sos.la.gov/media/p31j4zkc/elections-calendar-2026.pdf"
            ),
            "method": "verify Act 7 revised calendar",
            "outcome": "future_house_open_primary",
            "raw_evidence": "la/2026/elections-calendar-2026.pdf",
            "notes": "House contests moved to November 3, after the current cutoff.",
        },
        {
            "attempt_id": "ms-01-2024-statewide-recaps",
            "attempted_on": date,
            "state_code": "MS",
            "cycle": "2024",
            "source_url": MS_PUBLICATION,
            "method": "download Democratic, Republican, and runoff statewide recaps",
            "outcome": "success_partial_primary_complete_runoff",
            "raw_evidence": MS_SOURCES[0]["filename"],
            "notes": "House 3 Democratic has no published result row.",
        },
        {
            "attempt_id": "ms-02-2026-statewide-recaps",
            "attempted_on": date,
            "state_code": "MS",
            "cycle": "2026",
            "source_url": MS_PUBLICATION,
            "method": "download both statewide party recaps",
            "outcome": "success_complete_primary",
            "raw_evidence": MS_SOURCES[3]["filename"],
            "notes": "Senate and all four House districts recovered for both parties.",
        },
        {
            "attempt_id": "ok-01-official-page",
            "attempted_on": date,
            "state_code": "OK",
            "cycle": "2024-2026",
            "source_url": OK_PUBLICATION,
            "method": "follow official result-application links",
            "outcome": "success_election_dates",
            "raw_evidence": "discovery/ok-2026-runoff-page.html",
            "notes": "June and August election events identified separately.",
        },
        {
            "attempt_id": "ok-02-public-app-network",
            "attempted_on": date,
            "state_code": "OK",
            "cycle": "2024-2026",
            "source_url": OK_RESULTS,
            "method": "capture public application result JSON responses",
            "outcome": "success_official_primary_and_runoff",
            "raw_evidence": OK_SOURCES[2]["filename"],
            "notes": "All captured result payloads have isOfficial=true.",
        },
        {
            "attempt_id": "ok-03-candidate-rosters",
            "attempted_on": date,
            "state_code": "OK",
            "cycle": "2024-2026",
            "source_url": (
                "https://oklahoma.gov/elections/candidates/"
                "2026-candidate-filing-information.html"
            ),
            "method": "capture official filing grid and candidate book",
            "outcome": "success_roster_only_units",
            "raw_evidence": "ok/2026/ok-2026-roster-only.csv",
            "notes": "Includes 2026 Libertarian Senate and unopposed House units.",
        },
        {
            "attempt_id": "ok-04-withdrawal-resolution",
            "attempted_on": date,
            "state_code": "OK",
            "cycle": "2026",
            "source_url": (
                "https://oklahoma.gov/elections/candidates/"
                "2026-candidate-filing-information/2026-candidate-withdrawals.html"
            ),
            "method": "verify missing House 1 Republican runoff",
            "outcome": "runoff_cancelled_after_withdrawal",
            "raw_evidence": "ok/2026/ok-2026-candidate-withdrawals.html",
            "notes": "Prevents the June plurality from being treated as a runoff result.",
        },
    ]


def _validate_output(rows: list[dict]) -> None:
    if not rows:
        raise ValueError("No Gulf-state primary results generated")
    seen = set()
    for row in rows:
        if set(row) != set(FIELDS):
            raise ValueError(f"Output schema differs: {row.get('candidate_name')}")
        key = (row["contest_id"], row["candidate_source_id"])
        if key in seen:
            raise ValueError(f"Duplicate candidate result: {key}")
        seen.add(key)
        if not isinstance(row["votes"], int) or row["votes"] < 0:
            raise ValueError(f"Invalid vote total: {row}")
        if row["write_in"] and row["party"]:
            raise ValueError(f"Write-in row asserts party affiliation: {row}")
        if len(row["source_sha256"]) != 64 or not row["source_locators"]:
            raise ValueError(f"Incomplete provenance: {row}")


def generate(
    raw_root: Path | None = None,
    analysis_root: Path | None = None,
) -> list[dict]:
    raw_root = raw_root or RAW_DIR / RAW_SUBDIR
    analysis_root = analysis_root or ANALYSIS_DATA_DIR / "congressional"
    results = []
    for source in LA_SOURCES:
        results.extend(parse_louisiana(raw_root / source["filename"], source))
    for source in MS_SOURCES:
        results.extend(parse_mississippi(raw_root / source["filename"], source))
    for source in OK_SOURCES:
        results.extend(parse_oklahoma(raw_root / source["filename"], source))
    results.sort(key=lambda row: (
        row["state_code"],
        row["election_date"],
        row["stage"],
        row["chamber"],
        row["district"],
        row["primary_party"],
        row["candidate_name"],
    ))
    _validate_output(results)
    ok_actual_2024 = {
        (row["chamber"], row["district"], row["primary_party"])
        for row in results
        if row["state_code"] == "OK"
        and row["election_date"] == "2024-06-18"
    }
    ok_2024_roster, ok_2024_missing = parse_oklahoma_2024_roster(
        raw_root / "ok/2024/ok-2024-candidate-roster.json",
        ok_actual_2024,
    )
    ok_2026_roster = parse_oklahoma_2026_roster(
        raw_root / "ok/2026/ok-2026-roster-only.csv",
        raw_root / "ok/2026/ok-2026-candidate-list-book.pdf",
    )
    rosters = sorted(
        ok_2024_roster + ok_2026_roster,
        key=lambda row: (
            row["state_code"],
            row["cycle"],
            row["district"],
            row["primary_party"],
        ),
    )
    coverage = _coverage_rows(results, rosters, ok_2024_missing)
    _assert_future_scheduled_coverage(coverage, results)
    write_csv(analysis_root / RESULT_FILENAME, results, FIELDS)
    manifest_dir = analysis_root / MANIFEST_SUBDIR
    write_csv(
        manifest_dir / "source_manifest.csv",
        _source_manifest(raw_root, results, rosters),
        SOURCE_MANIFEST_FIELDS,
    )
    write_csv(manifest_dir / "attempt_ledger.csv", _attempt_rows(), ATTEMPT_FIELDS)
    write_csv(manifest_dir / "unresolved_gaps.csv", _gap_rows(coverage), GAP_FIELDS)
    write_csv(
        manifest_dir / "contest_manifest.csv",
        _contest_rows(results),
        CONTEST_FIELDS,
    )
    write_csv(manifest_dir / "roster_only.csv", rosters, ROSTER_FIELDS)
    write_csv(manifest_dir / "coverage_units.csv", coverage, COVERAGE_FIELDS)
    write_csv(
        manifest_dir / "source_observations.csv",
        _observations(),
        OBSERVATION_FIELDS,
    )
    _update_party_accounting_future_status(analysis_root, results)
    return results


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate official Gulf-state congressional primary results"
    )
    parser.parse_args()
    rows = generate()
    print(
        f"Wrote {len(rows)} candidate results to "
        f"{ANALYSIS_DATA_DIR / 'congressional' / RESULT_FILENAME}"
    )


if __name__ == "__main__":
    main()
