"""Recover official South Carolina and Tennessee congressional primaries."""

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .io import write_csv
from .paths import ANALYSIS_DATA_DIR, RAW_DIR
from .state_results import FIELDS

RAW_SUBDIR = "sc_tn_primaries"
RESULT_FILENAME = "sc_tn_primary_results.csv"
MANIFEST_SUBDIR = "sc_tn"

PARTIES = ("Democratic", "Republican")
TN_SOURCE_URLS = {
    "2024": (
        "https://sos-prod.tnsosgovfiles.com/s3fs-public/document/"
        "20240801AllbyPrecinct.xlsx"
    ),
    "2026": (
        "https://sos-prod.tnsosgovfiles.com/s3fs-public/document/"
        "20260806AllbyPrecinct.xlsx"
    ),
}

SOURCE_MANIFEST_FIELDS = [
    "source_id",
    "state_code",
    "election_date",
    "stage",
    "artifact_kind",
    "authority",
    "publication_url",
    "source_url",
    "raw_filename",
    "source_date",
    "sha256",
    "acquisition_status",
    "verification_status",
    "certification_status",
    "provisional_status",
    "candidate_rows",
    "contests",
    "coverage",
    "source_locators",
    "notes",
]

ATTEMPT_FIELDS = [
    "attempt_id",
    "attempted_on",
    "state_code",
    "cycle",
    "source_url",
    "method",
    "outcome",
    "raw_evidence",
    "notes",
]

GAP_FIELDS = [
    "gap_id",
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "primary_party",
    "coverage_status",
    "missing_scope",
    "blocker",
    "next_authoritative_step",
    "source_urls",
]

CONTEST_FIELDS = [
    "contest_id",
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "term",
    "primary_party",
    "candidate_count",
    "unopposed",
    "actual_zero_rows",
    "has_write_in",
    "result_status",
    "certification_status",
    "source_id",
]

COVERAGE_FIELDS = [
    "state_code",
    "cycle",
    "election_date",
    "stage",
    "chamber",
    "district",
    "term",
    "primary_party",
    "coverage_status",
    "result_rows",
    "missing_reason",
    "source_ids",
]

SC_PUBLICATION = "https://scvotes.gov/elections-statistics/election-results/"
SC_DATABASE = "https://electionhistory.scvotes.gov"
SC_CIVERA = "https://sc.elstats.civera.com/api"
SC_ENR = "https://www.enr-scvotes.org"
TN_PUBLICATION = "https://sos.tn.gov/elections/results"

SC_2024_CONTEST_IDS = tuple(range(6898, 6907)) + (7067,)

SC_2026_SOURCES = (
    {
        "source_id": "sc-2026-statewide-primary-enr",
        "election_id": "126294",
        "version": "375593",
        "date": "2026-06-09",
        "stage": "primary",
        "term": "regular",
        "summary": "sc/2026/sc-126294-en-summary.json",
        "config": "sc/2026/sc-126294-config.json",
        "certification": "sc/2026/sc-2026-primary-signed-certification.pdf",
    },
    {
        "source_id": "sc-2026-statewide-primary-runoff-enr",
        "election_id": "126718",
        "version": "376464",
        "date": "2026-06-23",
        "stage": "primary_runoff",
        "term": "regular",
        "summary": "sc/2026/sc-126718-en-summary.json",
        "config": "sc/2026/sc-126718-config.json",
        "certification": "sc/2026/sc-2026-runoff-signed-certification.pdf",
    },
    {
        "source_id": "sc-2026-special-senate-primary-enr",
        "election_id": "127017",
        "version": "378537",
        "date": "2026-08-11",
        "stage": "special_primary",
        "term": "special",
        "summary": "sc/2026/sc-127017-en-summary.json",
        "config": "sc/2026/sc-127017-config.json",
        "certification": (
            "sc/2026/"
            "sc-2026-special-senate-primary-signed-certification.pdf"
        ),
    },
    {
        "source_id": "sc-2026-special-senate-primary-runoff-enr",
        "election_id": "127179",
        "version": "379277",
        "date": "2026-08-25",
        "stage": "special_primary_runoff",
        "term": "special",
        "summary": "sc/2026/sc-127179-en-summary.json",
        "config": "sc/2026/sc-127179-config.json",
        "certification": (
            "sc/2026/"
            "sc-2026-special-senate-runoff-signed-certification.pdf"
        ),
    },
)

SC_PRIMARY_EXPECTED = {
    ("2024-06-11", "primary"): {
        ("House", f"{district:02}", party)
        for district in range(1, 8)
        for party in PARTIES
    },
    ("2026-06-09", "primary"): {
        *{
            ("House", f"{district:02}", party)
            for district in range(1, 8)
            for party in PARTIES
        },
        ("Senate", "S", "Democratic"),
        ("Senate", "S", "Republican"),
    },
}

SC_RUNOFF_EXPECTED = {
    ("2024-06-25", "primary_runoff"): {
        ("House", "03", "Republican"),
    },
    ("2026-06-23", "primary_runoff"): {
        ("House", "01", "Democratic"),
        ("House", "01", "Republican"),
        ("House", "02", "Democratic"),
    },
    ("2026-08-11", "special_primary"): {
        ("Senate", "S", "Democratic"),
        ("Senate", "S", "Republican"),
    },
    ("2026-08-25", "special_primary_runoff"): {
        ("Senate", "S", "Republican"),
    },
}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _integer(value: object) -> int:
    text = str(value).strip().replace(",", "")
    if re.fullmatch(r"\d+", text) is None:
        raise ValueError(f"Invalid vote total: {value!r}")
    return int(text)


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")


def _contest_id(
    state: str,
    date: str,
    chamber: str,
    district: str,
    term: str,
    party: str,
) -> str:
    return "-".join((
        state.casefold(),
        date,
        chamber.casefold(),
        district.casefold(),
        term,
        _slug(party),
    ))


def _locator_ranges(filename: str, numbers: list[int]) -> str:
    ranges = []
    start = end = None
    for number in sorted(set(numbers)):
        if start is None:
            start = end = number
        elif number == end + 1:
            end = number
        else:
            ranges.append(str(start) if start == end else f"{start}-{end}")
            start = end = number
    if start is not None:
        ranges.append(str(start) if start == end else f"{start}-{end}")
    return f"{filename}:{','.join(ranges)}"


def parse_south_carolina_civera(path: Path) -> list[dict]:
    source_hash = _sha256(path)
    contest_number = path.stem.rsplit("-", 1)[-1]
    source_id = f"sc-2024-contest-{contest_number}"
    with path.open(newline="", encoding="utf-8-sig") as handle:
        table = list(csv.reader(handle))
    if len(table) < 4 or len(table[0]) != len(table[1]):
        raise ValueError("South Carolina election-history table schema changed")
    names = table[0][2:]
    party_values = table[1][2:]
    district_total = table[2]
    if district_total[0] != "Congressional District":
        raise ValueError("South Carolina congressional total row is missing")
    district = district_total[1].zfill(2)
    candidate_columns = [
        index for index, name in enumerate(names, 2)
        if name not in {
            "Total Votes Cast",
            "Overvotes/Undervotes",
            "Total Ballots Cast",
        }
    ]
    if not candidate_columns:
        raise ValueError("South Carolina contest has no candidate columns")
    parties = {
        party_values[index - 2]
        for index in candidate_columns
        if party_values[index - 2]
    }
    if len(parties) != 1:
        raise ValueError(f"South Carolina primary party differs: {parties}")
    party = next(iter(parties))
    if party not in PARTIES:
        raise ValueError(f"Unknown South Carolina primary party: {party}")
    date = "2024-06-25" if contest_number == "7067" else "2024-06-11"
    stage = "primary_runoff" if contest_number == "7067" else "primary"
    precinct_rows = [
        (line_number, row)
        for line_number, row in enumerate(table[3:], 4)
        if row and row[0] == "Precinct"
    ]
    if not precinct_rows:
        raise ValueError("South Carolina contest has no precinct rows")
    output = []
    for candidate_number, column in enumerate(candidate_columns, 1):
        name = names[column - 2]
        write_in = name.casefold() in {"write-in", "scattering"}
        precinct_votes = [_integer(row[column]) for _, row in precinct_rows]
        official_total = _integer(district_total[column])
        if sum(precinct_votes) != official_total:
            raise ValueError(
                f"South Carolina precincts differ from district total: {name}"
            )
        output.append({
            "contest_id": _contest_id(
                "SC", date, "House", district, "regular", party
            ),
            "cycle": "2024",
            "election_date": date,
            "stage": stage,
            "state_code": "SC",
            "chamber": "House",
            "district": district,
            "term": "regular",
            "election_system": "partisan_primary",
            "primary_party": party,
            "candidate_name": name,
            "candidate_source_id": f"{contest_number}:{candidate_number}",
            "party": "" if write_in else party,
            "write_in": write_in,
            "votes": official_total,
            "source_id": source_id,
            "source_url": (
                f"{SC_CIVERA}/download_contest/{contest_number}_table.csv?"
                "split_party=false"
            ),
            "source_sha256": source_hash,
            "source_locators": (
                f"{path.name}:candidate-column-{column + 1}; "
                + _locator_ranges(
                    path.name, [line for line, _ in precinct_rows]
                )
            ),
            "reporting_units": len(precinct_rows),
        })
    if not output:
        raise ValueError(f"No South Carolina candidates parsed from {path}")
    return output


def _sc_enr_contest(name: str) -> tuple[str, str, str] | None:
    normalized = " ".join(name.split())
    senate = re.fullmatch(r"U\.S\. Senate - (DEM|REP)", normalized)
    if senate:
        return "Senate", "S", {
            "DEM": "Democratic",
            "REP": "Republican",
        }[senate[1]]
    house = re.fullmatch(
        r"U\.S\. House of Representatives, District (\d+) - (DEM|REP)",
        normalized,
    )
    if house:
        return "House", house[1].zfill(2), {
            "DEM": "Democratic",
            "REP": "Republican",
        }[house[2]]
    return None


def parse_south_carolina_enr(
    summary_path: Path,
    config_path: Path,
    source: dict,
) -> list[dict]:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    if config.get("isPublished") is not True:
        raise ValueError(f'{source["source_id"]}: ENR source is not published')
    contests = json.loads(summary_path.read_text(encoding="utf-8"))
    if not isinstance(contests, list):
        raise ValueError("South Carolina ENR summary schema changed")
    source_hash = _sha256(summary_path)
    output = []
    for contest_index, contest in enumerate(contests):
        parsed = _sc_enr_contest(contest.get("C", ""))
        if parsed is None:
            continue
        chamber, district, party = parsed
        names = contest.get("CH", [])
        votes = contest.get("V", [])
        parties = contest.get("P", [])
        if not names or not (len(names) == len(votes) == len(parties)):
            raise ValueError("South Carolina ENR candidate arrays differ")
        for candidate_index, (name, vote, party_code) in enumerate(
            zip(names, votes, parties, strict=True),
            1,
        ):
            expected_code = "DEM" if party == "Democratic" else "REP"
            write_in = "WRITE-IN" in name.upper() or name.upper() == "WRITE-IN"
            if not write_in and party_code != expected_code:
                raise ValueError(f"South Carolina ENR party differs: {name}")
            output.append({
                "contest_id": _contest_id(
                    "SC",
                    source["date"],
                    chamber,
                    district,
                    source["term"],
                    party,
                ),
                "cycle": "2026",
                "election_date": source["date"],
                "stage": source["stage"],
                "state_code": "SC",
                "chamber": chamber,
                "district": district,
                "term": source["term"],
                "election_system": "partisan_primary",
                "primary_party": party,
                "candidate_name": " ".join(name.split()),
                "candidate_source_id": (
                    f'{contest["K"]}:{candidate_index}'
                ),
                "party": "" if write_in else party,
                "write_in": write_in,
                "votes": _integer(vote),
                "source_id": source["source_id"],
                "source_url": (
                    f"{SC_ENR}/SC/{source['election_id']}/"
                    f"{source['version']}/json/en/summary.json"
                ),
                "source_sha256": source_hash,
                "source_locators": (
                    f"{summary_path.name}:$[{contest_index}]"
                    f".CH[{candidate_index - 1}]/V[{candidate_index - 1}]"
                ),
                "reporting_units": _integer(contest.get("TP", 1)) or 1,
            })
    if not output:
        raise ValueError(f'{source["source_id"]}: no federal candidates parsed')
    return output


def parse_tennessee(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["SOFFICEL"]
        values = sheet.iter_rows(values_only=True)
        header = next(values)
        required = {
            "COUNTY",
            "PRCTSEQ",
            "PRECINCT",
            "OFFICENAME",
            "ELECTDATE",
            "ELECTTYPE",
        }
        if not required <= set(header):
            raise ValueError("Tennessee workbook schema changed")
        indexes = {name: index for index, name in enumerate(header)}
        groups = defaultdict(list)
        seen = set()
        contest_identities = {}
        for line_number, values in enumerate(values, 2):
            office = str(values[indexes["OFFICENAME"]] or "")
            if office == "United States Senate":
                chamber, district = "Senate", "S"
            else:
                match = re.fullmatch(
                    r"United States House of Representatives District (\d+)",
                    office,
                )
                if match is None:
                    continue
                chamber, district = "House", match[1].zfill(2)
            election_type = str(values[indexes["ELECTTYPE"]] or "")
            if election_type not in {
                "Democratic Primary",
                "Republican Primary",
            }:
                continue
            party = election_type.removesuffix(" Primary")
            election_date = str(values[indexes["ELECTDATE"]])
            if election_date != source["date"]:
                raise ValueError("Tennessee workbook contains a different date")
            unit = (
                f'{values[indexes["COUNTY"]]}:'
                f'{values[indexes["PRCTSEQ"]]}:'
                f'{values[indexes["PRECINCT"]]}'
            )
            contest_key = (chamber, district, party)
            for candidate_number in range(1, 11):
                name = values[indexes[f"RNAME{candidate_number}"]]
                if name is None:
                    continue
                vote = values[indexes[f"PVTALLY{candidate_number}"]]
                candidate_party = values[indexes[f"PARTY{candidate_number}"]]
                name = " ".join(str(name).split())
                write_in = name.upper().startswith("WRITE-IN")
                if not write_in and str(candidate_party) != party:
                    raise ValueError(f"Tennessee candidate party differs: {name}")
                duplicate = (*contest_key, candidate_number, unit)
                if duplicate in seen:
                    raise ValueError(f"Duplicate Tennessee precinct row: {duplicate}")
                seen.add(duplicate)
                identity = (name, write_in)
                identity_key = (*contest_key, candidate_number)
                if (
                    identity_key in contest_identities
                    and contest_identities[identity_key] != identity
                ):
                    raise ValueError(f"Tennessee candidate column changed: {identity_key}")
                contest_identities[identity_key] = identity
                groups[identity_key].append({
                    "name": name,
                    "write_in": write_in,
                    "votes": _integer(vote),
                    "line": line_number,
                    "unit": unit,
                })
    finally:
        workbook.close()
    source_hash = _sha256(path)
    output = []
    for (chamber, district, party, candidate_number), records in sorted(groups.items()):
        name, write_in = records[0]["name"], records[0]["write_in"]
        output.append({
            "contest_id": _contest_id(
                "TN", source["date"], chamber, district, "regular", party
            ),
            "cycle": source["date"][:4],
            "election_date": source["date"],
            "stage": "primary",
            "state_code": "TN",
            "chamber": chamber,
            "district": district,
            "term": "regular",
            "election_system": "partisan_primary",
            "primary_party": party,
            "candidate_name": name,
            "candidate_source_id": (
                f"tn-{source['date'][:4]}-{chamber.casefold()}-"
                f"{district.casefold()}-{party.casefold()}-{candidate_number}"
            ),
            "party": "" if write_in else party,
            "write_in": write_in,
            "votes": sum(record["votes"] for record in records),
            "source_id": source["source_id"],
            "source_url": source["url"],
            "source_sha256": source_hash,
            "source_locators": _locator_ranges(
                f"{path.name}:{sheet.title}",
                [record["line"] for record in records],
            ),
            "reporting_units": len({record["unit"] for record in records}),
        })
    actual_units = {
        (row["chamber"], row["district"], row["primary_party"])
        for row in output
    }
    expected_units = source.get("expected_units") or {
        *{("Senate", "S", party) for party in PARTIES},
        *{
            ("House", f"{district:02}", party)
            for district in range(1, 10)
            for party in PARTIES
        },
    }
    if actual_units != expected_units:
        raise ValueError(
            f"Tennessee district-party coverage differs: "
            f"{actual_units ^ expected_units}"
        )
    return output


def _result_unit(row: dict) -> tuple[str, str, str, str, str, str]:
    return (
        row["state_code"],
        row["election_date"],
        row["stage"],
        row["chamber"],
        row["district"],
        row["primary_party"],
    )


def _coverage_rows(rows: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in rows:
        groups[_result_unit(row)].append(row)
    expected = []
    expected.extend(
        ("SC", "2024-06-11", "primary", "House", f"{district:02}", "regular", party)
        for district in range(1, 8)
        for party in PARTIES
    )
    expected.extend(
        ("SC", "2024-06-25", "primary_runoff", chamber, district, "regular", party)
        for chamber, district, party in SC_RUNOFF_EXPECTED[
            ("2024-06-25", "primary_runoff")
        ]
    )
    expected.extend(
        (
            "SC", "2026-06-09", "primary",
            chamber, district, "regular", party,
        )
        for chamber, district, party in SC_PRIMARY_EXPECTED[
            ("2026-06-09", "primary")
        ]
    )
    expected.extend(
        (
            "SC", "2026-06-23", "primary_runoff",
            chamber, district, "regular", party,
        )
        for chamber, district, party in SC_RUNOFF_EXPECTED[
            ("2026-06-23", "primary_runoff")
        ]
    )
    expected.extend(
        (
            "SC", "2026-08-11", "special_primary",
            "Senate", "S", "special", party,
        )
        for party in PARTIES
    )
    expected.append((
        "SC", "2026-08-25", "special_primary_runoff",
        "Senate", "S", "special", "Republican",
    ))
    for year, date in (("2024", "2024-08-01"), ("2026", "2026-08-06")):
        expected.extend(
            ("TN", date, "primary", "Senate", "S", "regular", party)
            for party in PARTIES
        )
        expected.extend(
            (
                "TN", date, "primary", "House",
                f"{district:02}", "regular", party,
            )
            for district in range(1, 10)
            for party in PARTIES
        )
    output = []
    for state, date, stage, chamber, district, term, party in expected:
        key = (state, date, stage, chamber, district, party)
        result_rows = groups.get(key, [])
        status = "actual_results" if result_rows else "no_published_primary_result"
        output.append({
            "state_code": state,
            "cycle": date[:4],
            "election_date": date,
            "stage": stage,
            "chamber": chamber,
            "district": district,
            "term": term,
            "primary_party": party,
            "coverage_status": status,
            "result_rows": len(result_rows),
            "missing_reason": (
                "" if result_rows else (
                    "official result sources publish no contest; candidate status "
                    "was not inferred from the general-election ballot"
                )
            ),
            "source_ids": " | ".join(sorted({
                row["source_id"] for row in result_rows
            })),
        })
    return output


def _contest_rows(rows: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in rows:
        groups[row["contest_id"]].append(row)
    output = []
    for contest_id, candidates in sorted(groups.items()):
        first = candidates[0]
        output.append({
            "contest_id": contest_id,
            "state_code": first["state_code"],
            "cycle": first["cycle"],
            "election_date": first["election_date"],
            "stage": first["stage"],
            "chamber": first["chamber"],
            "district": first["district"],
            "term": first["term"],
            "primary_party": first["primary_party"],
            "candidate_count": len(candidates),
            "unopposed": len(candidates) == 1,
            "actual_zero_rows": sum(candidate["votes"] == 0 for candidate in candidates),
            "has_write_in": any(candidate["write_in"] for candidate in candidates),
            "result_status": "actual_numeric",
            "certification_status": "official",
            "source_id": first["source_id"],
        })
    return output


def _gap_rows(coverage: list[dict]) -> list[dict]:
    output = []
    for row in coverage:
        if row["coverage_status"] == "actual_results":
            continue
        output.append({
            "gap_id": "-".join((
                row["state_code"].casefold(),
                row["election_date"],
                row["stage"],
                row["chamber"].casefold(),
                row["district"].casefold(),
                _slug(row["primary_party"]),
            )),
            "state_code": row["state_code"],
            "cycle": row["cycle"],
            "election_date": row["election_date"],
            "stage": row["stage"],
            "chamber": row["chamber"],
            "district": row["district"],
            "primary_party": row["primary_party"],
            "coverage_status": row["coverage_status"],
            "missing_scope": "numeric candidate totals and primary roster status",
            "blocker": row["missing_reason"],
            "next_authoritative_step": (
                "obtain a contemporaneous certified primary candidate roster"
            ),
            "source_urls": SC_PUBLICATION,
        })
    return output


def _source_counts(rows: list[dict], source_id: str) -> tuple[int, int]:
    selected = [row for row in rows if row["source_id"] == source_id]
    return len(selected), len({row["contest_id"] for row in selected})


def _manifest_row(
    raw_root: Path,
    path: Path,
    *,
    source_id: str,
    state_code: str,
    election_date: str,
    stage: str,
    artifact_kind: str,
    authority: str,
    publication_url: str,
    source_url: str,
    source_date: str,
    verification_status: str,
    certification_status: str,
    provisional_status: str,
    candidate_rows: int,
    contests: int,
    coverage: str,
    source_locators: str,
    notes: str,
) -> dict:
    return {
        "source_id": source_id,
        "state_code": state_code,
        "election_date": election_date,
        "stage": stage,
        "artifact_kind": artifact_kind,
        "authority": authority,
        "publication_url": publication_url,
        "source_url": source_url,
        "raw_filename": path.relative_to(raw_root).as_posix(),
        "source_date": source_date,
        "sha256": _sha256(path),
        "acquisition_status": "captured",
        "verification_status": verification_status,
        "certification_status": certification_status,
        "provisional_status": provisional_status,
        "candidate_rows": candidate_rows,
        "contests": contests,
        "coverage": coverage,
        "source_locators": source_locators,
        "notes": notes,
    }


def _source_manifest(raw_root: Path, rows: list[dict]) -> list[dict]:
    manifest = []
    for contest_id in SC_2024_CONTEST_IDS:
        source_id = f"sc-2024-contest-{contest_id}"
        candidate_rows, contests = _source_counts(rows, source_id)
        path = raw_root / f"sc/2024/sc-contest-{contest_id}.csv"
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source_id,
            state_code="SC",
            election_date=(
                "2024-06-25" if contest_id == 7067 else "2024-06-11"
            ),
            stage="primary_runoff" if contest_id == 7067 else "primary",
            artifact_kind="official_election_database_contest_csv",
            authority="South Carolina State Election Commission",
            publication_url=SC_PUBLICATION,
            source_url=(
                f"{SC_CIVERA}/download_contest/{contest_id}_table.csv?"
                "split_party=false"
            ),
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="official",
            provisional_status="final_official_totals",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="one contested federal primary or runoff",
            source_locators="CSV source line ranges recorded per result row",
            notes="Accounting rows are excluded; candidate zeros are retained.",
        ))
    for source in SC_2026_SOURCES:
        summary_path = raw_root / source["summary"]
        config_path = raw_root / source["config"]
        candidate_rows, contests = _source_counts(rows, source["source_id"])
        summary_url = (
            f"{SC_ENR}/SC/{source['election_id']}/{source['version']}/"
            "json/en/summary.json"
        )
        manifest.append(_manifest_row(
            raw_root,
            summary_path,
            source_id=source["source_id"],
            state_code="SC",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="official_enr_summary_json",
            authority="South Carolina State Election Commission linked ENR",
            publication_url=SC_PUBLICATION,
            source_url=summary_url,
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="certified",
            provisional_status="final_official_totals",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="all federal contests published for this election event",
            source_locators="JSON array indexes recorded per result row",
            notes="Captured through the public browser application after its WAF challenge.",
        ))
        manifest.append(_manifest_row(
            raw_root,
            config_path,
            source_id=f"{source['source_id']}-config",
            state_code="SC",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="official_enr_publication_metadata",
            authority="South Carolina State Election Commission linked ENR",
            publication_url=SC_PUBLICATION,
            source_url=(
                f"{SC_ENR}/SC/{source['election_id']}/{source['version']}/"
                "json/config.json"
            ),
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="isPublished=true",
            provisional_status="metadata_only",
            candidate_rows=0,
            contests=0,
            coverage="publication status and language metadata",
            source_locators="$.isPublished",
            notes="The parser requires isPublished=true.",
        ))
        certification_path = raw_root / source["certification"]
        manifest.append(_manifest_row(
            raw_root,
            certification_path,
            source_id=f"{source['source_id']}-certification",
            state_code="SC",
            election_date=source["date"],
            stage=source["stage"],
            artifact_kind="signed_state_canvass_pdf",
            authority="South Carolina Board of State Canvassers",
            publication_url=(
                "https://scvotes.gov/about-the-sec/public-notices/"
            ),
            source_url=SC_PUBLICATION,
            source_date=(
                "2026-06-12" if source["date"] == "2026-06-09"
                else "2026-06-26" if source["date"] == "2026-06-23"
                else "2026-08-14" if source["date"] == "2026-08-11"
                else "2026-08-28"
            ),
            verification_status="verified_official",
            certification_status="signed_certification",
            provisional_status="certification_only",
            candidate_rows=0,
            contests=contests,
            coverage="certification date and election identity",
            source_locators="signed canvass pages",
            notes="Vote totals come from the linked ENR JSON, not OCR.",
        ))
    audit_path = raw_root / "sc/2026/sc-2026-primary-contest-results-comparison.pdf"
    manifest.append(_manifest_row(
        raw_root,
        audit_path,
        source_id="sc-2026-primary-results-verification-audit",
        state_code="SC",
        election_date="2026-06-09",
        stage="primary",
        artifact_kind="official_results_verification_audit_pdf",
        authority="South Carolina State Election Commission",
        publication_url="https://scvotes.gov/elections-statistics/election-audits/",
        source_url=(
            "https://scvotes.gov/wp-content/uploads/2026/06/"
            "2.Contest-Results-Comparison.pdf"
        ),
        source_date="2026-06-26",
        verification_status="verified_official",
        certification_status="post_election_audit",
        provisional_status="supporting_artifact",
        candidate_rows=0,
        contests=11,
        coverage="all contested June federal primary offices",
        source_locators="contest comparison tables",
        notes="Independent supporting check; ENR totals remain the result source.",
    ))
    for year, filename, date in (
        ("2024", "20240801AllbyPrecinct.xlsx", "2024-08-01"),
        ("2026", "20260806AllbyPrecinct.xlsx", "2026-08-06"),
    ):
        source_id = f"tn-{year}-statewide-primary-precinct-workbook"
        candidate_rows, contests = _source_counts(rows, source_id)
        path = raw_root / "tn" / year / filename
        manifest.append(_manifest_row(
            raw_root,
            path,
            source_id=source_id,
            state_code="TN",
            election_date=date,
            stage="primary",
            artifact_kind="official_statewide_precinct_workbook",
            authority="Tennessee Secretary of State, Division of Elections",
            publication_url=TN_PUBLICATION,
            source_url=TN_SOURCE_URLS[year],
            source_date="2026-09-22",
            verification_status="verified_official",
            certification_status="official",
            provisional_status="final_precinct_totals",
            candidate_rows=candidate_rows,
            contests=contests,
            coverage="all Senate and House district-party primary units",
            source_locators="workbook sheet and row ranges recorded per result",
            notes=(
                "Named write-ins are retained with party blank. "
                "The 2026 workbook uses the revised May 2026 congressional map."
            ),
        ))
    redistricting = raw_root / "tn/2026/tn-2026-congressional-redistricting.html"
    manifest.append(_manifest_row(
        raw_root,
        redistricting,
        source_id="tn-2026-congressional-redistricting",
        state_code="TN",
        election_date="2026-08-06",
        stage="primary",
        artifact_kind="official_redistricting_notice",
        authority="Tennessee Secretary of State",
        publication_url=TN_PUBLICATION,
        source_url=(
            "https://sos.tn.gov/announcements/2026-congressional-redistricting"
        ),
        source_date="2026-05-01",
        verification_status="verified_official",
        certification_status="map_inventory",
        provisional_status="not_vote_results",
        candidate_rows=0,
        contests=9,
        coverage="revised congressional districts adopted in May 2026",
        source_locators="page title and redistricting notice body",
        notes="Prevents superseded district geography from being used for 2026.",
    ))
    return manifest


def _attempt_rows() -> list[dict]:
    date = "2026-09-22"
    return [
        {
            "attempt_id": "sc-01-results-inventory",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2024-2026",
            "source_url": SC_PUBLICATION,
            "method": "inspect official election event inventory",
            "outcome": "success_dates_and_event_ids",
            "raw_evidence": "",
            "notes": "Identified regular, runoff, and special Senate events.",
        },
        {
            "attempt_id": "sc-02-2024-election-database",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2024",
            "source_url": SC_DATABASE,
            "method": "download official contest-level result CSVs",
            "outcome": "success_contested_primary_and_runoff",
            "raw_evidence": "sc/2024/sc-contest-6898.csv",
            "notes": "Nine primary contests and one congressional runoff.",
        },
        {
            "attempt_id": "sc-03-enr-direct",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2026",
            "source_url": f"{SC_ENR}/SC/126294/",
            "method": "request linked ENR directly",
            "outcome": "blocked_aws_waf_challenge",
            "raw_evidence": "",
            "notes": "No repeated loop; a public browser session was used.",
        },
        {
            "attempt_id": "sc-04-enr-browser-json",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2026",
            "source_url": f"{SC_ENR}/SC/126294/",
            "method": "capture versioned JSON through public browser application",
            "outcome": "success_regular_and_runoff",
            "raw_evidence": "sc/2026/sc-126294-en-summary.json",
            "notes": "Includes June primary and runoff candidate totals.",
        },
        {
            "attempt_id": "sc-05-special-senate-json",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2026",
            "source_url": f"{SC_ENR}/SC/127017/",
            "method": "capture special Senate primary and runoff JSON",
            "outcome": "success_special_primary_and_runoff",
            "raw_evidence": "sc/2026/sc-127017-en-summary.json",
            "notes": "Republican special primary and runoff retained separately.",
        },
        {
            "attempt_id": "sc-06-certification",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2026",
            "source_url": "https://scvotes.gov/about-the-sec/public-notices/",
            "method": "capture signed canvass and results-verification PDFs",
            "outcome": "success_certification_support",
            "raw_evidence": "sc/2026/sc-2026-primary-signed-certification.pdf",
            "notes": "Certification dates verified independently from ENR.",
        },
        {
            "attempt_id": "sc-07-candidate-roster",
            "attempted_on": date,
            "state_code": "SC",
            "cycle": "2024-2026",
            "source_url": "https://scvotes.gov/candidates/",
            "method": "seek contemporaneous certified federal primary roster",
            "outcome": "not_recovered",
            "raw_evidence": "",
            "notes": "Missing units remain explicit; no general nominee substitution.",
        },
        {
            "attempt_id": "tn-01-results-page",
            "attempted_on": date,
            "state_code": "TN",
            "cycle": "2024-2026",
            "source_url": TN_PUBLICATION,
            "method": "request official result inventory with browser headers",
            "outcome": "success_download_links",
            "raw_evidence": "discovery/tn-election-results.html",
            "notes": "Ordinary requests were blocked; browser-equivalent headers succeeded.",
        },
        {
            "attempt_id": "tn-02-2024-workbook",
            "attempted_on": date,
            "state_code": "TN",
            "cycle": "2024",
            "source_url": TN_SOURCE_URLS["2024"],
            "method": "download statewide results-by-precinct workbook",
            "outcome": "success_complete_primary",
            "raw_evidence": "tn/2024/20240801AllbyPrecinct.xlsx",
            "notes": "All 20 federal district-party units and named write-ins.",
        },
        {
            "attempt_id": "tn-03-2026-workbook",
            "attempted_on": date,
            "state_code": "TN",
            "cycle": "2026",
            "source_url": TN_SOURCE_URLS["2026"],
            "method": "download statewide results-by-precinct workbook",
            "outcome": "success_complete_primary",
            "raw_evidence": "tn/2026/20260806AllbyPrecinct.xlsx",
            "notes": "All 20 federal units under the revised district map.",
        },
        {
            "attempt_id": "tn-04-redistricting",
            "attempted_on": date,
            "state_code": "TN",
            "cycle": "2026",
            "source_url": (
                "https://sos.tn.gov/announcements/2026-congressional-redistricting"
            ),
            "method": "capture official revised-map notice",
            "outcome": "success_map_verified",
            "raw_evidence": "tn/2026/tn-2026-congressional-redistricting.html",
            "notes": "Confirms May 2026 map revision before the August primary.",
        },
    ]


def _validate_output(rows: list[dict]) -> None:
    if not rows:
        raise ValueError("No South Carolina/Tennessee primary results generated")
    seen = set()
    for row in rows:
        if set(row) != set(FIELDS):
            raise ValueError(f"Output schema differs: {row.get('candidate_name')}")
        key = (row["contest_id"], row["candidate_source_id"])
        if key in seen:
            raise ValueError(f"Duplicate candidate result: {key}")
        seen.add(key)
        if not isinstance(row["votes"], int) or row["votes"] < 0:
            raise ValueError(f"Invalid vote total: {row}")
        if row["write_in"] and row["party"]:
            raise ValueError(f"Write-in row asserts party affiliation: {row}")
        if len(row["source_sha256"]) != 64 or not row["source_locators"]:
            raise ValueError(f"Incomplete provenance: {row}")


def generate(
    raw_root: Path | None = None,
    analysis_root: Path | None = None,
) -> list[dict]:
    raw_root = raw_root or RAW_DIR / RAW_SUBDIR
    analysis_root = analysis_root or ANALYSIS_DATA_DIR / "congressional"
    rows = []
    for contest_id in SC_2024_CONTEST_IDS:
        rows.extend(parse_south_carolina_civera(
            raw_root / f"sc/2024/sc-contest-{contest_id}.csv"
        ))
    for source in SC_2026_SOURCES:
        rows.extend(parse_south_carolina_enr(
            raw_root / source["summary"],
            raw_root / source["config"],
            source,
        ))
    for year, date in (("2024", "2024-08-01"), ("2026", "2026-08-06")):
        filename = f"{date.replace('-', '')}AllbyPrecinct.xlsx"
        rows.extend(parse_tennessee(
            raw_root / "tn" / year / filename,
            {
                "date": date,
                "source_id": f"tn-{year}-statewide-primary-precinct-workbook",
                "url": TN_SOURCE_URLS[year],
            },
        ))
    rows.sort(key=lambda row: (
        row["state_code"],
        row["election_date"],
        row["stage"],
        row["chamber"],
        row["district"],
        row["primary_party"],
        row["candidate_name"],
    ))
    _validate_output(rows)
    coverage = _coverage_rows(rows)
    write_csv(analysis_root / RESULT_FILENAME, rows, FIELDS)
    manifest_dir = analysis_root / MANIFEST_SUBDIR
    write_csv(
        manifest_dir / "source_manifest.csv",
        _source_manifest(raw_root, rows),
        SOURCE_MANIFEST_FIELDS,
    )
    write_csv(manifest_dir / "attempt_ledger.csv", _attempt_rows(), ATTEMPT_FIELDS)
    write_csv(manifest_dir / "unresolved_gaps.csv", _gap_rows(coverage), GAP_FIELDS)
    write_csv(
        manifest_dir / "contest_manifest.csv",
        _contest_rows(rows),
        CONTEST_FIELDS,
    )
    write_csv(
        manifest_dir / "coverage_units.csv",
        coverage,
        COVERAGE_FIELDS,
    )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate official South Carolina/Tennessee primary results"
    )
    parser.parse_args()
    rows = generate()
    print(
        f"Wrote {len(rows)} candidate results to "
        f"{ANALYSIS_DATA_DIR / 'congressional' / RESULT_FILENAME}"
    )


if __name__ == "__main__":
    main()
