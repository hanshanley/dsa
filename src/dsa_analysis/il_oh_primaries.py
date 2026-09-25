"""Parse official Illinois and Ohio congressional primary results."""

import argparse
import csv
import hashlib
import html
import re
from collections import defaultdict
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .io import write_csv
from .paths import ANALYSIS_DATA_DIR, RAW_DIR
from .state_results import FIELDS

RAW_SUBDIR = "il_oh_primaries"
RESULT_FILENAME = "il_oh_primary_results.csv"

IL_SOURCES = (
    {
        "source_id": "il-2024-primary-federal",
        "date": "2024-03-19",
        "filename": "illinois/il-2024-primary-federal.html",
        "url": (
            "https://www.elections.il.gov/electionoperations/"
            "ElectionVoteTotals.aspx?ID=rfZ%2BuidMSDY%3D"
        ),
    },
    {
        "source_id": "il-2026-primary-federal",
        "date": "2026-03-17",
        "filename": "illinois/il-2026-primary-federal.html",
        "url": (
            "https://www.elections.il.gov/electionoperations/"
            "ElectionVoteTotals.aspx?ID=Z2J%2FvYpKX8w%3D"
        ),
    },
)

OH_2024_SOURCES = (
    {
        "source_id": "oh-2024-primary-democratic",
        "date": "2024-03-19",
        "party": "Democratic",
        "filename": (
            "ohio/archive/"
            "summary-level-official-results-primary-2024---democratic.xlsx"
        ),
        "url": (
            "https://www.ohiosos.gov/globalassets/elections/2024/pri/official/"
            "summary-level-official-results-primary-2024---democratic.xlsx"
        ),
        "archive_url": (
            "https://web.archive.org/web/20240420171450id_/"
            "https://www.ohiosos.gov/globalassets/elections/2024/pri/official/"
            "summary-level-official-results-primary-2024---democratic.xlsx"
        ),
    },
    {
        "source_id": "oh-2024-primary-republican",
        "date": "2024-03-19",
        "party": "Republican",
        "filename": (
            "ohio/archive/"
            "summary-level-results-primary-2024---republican.xlsx"
        ),
        "url": (
            "https://www.ohiosos.gov/globalassets/elections/2024/pri/official/"
            "summary-level-results-primary-2024---republican.xlsx"
        ),
        "archive_url": (
            "https://web.archive.org/web/20240420171458id_/"
            "https://www.ohiosos.gov/globalassets/elections/2024/pri/official/"
            "summary-level-results-primary-2024---republican.xlsx"
        ),
    },
)

OH_2026_FILES = (
    {
        "party": "Democratic",
        "blob_path": (
            "past-elections/2026/Primary+Special Election - May 5, 2026/group1/"
            "summary-level-official-results-2026-primary---democratic.xlsx"
        ),
    },
    {
        "party": "Libertarian",
        "blob_path": (
            "past-elections/2026/Primary+Special Election - May 5, 2026/group1/"
            "summary-level-official-results-2026-primary---libertarian.xlsx"
        ),
    },
    {
        "party": "Republican",
        "blob_path": (
            "past-elections/2026/Primary+Special Election - May 5, 2026/group1/"
            "summary-level-official-results-2026-primary---republican.xlsx"
        ),
    },
)

MANIFEST_FIELDS = [
    "source_id",
    "state_code",
    "election_date",
    "artifact_kind",
    "authority",
    "publication_url",
    "source_url",
    "raw_filename",
    "sha256",
    "acquisition_status",
    "verification_status",
    "candidate_rows",
    "contests",
    "coverage",
    "notes",
]

ACCOUNTING_FIELDS = [
    "cycle",
    "state_code",
    "chamber",
    "district",
    "term",
    "primary_party",
    "status",
    "candidate_rows",
    "source_ids",
    "notes",
]

ATTEMPT_FIELDS = [
    "attempt_id",
    "attempted_on",
    "state_code",
    "cycle",
    "source_url",
    "method",
    "outcome",
    "raw_evidence",
    "notes",
]

GAP_FIELDS = [
    "gap_id",
    "state_code",
    "cycle",
    "chamber",
    "requested_artifact",
    "coverage_status",
    "missing_scope",
    "blocker",
    "next_authoritative_step",
    "source_urls",
]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _combined_sha256(paths: list[Path]) -> str:
    digest = hashlib.sha256()
    for path in sorted(paths):
        digest.update(path.read_bytes())
    return digest.hexdigest()


def _votes(value: object) -> int:
    text = str(value or "").strip().replace(",", "")
    if not text or not text.isdigit():
        raise ValueError(f"Invalid vote total: {value!r}")
    return int(text)


class _IllinoisHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title_depth = 0
        self.title_parts: list[str] = []
        self.pending_title = ""
        self.in_table = False
        self.in_row = False
        self.in_cell = False
        self.cell_parts: list[str] = []
        self.cell_href = ""
        self.cells: list[dict[str, str]] = []
        self.row_number = 0
        self.rows: list[dict[str, object]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        classes = (attributes.get("class") or "").split()
        if tag == "div" and "gridview-title-bar" in classes:
            self.title_depth = 1
            self.title_parts = []
        elif tag == "div" and self.title_depth:
            self.title_depth += 1
        if tag == "table" and self.pending_title:
            self.in_table = True
            self.row_number = 0
        elif tag == "tr" and self.in_table:
            self.in_row = True
            self.cells = []
            self.row_number += 1
        elif tag == "td" and self.in_row:
            self.in_cell = True
            self.cell_parts = []
            self.cell_href = ""
        elif tag == "a" and self.in_cell:
            self.cell_href = attributes.get("href") or ""

    def handle_endtag(self, tag: str) -> None:
        if tag == "div" and self.title_depth:
            self.title_depth -= 1
            if self.title_depth == 0:
                title = " ".join(" ".join(self.title_parts).split())
                self.pending_title = title.removesuffix(" Download Results").strip()
        elif tag == "td" and self.in_cell:
            self.cells.append({
                "text": " ".join(" ".join(self.cell_parts).split()),
                "href": self.cell_href,
            })
            self.in_cell = False
        elif tag == "tr" and self.in_row:
            if (
                len(self.cells) >= 4
                and self.cells[0]["text"] not in {"Candidate", "Vote Totals"}
            ):
                self.rows.append({
                    "contest": self.pending_title,
                    "cells": self.cells,
                    "row_number": self.row_number,
                })
            self.in_row = False
        elif tag == "table" and self.in_table:
            self.in_table = False
            self.pending_title = ""

    def handle_data(self, data: str) -> None:
        if self.title_depth:
            self.title_parts.append(data)
        if self.in_cell:
            self.cell_parts.append(data)


def _illinois_candidate_details(path: Path) -> dict[tuple[str, str], dict]:
    output = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row_number, row in enumerate(csv.DictReader(handle), 2):
            key = (row["cycle"], row["candidate_source_id"])
            if key in output:
                raise ValueError(f"Duplicate Illinois candidate detail: {key}")
            output[key] = {**row, "row_number": row_number}
    return output


def parse_illinois_results(
    path: Path,
    details_path: Path,
    source: dict,
) -> list[dict]:
    parser = _IllinoisHTMLParser()
    parser.feed(path.read_text(encoding="utf-8"))
    details = _illinois_candidate_details(details_path)
    source_hash = _combined_sha256([path, details_path])
    output = []
    for parsed in parser.rows:
        contest_name = str(parsed["contest"])
        senate = contest_name == "UNITED STATES SENATOR"
        house = re.fullmatch(r"(\d+)(?:ST|ND|RD|TH) CONGRESS", contest_name)
        if not senate and house is None:
            continue
        cells = parsed["cells"]
        candidate_href = html.unescape(cells[0]["href"])
        if not candidate_href:
            raise ValueError(f"Illinois candidate lacks detail link: {contest_name}")
        candidate_id = parse_qs(urlparse(candidate_href).query)["CandidateID"][0]
        detail = details[(source["date"][:4], candidate_id)]
        candidate_name = cells[0]["text"]
        party = cells[1]["text"].title()
        if detail["candidate_name"].casefold() != candidate_name.casefold():
            raise ValueError(f"Illinois candidate detail mismatch: {candidate_name}")
        if detail["party"].title() != party:
            raise ValueError(f"Illinois candidate party mismatch: {candidate_name}")
        write_in = detail["write_in"] == "Yes"
        chamber = "Senate" if senate else "House"
        district = "S" if senate else house[1].zfill(2)
        contest_id = (
            f"il-{source['date']}-{chamber.lower()}-{district.lower()}-regular-"
            f"{party.lower().replace(' ', '-')}"
        )
        output.append({
            "contest_id": contest_id,
            "cycle": source["date"][:4],
            "election_date": source["date"],
            "stage": "primary",
            "state_code": "IL",
            "chamber": chamber,
            "district": district,
            "term": "regular",
            "election_system": "partisan_primary",
            "primary_party": party,
            "candidate_name": candidate_name,
            "candidate_source_id": candidate_id,
            "party": "" if write_in else party,
            "write_in": write_in,
            "votes": _votes(cells[2]["text"]),
            "source_id": source["source_id"],
            "source_url": source["url"],
            "source_sha256": source_hash,
            "source_locators": (
                f"{path.name}:{contest_name}:row-{parsed['row_number']} | "
                f"{details_path.name}:{detail['row_number']}"
            ),
            "reporting_units": 1,
        })
    house_districts = {row["district"] for row in output if row["chamber"] == "House"}
    if house_districts != {f"{district:02}" for district in range(1, 18)}:
        raise ValueError("Illinois result page is missing congressional districts")
    senate_rows = [row for row in output if row["chamber"] == "Senate"]
    if source["date"].startswith("2026") and not senate_rows:
        raise ValueError("Illinois 2026 results are missing U.S. Senate")
    if source["date"].startswith("2024") and senate_rows:
        raise ValueError("Illinois 2024 page unexpectedly contains U.S. Senate")
    return output


def parse_ohio_results(path: Path, source: dict) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["U.S. Congress"]
        office_row = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        candidate_row = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        total_row = list(next(sheet.iter_rows(min_row=3, max_row=3, values_only=True)))
    finally:
        workbook.close()
    current_office = ""
    output = []
    source_hash = _sha256(path)
    for index in range(6, len(candidate_row)):
        if office_row[index]:
            current_office = " ".join(str(office_row[index]).split())
        candidate_label = str(candidate_row[index] or "").strip()
        if not candidate_label:
            continue
        senate = current_office.startswith("U.S. Senator")
        house = re.search(
            r"Representative to Congress - District (\d+)",
            current_office,
            re.I,
        )
        if not senate and house is None:
            continue
        write_in = candidate_label.endswith("(WI)")
        if write_in:
            candidate_name = candidate_label.removesuffix("(WI)").strip()
        else:
            candidate_name = re.sub(r"\s+\([DRL]\)$", "", candidate_label).strip()
        chamber = "Senate" if senate else "House"
        district = "S" if senate else house[1].zfill(2)
        term = "special" if "Unexpired Term" in current_office else "regular"
        party = source["party"]
        contest_id = (
            f"oh-{source['date']}-{chamber.lower()}-{district.lower()}-{term}-"
            f"{party.lower()}"
        )
        output.append({
            "contest_id": contest_id,
            "cycle": source["date"][:4],
            "election_date": source["date"],
            "stage": "primary",
            "state_code": "OH",
            "chamber": chamber,
            "district": district,
            "term": term,
            "election_system": "partisan_primary",
            "primary_party": party,
            "candidate_name": candidate_name,
            "candidate_source_id": (
                f"oh-{source['date']}-{party.lower()}-{get_column_letter(index + 1)}"
            ),
            "party": "" if write_in else party,
            "write_in": write_in,
            "votes": _votes(total_row[index]),
            "source_id": source["source_id"],
            "source_url": source["url"],
            "source_sha256": source_hash,
            "source_locators": f"{sheet.title}:{get_column_letter(index + 1)}1-3",
            "reporting_units": 1,
        })
    if not output:
        raise ValueError(f"No Ohio federal candidates parsed from {path}")
    return output


def _contest_accounting(results: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in results:
        groups[row["contest_id"]].append(row)
    output = []
    for group in groups.values():
        first = group[0]
        output.append({
            "cycle": first["cycle"],
            "state_code": first["state_code"],
            "chamber": first["chamber"],
            "district": first["district"],
            "term": first["term"],
            "primary_party": first["primary_party"],
            "status": (
                "complete_unopposed_result"
                if len(group) == 1
                else "complete_official_results"
            ),
            "candidate_rows": len(group),
            "source_ids": " | ".join(sorted({row["source_id"] for row in group})),
            "notes": (
                "The official result source reports one candidate in this party contest."
                if len(group) == 1
                else ""
            ),
        })
    observed = {
        (
            row["cycle"],
            row["state_code"],
            row["chamber"],
            row["district"],
            row["term"],
            row["primary_party"],
        )
        for row in output
    }
    expected_slots = []
    for cycle in ("2024", "2026"):
        expected_slots.extend(
            (cycle, "IL", "House", f"{district:02}", "regular", party)
            for district in range(1, 18)
            for party in ("Democratic", "Republican")
        )
    expected_slots.extend(
        ("2026", "IL", "Senate", "S", "regular", party)
        for party in ("Democratic", "Republican")
    )
    expected_slots.extend(
        ("2024", "OH", "House", f"{district:02}", "regular", party)
        for district in range(1, 16)
        for party in ("Democratic", "Republican")
    )
    expected_slots.extend(
        ("2024", "OH", "Senate", "S", "regular", party)
        for party in ("Democratic", "Republican")
    )
    expected_slots.extend(
        ("2024", "OH", "House", "06", "special", party)
        for party in ("Democratic", "Republican")
    )
    for slot in expected_slots:
        if slot in observed:
            continue
        cycle, state, chamber, district, term, party = slot
        output.append({
            "cycle": cycle,
            "state_code": state,
            "chamber": chamber,
            "district": district,
            "term": term,
            "primary_party": party,
            "status": "no_reported_contest_not_inferred",
            "candidate_rows": 0,
            "source_ids": (
                f"{state.lower()}-{cycle}-official-result-source"
            ),
            "notes": (
                "The complete official result source has no row for this district-party "
                "slot. This is not labeled no-primary or uncontested without explicit evidence."
            ),
        })
    for party in ("Democratic", "Libertarian", "Republican"):
        for district in range(1, 16):
            output.append({
                "cycle": "2026",
                "state_code": "OH",
                "chamber": "House",
                "district": f"{district:02}",
                "term": "regular",
                "primary_party": party,
                "status": "blocked_certified_workbook_unretrieved",
                "candidate_rows": "",
                "source_ids": "oh-2026-files-index",
                "notes": (
                    "The party workbook is listed by the official portal but blocked by "
                    "the current security layer. This is not evidence that a primary was absent."
                ),
            })
        output.append({
            "cycle": "2026",
            "state_code": "OH",
            "chamber": "Senate",
            "district": "S",
            "term": "special",
            "primary_party": party,
            "status": "blocked_certified_workbook_unretrieved",
            "candidate_rows": "",
            "source_ids": "oh-2026-files-index",
            "notes": (
                "Unexpired term ending January 3, 2029. Candidate and vote rows remain "
                "unknown until the certified party workbook is recovered."
            ),
        })
    return sorted(output, key=lambda row: (
        row["cycle"],
        row["state_code"],
        row["chamber"],
        row["district"],
        row["term"],
        row["primary_party"],
    ))


def _manifest(raw_dir: Path, results: list[dict]) -> list[dict]:
    result_groups = defaultdict(list)
    for row in results:
        result_groups[row["source_id"]].append(row)
    entries = []
    details_path = raw_dir / "illinois" / "il-candidate-details.csv"
    for source in IL_SOURCES:
        path = raw_dir / source["filename"]
        rows = result_groups[source["source_id"]]
        entries.append({
            "source_id": source["source_id"],
            "state_code": "IL",
            "election_date": source["date"],
            "artifact_kind": "official_statewide_results",
            "authority": "Illinois State Board of Elections",
            "publication_url": source["url"],
            "source_url": source["url"],
            "raw_filename": source["filename"],
            "sha256": _combined_sha256([path, details_path]),
            "acquisition_status": "recovered",
            "verification_status": "complete_all_congressional_districts",
            "candidate_rows": len(rows),
            "contests": len({row["contest_id"] for row in rows}),
            "coverage": (
                "all_17_US_House_districts"
                + ("_and_US_Senate" if source["date"].startswith("2026") else "")
            ),
            "notes": (
                "Candidate-detail pages identify write-ins. Named write-ins retain "
                "primary_party as the ballot category and have blank party."
            ),
        })
    entries.append({
        "source_id": "il-primary-candidate-details",
        "state_code": "IL",
        "election_date": "2024-03-19 | 2026-03-17",
        "artifact_kind": "candidate_status_support",
        "authority": "Illinois State Board of Elections",
        "publication_url": (
            "https://www.elections.il.gov/ElectionOperations/CandidateFilingSearch.aspx"
        ),
        "source_url": "individual_candidate_detail_urls_recorded_in_raw_csv",
        "raw_filename": "illinois/il-candidate-details.csv",
        "sha256": _sha256(details_path),
        "acquisition_status": "recovered",
        "verification_status": "write_in_status_verified",
        "candidate_rows": sum(1 for _ in details_path.open()) - 1,
        "contests": "",
        "coverage": "all_candidates_in_the_two_federal_result_pages",
        "notes": "Individual raw detail pages are persisted under candidate_details/.",
    })
    for source in OH_2024_SOURCES:
        path = raw_dir / source["filename"]
        rows = result_groups[source["source_id"]]
        entries.append({
            "source_id": source["source_id"],
            "state_code": "OH",
            "election_date": source["date"],
            "artifact_kind": "archived_official_statewide_results",
            "authority": "Ohio Secretary of State",
            "publication_url": source["url"],
            "source_url": source["archive_url"],
            "raw_filename": source["filename"],
            "sha256": _sha256(path),
            "acquisition_status": "recovered_from_archived_official_publication",
            "verification_status": "complete_official_party_workbook",
            "candidate_rows": len(rows),
            "contests": len({row["contest_id"] for row in rows}),
            "coverage": (
                "US_Senate_all_15_US_House_districts_and_CD06_special_primary"
            ),
            "notes": "Archived capture preserves the original Secretary of State workbook.",
        })
    index_path = raw_dir / "ohio" / "ohio-files-index.json"
    entries.append({
        "source_id": "oh-2026-files-index",
        "state_code": "OH",
        "election_date": "2026-05-05",
        "artifact_kind": "official_download_index",
        "authority": "Ohio Secretary of State",
        "publication_url": "https://data.ohiosos.gov/portal/past-election-results",
        "source_url": (
            "https://publicfiles.ohiosos.gov/election-results/files-index.json"
        ),
        "raw_filename": "ohio/ohio-files-index.json",
        "sha256": _sha256(index_path),
        "acquisition_status": "recovered_in_full_browser_session",
        "verification_status": "certified_workbook_urls_verified_but_files_blocked",
        "candidate_rows": 0,
        "contests": 0,
        "coverage": "Democratic_Libertarian_Republican_2026_party_workbook_routes",
        "notes": "The index is recovered; the three workbook requests hit the security layer.",
    })
    blocked_path = raw_dir / "ohio" / "oh-2026-primary-democratic-blocked.html"
    for source in OH_2026_FILES:
        party_slug = source["party"].lower()
        entries.append({
            "source_id": f"oh-2026-primary-{party_slug}",
            "state_code": "OH",
            "election_date": "2026-05-05",
            "artifact_kind": "certified_party_workbook",
            "authority": "Ohio Secretary of State",
            "publication_url": "https://data.ohiosos.gov/portal/past-election-results",
            "source_url": (
                "https://publicfiles.ohiosos.gov/election-results/"
                + source["blob_path"]
            ),
            "raw_filename": (
                "ohio/oh-2026-primary-democratic-blocked.html"
                if party_slug == "democratic"
                else ""
            ),
            "sha256": _sha256(blocked_path) if party_slug == "democratic" else "",
            "acquisition_status": "blocked_cloudflare_security_layer",
            "verification_status": "unresolved_no_result_rows_emitted",
            "candidate_rows": 0,
            "contests": 0,
            "coverage": "unknown_until_workbook_recovered",
            "notes": "No absent contest or unopposed status is inferred from this blocker.",
        })
    return entries


def _attempts() -> list[dict]:
    return [
        {
            "attempt_id": "il-01",
            "attempted_on": "2026-09-22",
            "state_code": "IL",
            "cycle": "2024",
            "source_url": IL_SOURCES[0]["url"],
            "method": "direct_official_html",
            "outcome": "success",
            "raw_evidence": IL_SOURCES[0]["filename"],
            "notes": "All 17 congressional district result tables recovered.",
        },
        {
            "attempt_id": "il-02",
            "attempted_on": "2026-09-22",
            "state_code": "IL",
            "cycle": "2026",
            "source_url": IL_SOURCES[1]["url"],
            "method": "direct_official_html",
            "outcome": "success",
            "raw_evidence": IL_SOURCES[1]["filename"],
            "notes": "U.S. Senate and all 17 congressional districts recovered.",
        },
        {
            "attempt_id": "il-03",
            "attempted_on": "2026-09-22",
            "state_code": "IL",
            "cycle": "2024-2026",
            "source_url": (
                "https://www.elections.il.gov/ElectionOperations/CandidateDetailEO.aspx"
            ),
            "method": "individual_official_candidate_details",
            "outcome": "success",
            "raw_evidence": "illinois/il-candidate-details.csv",
            "notes": "Verified write-in status for every reported federal candidate.",
        },
        {
            "attempt_id": "oh-01",
            "attempted_on": "2026-09-22",
            "state_code": "OH",
            "cycle": "2024",
            "source_url": OH_2024_SOURCES[0]["url"],
            "method": "direct_download",
            "outcome": "blocked_security_layer",
            "raw_evidence": "ohio/ohio-past-election-results-blocked.html",
            "notes": "Live host returned the Secretary of State maintenance/security page.",
        },
        {
            "attempt_id": "oh-02",
            "attempted_on": "2026-09-22",
            "state_code": "OH",
            "cycle": "2024",
            "source_url": OH_2024_SOURCES[0]["archive_url"],
            "method": "archived_official_publication",
            "outcome": "success",
            "raw_evidence": "ohio/archive/",
            "notes": "Both official party summary workbooks recovered.",
        },
        {
            "attempt_id": "oh-03",
            "attempted_on": "2026-09-22",
            "state_code": "OH",
            "cycle": "2026",
            "source_url": "https://data.ohiosos.gov/portal/past-election-results",
            "method": "full_browser_session",
            "outcome": "index_success_workbooks_blocked",
            "raw_evidence": "ohio/ohio-files-index.json",
            "notes": "Verified all three certified party workbook paths.",
        },
    ]


def _gaps() -> list[dict]:
    return [
        {
            "gap_id": "oh-2026-certified-primary-results",
            "state_code": "OH",
            "cycle": "2026",
            "chamber": "House and Senate",
            "requested_artifact": "certified_primary_results",
            "coverage_status": "unresolved",
            "missing_scope": (
                "All Democratic, Libertarian, and Republican congressional candidate rows"
            ),
            "blocker": (
                "The official portal index was recovered, but each certified workbook "
                "request is intercepted by the current security/maintenance layer."
            ),
            "next_authoritative_step": (
                "Retry the exact publicfiles workbook URLs or obtain an archived official capture."
            ),
            "source_urls": " | ".join(
                "https://publicfiles.ohiosos.gov/election-results/" + source["blob_path"]
                for source in OH_2026_FILES
            ),
        },
    ]


def collect_il_oh_primaries(
    raw_dir: Path | None = None,
    output_dir: Path | None = None,
) -> dict:
    raw_dir = raw_dir or RAW_DIR / RAW_SUBDIR
    output_dir = output_dir or ANALYSIS_DATA_DIR / "congressional"
    ledger_dir = output_dir / "il_oh"
    output_dir.mkdir(parents=True, exist_ok=True)
    ledger_dir.mkdir(parents=True, exist_ok=True)

    details_path = raw_dir / "illinois" / "il-candidate-details.csv"
    results = []
    for source in IL_SOURCES:
        results.extend(
            parse_illinois_results(raw_dir / source["filename"], details_path, source)
        )
    for source in OH_2024_SOURCES:
        results.extend(parse_ohio_results(raw_dir / source["filename"], source))

    il_by_cycle = defaultdict(list)
    oh_rows = []
    for row in results:
        if row["state_code"] == "IL":
            il_by_cycle[row["cycle"]].append(row)
        else:
            oh_rows.append(row)
    for cycle, rows in il_by_cycle.items():
        districts = {row["district"] for row in rows if row["chamber"] == "House"}
        if districts != {f"{district:02}" for district in range(1, 18)}:
            raise ValueError(f"Illinois {cycle} district coverage is incomplete")
    oh_districts = {row["district"] for row in oh_rows if row["chamber"] == "House"}
    if oh_districts != {f"{district:02}" for district in range(1, 16)}:
        raise ValueError("Ohio 2024 district coverage is incomplete")
    if not any(row["chamber"] == "Senate" for row in oh_rows):
        raise ValueError("Ohio 2024 U.S. Senate primary is missing")
    special = {
        (row["primary_party"], row["district"])
        for row in oh_rows
        if row["term"] == "special"
    }
    if special != {("Democratic", "06"), ("Republican", "06")}:
        raise ValueError("Ohio 2024 CD06 special-primary coverage is incomplete")

    results.sort(key=lambda row: (
        row["cycle"],
        row["state_code"],
        row["chamber"],
        row["district"],
        row["term"],
        row["primary_party"],
        row["candidate_name"],
    ))
    write_csv(output_dir / RESULT_FILENAME, results, FIELDS)
    accounting = _contest_accounting(results)
    write_csv(ledger_dir / "contest_accounting.csv", accounting, ACCOUNTING_FIELDS)
    manifest = _manifest(raw_dir, results)
    write_csv(ledger_dir / "source_manifest.csv", manifest, MANIFEST_FIELDS)
    write_csv(ledger_dir / "attempt_ledger.csv", _attempts(), ATTEMPT_FIELDS)
    gaps = _gaps()
    write_csv(ledger_dir / "unresolved_gaps.csv", gaps, GAP_FIELDS)
    return {
        "result_rows": len(results),
        "result_contests": len({row["contest_id"] for row in results}),
        "accounting_rows": len(accounting),
        "manifest_sources": len(manifest),
        "unresolved_gaps": len(gaps),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", type=Path)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    summary = collect_il_oh_primaries(args.raw_dir, args.output_dir)
    print(" ".join(f"{key}={value}" for key, value in summary.items()))


if __name__ == "__main__":
    main()
