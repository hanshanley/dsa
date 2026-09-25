"""Nationwide congressional source inventory, independent of DSA endorsements."""

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError

from openpyxl import load_workbook

from .fec_presidential import _download_workbook
from .document_corpus import _extract_pdf_text
from .io import read_csv, read_json, write_csv
from .paths import ANALYSIS_DATA_DIR, CONFIG_DIR, RAW_DIR
from .race_registry import STATE_NAME_BY_CODE
from .state_results import PRIMARY_RESULT_STAGES
from .congressional_accounting import export_party_accounting
from .congressional_reconciliation import reconcile_dsa_congressional

FEC_INDEX = (
    "https://www.fec.gov/introduction-campaign-finance/"
    "election-results-and-voting-information/"
)
TERRITORIES = {"AS", "DC", "GU", "MP", "PR", "VI"}
STATE_CODES = set(STATE_NAME_BY_CODE) | TERRITORIES
STAGES = {"primary": "PRIMARY VOTES", "primary_runoff": "RUNOFF VOTES",
          "general": "GENERAL VOTES"}
SPECIAL_URL = "https://www.fec.gov/resources/cms-content/documents/federalspecialelectionslist.pdf"
BALLOT_2024_URL = "https://www.fec.gov/documents/5548/2024congressgecands.xlsx"
SENATE_CLASSES = {1: "I", 2: "II", 3: "III"}
RESULT_FIELDS = [
    "source_record_id", "workbook_cycle", "chamber", "state_code", "state",
    "district", "district_raw", "term", "candidate_name", "fec_candidate_id",
    "party_raw", "row_kind", "primary_votes", "primary_votes_raw",
    "primary_runoff_votes", "primary_runoff_votes_raw", "general_votes",
    "general_votes_raw", "general_runoff_votes", "general_runoff_votes_raw",
    "primary_winner_raw", "general_winner_raw", "footnotes", "source_url",
    "source_sha256", "source_sheet", "source_row", "raw_cells_json",
]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _number(value: object) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value < 0 or int(value) != value:
            raise ValueError(f"Invalid vote count: {value!r}")
        return str(int(value))
    return ""


def _district(value: str) -> str:
    match = re.match(r"^(S|\d{1,2})(?:\b|[- ])", value, re.I)
    if not match:
        return ""
    return match[1].upper() if match[1].upper() == "S" else match[1].zfill(2)


def parse_congressional_workbook(path: Path, cycle: int, source_url: str) -> list[dict]:
    workbook_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    rows = []
    with path.open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                values = sheet.iter_rows(values_only=True)
                header = next(values, ())
                labels = [_text(value).upper() for value in header]
                if not {"STATE ABBREVIATION", "CANDIDATE NAME", "PRIMARY VOTES"} <= set(labels):
                    continue
                if "PRES" in sheet.title.upper():
                    continue
                indices = {label: index for index, label in enumerate(labels) if label}
                district_column = "DISTRICT" if "DISTRICT" in indices else "D"
                if district_column not in indices:
                    raise ValueError(f"{path.name}:{sheet.title}: unrecognized district column")
                for row_number, values in enumerate(values, 2):
                    def get(label: str) -> object:
                        return values[indices[label]] if label in indices else None

                    state = _text(get("STATE ABBREVIATION"))
                    name = _text(get("CANDIDATE NAME"))
                    party = _text(get("PARTY"))
                    if state not in STATE_CODES or not name or not party:
                        continue
                    raw_district = _text(get(district_column))
                    district = _district(raw_district)
                    footnotes = " | ".join(
                        _text(values[index]) for label, index in indices.items()
                        if "FOOTNOTE" in label and values[index] is not None
                    )
                    special_sheet = (
                        "SPECIAL" in sheet.title.upper() or "NEW ELECTION" in sheet.title.upper()
                    )
                    term = (
                        "special" if special_sheet or "UNEXPIRED" in raw_district.upper()
                        or "SPECIAL ELECTION" in footnotes.upper()
                        else "regular"
                    )
                    if district == "S":
                        chamber = "Senate"
                    elif "SENATE RESULTS" in sheet.title.upper():
                        chamber = "Senate"
                    else:
                        chamber = "House"
                    record_id = hashlib.sha256(
                        f"{cycle}\n{sheet.title}\n{row_number}".encode()
                    ).hexdigest()[:24]
                    record = {
                        "source_record_id": record_id,
                        "workbook_cycle": cycle,
                        "chamber": chamber,
                        "state_code": state,
                        "state": _text(get("STATE")),
                        "district": district,
                        "district_raw": raw_district,
                        "term": term,
                        "candidate_name": name,
                        "fec_candidate_id": _text(get("FEC ID") or get("FEC ID#")),
                        "party_raw": party,
                        "row_kind": (
                            "aggregate" if re.search(
                                r"^(scattered|other|write[- ]?ins?|none of (these|the above)"
                                r"|all (other|write)|total)\b", name, re.I
                            ) else "candidate"
                        ),
                        "footnotes": footnotes,
                        "source_url": source_url,
                        "source_sha256": workbook_hash,
                        "source_sheet": sheet.title,
                        "source_row": row_number,
                        "raw_cells_json": json.dumps(
                            [_text(value) for value in values], ensure_ascii=True
                        ),
                    }
                    for stage, label in STAGES.items():
                        record[f"{stage}_votes"] = _number(get(label))
                        record[f"{stage}_votes_raw"] = _text(get(label))
                    runoff_label = next(
                        (label for label in indices if label.startswith("GE RUNOFF ELECTION VOTES")),
                        "",
                    )
                    record["general_runoff_votes"] = _number(get(runoff_label))
                    record["general_runoff_votes_raw"] = _text(get(runoff_label))
                    for stage, prefix in (("primary", "PE"), ("general", "GE")):
                        winner_label = next(
                            (label for label in indices if label == f"{prefix} WINNER INDICATOR"
                             or (label.startswith(f"SPECIAL {prefix}") and "WINNER" in label)),
                            "",
                        )
                        record[f"{stage}_winner_raw"] = _text(get(winner_label))
                    rows.append(record)
        finally:
            workbook.close()
    if not rows:
        raise ValueError(f"{path.name}: no congressional result records found")
    return rows


def _regular_seats(rows: list[dict], cycle: int, chamber: str) -> set[tuple[str, str]]:
    return {
        (row["state_code"], row["district"])
        for row in rows
        if row["workbook_cycle"] == cycle and row["term"] == "regular"
        and row["chamber"] == chamber and row["district"]
    }


def parse_senate_class(html: str, senate_class: int) -> set[str]:
    states = set(re.findall(r"\([RDI]-([A-Z]{2})\)", html))
    expected = 34 if senate_class == 3 else 33
    if len(states) != expected or not states <= STATE_CODES - TERRITORIES:
        raise ValueError(f"Senate Class {senate_class}: expected {expected} distinct states")
    return states


def parse_general_ballot_workbook(path: Path) -> list[dict]:
    output = []
    with path.open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                columns = {_text(value).upper(): i for i, value in enumerate(next(rows, ()))}
                required = {"STATE ABBREVIATION", "DISTRICT", "CANDIDATE NAME", "PARTY"}
                if not required <= columns.keys():
                    continue
                for number, values in enumerate(rows, 2):
                    state = _text(values[columns["STATE ABBREVIATION"]])
                    candidate = _text(values[columns["CANDIDATE NAME"]])
                    if state not in STATE_CODES or not candidate:
                        continue
                    district_raw = _text(values[columns["DISTRICT"]])
                    output.append({
                        "cycle": 2024, "stage": "general_ballot_only",
                        "state_code": state,
                        "district": _district(district_raw),
                        "district_raw": district_raw,
                        "candidate_name": candidate,
                        "fec_candidate_id": _text(values[columns["FEC ID#"]]),
                        "party_raw": _text(values[columns["PARTY"]]),
                        "source_url": BALLOT_2024_URL,
                        "source_sheet": sheet.title,
                        "source_row": number,
                        "raw_cells_json": json.dumps([_text(value) for value in values]),
                    })
        finally:
            workbook.close()
    if not output:
        raise ValueError("No congressional candidates in the 2024 general ballot workbook")
    return output


def source_identity_issues(rows: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in rows:
        for candidate_id in re.findall(r"\b[HS]\d[A-Z0-9]{7}\b", row["fec_candidate_id"]):
            groups[(row["source_url"], candidate_id)].append(row)
    issues = []
    for (source, candidate_id), group in sorted(groups.items()):
        names = {" ".join(row["candidate_name"].casefold().split()) for row in group}
        if len(names) < 2:
            continue
        issues.append({
            "issue": "fec_id_with_multiple_name_spellings_requires_review",
            "fec_candidate_id": candidate_id,
            "source_url": source,
            "names": " | ".join(sorted({row["candidate_name"] for row in group})),
            "locators": " | ".join(sorted({
                f'{row["source_sheet"]}:{row["source_row"]}' for row in group
            })),
        })
    return issues


def export_longitudinal_primary_records(
    output_dir: Path, fec_rows: list[dict], state_rows: list[dict],
) -> int:
    fields = [
        "observation_id", "cycle", "cycle_basis", "election_date", "election_date_status",
        "state_code", "chamber", "district", "term", "stage", "primary_party",
        "party_raw", "candidate_name", "candidate_source_id", "ballot_option_kind",
        "votes", "votes_raw", "observation_status", "source_url", "source_sha256",
        "source_locators", "notes",
    ]
    output = []
    for row in fec_rows:
        for stage in ("primary", "primary_runoff"):
            raw = row[f"{stage}_votes_raw"]
            if not raw:
                continue
            numeric = row[f"{stage}_votes"]
            output.append({
                "observation_id": f'fec-{row["source_record_id"]}-{stage}',
                "cycle": row["workbook_cycle"], "cycle_basis": "FEC_publication_cycle",
                "election_date": "", "election_date_status": "not_normalized_from_source",
                "state_code": row["state_code"], "chamber": row["chamber"],
                "district": row["district"], "term": row["term"], "stage": stage,
                "primary_party": "", "party_raw": row["party_raw"],
                "candidate_name": row["candidate_name"],
                "candidate_source_id": row["fec_candidate_id"],
                "ballot_option_kind": row["row_kind"],
                "votes": numeric, "votes_raw": raw,
                "observation_status": "official_numeric_return" if numeric != "" else "source_nomination_flag",
                "source_url": row["source_url"], "source_sha256": row["source_sha256"],
                "source_locators": f'{row["source_sheet"]}:{row["source_row"]}',
                "notes": row["footnotes"],
            })
    for index, row in enumerate(state_rows, 2):
        key = f'{row["source_id"]}\n{row["contest_id"]}\n{row["candidate_source_id"]}'
        output.append({
            "observation_id": "state-" + hashlib.sha256(key.encode()).hexdigest()[:24],
            "cycle": row["cycle"], "cycle_basis": "dated_election",
            "election_date": row["election_date"], "election_date_status": "source_verified",
            "state_code": row["state_code"], "chamber": row["chamber"],
            "district": row["district"], "term": row["term"], "stage": row["stage"],
            "primary_party": row["primary_party"], "party_raw": row["party"],
            "candidate_name": row["candidate_name"],
            "candidate_source_id": row["candidate_source_id"],
            "ballot_option_kind": "write_in_option" if str(row["write_in"]).lower() == "true" else "reported_candidate",
            "votes": row["votes"], "votes_raw": row["votes"],
            "observation_status": "superseded_numeric_return" if row["stage"] == "superseded_primary" else "official_numeric_return",
            "source_url": row["source_url"], "source_sha256": row["source_sha256"],
            "source_locators": row["source_locators"],
            "notes": f"state_primary_results.csv:{index}; source observations are not deduplicated races",
        })
    if len({row["observation_id"] for row in output}) != len(output):
        raise ValueError("Longitudinal primary records have duplicate observation IDs")
    write_csv(output_dir / "all_primary_records.csv", output, fields)
    return len(output)


def parse_special_elections(text: str, start: int, cutoff: str) -> list[dict]:
    updated = re.search(r"Updated by:[^\n]*,\s*(\d{1,2}/\d{1,2}/\d{4})", text)
    if not updated:
        raise ValueError("Special-election calendar lacks an identifiable update date")
    updated_on = datetime.strptime(updated[1], "%m/%d/%Y").date().isoformat()
    year = None
    page = ""
    output = []
    in_notes = False
    for line in text.splitlines():
        if line.startswith("[PDF page "):
            page = line.split("]", 1)[0].lstrip("[")
        if line.strip() == "NOTES:":
            in_notes = True
        if in_notes:
            continue
        if re.fullmatch(r"20\d{2}", line.strip()):
            year = int(line.strip())
            continue
        match = re.match(
            r"^([A-Z]{2})/(S|AL|\d{1,2})([*#]*)\s+(?:##\s+)?"
            r"(\d{1,2}/\d{1,2})([*#]*)\s+(.*)", line,
        )
        if not match or year is None or year < start:
            continue
        state, district, seat_markers, month_day, markers, detail = match.groups()
        election_date = datetime.strptime(f"{year}/{month_day}", "%Y/%m/%d").date().isoformat()
        if election_date > cutoff:
            continue
        output.append({
            "special_election_id": f"{state.lower()}-{district.lower()}-{election_date}",
            "state_code": state,
            "chamber": "Senate" if district == "S" else "House",
            "district": "00" if district == "AL" else district.zfill(2) if district != "S" else "S",
            "reported_election_date": election_date,
            "date_type": (
                "special_general_runoff" if "**" in seat_markers + markers else "special_general"
            ),
            "primary_date": "",
            "primary_roster_status": "not_reconciled",
            "source_updated_on": updated_on,
            "source_url": SPECIAL_URL,
            "locator": page,
            "source_text": line,
            "detail": detail,
        })
    if not output:
        raise ValueError("No in-window federal special elections parsed")
    return output


def build_congressional_inventory(
    *, download: bool = False,
    raw_dir: Path = RAW_DIR / "fec",
    output_dir: Path = ANALYSIS_DATA_DIR / "congressional",
) -> dict:
    config = read_json(CONFIG_DIR / "sources.json")
    start = int(config["study_start"][:4])
    end = int(config["research_cutoff"][:4])
    cycles = list(range(start + start % 2, end + 1, 2))
    raw_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    sources = []
    manifest_path = output_dir / "source_manifest.csv"
    previous_sources = {
        int(row["cycle"]): row for row in read_csv(manifest_path)
    } if manifest_path.exists() else {}
    records = []
    gaps = []
    senate_classes = {}
    for number, roman in SENATE_CLASSES.items():
        url = f"https://www.senate.gov/senators/Class_{roman}.htm"
        path = raw_dir / f"senate_class_{roman.lower()}.html"
        if download:
            payload = _download_workbook(url)
            parse_senate_class(payload.decode(), number)
            path.write_bytes(payload)
        if path.exists():
            senate_classes[number] = parse_senate_class(path.read_text(), number)
        else:
            gaps.append({
                "cycle": "", "state_code": "", "chamber": "Senate", "district": "",
                "issue": "missing_independent_senate_class_inventory",
                "source_url": url, "detail": f"Class {number}",
            })
    if len(senate_classes) == 3:
        memberships = Counter(state for states in senate_classes.values() for state in states)
        if set(memberships.values()) != {2} or len(memberships) != 50:
            raise ValueError("Senate class inventories must cover exactly two seats in all 50 states")
    ballot_path = raw_dir / "2024congressgecands.xlsx"
    if download:
        ballot_path.write_bytes(_download_workbook(BALLOT_2024_URL))
    ballot_rows = parse_general_ballot_workbook(ballot_path) if ballot_path.exists() else []
    if ballot_rows:
        write_csv(output_dir / "general_ballot_2024.csv", ballot_rows, list(ballot_rows[0]))
    for cycle in cycles:
        url = f"https://www.fec.gov/resources/cms-content/documents/federalelections{cycle}.xlsx"
        path = raw_dir / f"federalelections{cycle}.xlsx"
        error = ""
        downloaded = False
        if download:
            try:
                payload = _download_workbook(url)
                # Validate before replacing any previously usable source.
                temporary = path.with_suffix(".download.xlsx")
                temporary.write_bytes(payload)
                try:
                    parse_congressional_workbook(temporary, cycle, url)
                    temporary.replace(path)
                finally:
                    temporary.unlink(missing_ok=True)
                downloaded = True
            except (HTTPError, URLError, TimeoutError) as caught:
                error = f"{type(caught).__name__}: {caught}"
        status = (
            "downloaded" if downloaded else "cached_refresh_failed" if error and path.exists()
            else "cached" if path.exists() else "unavailable" if error else "not_downloaded"
        )
        source = {
            "cycle": cycle, "source_url": url, "status": status, "error": error,
            "checked_at": datetime.now(UTC).isoformat() if download else "",
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else "",
        }
        if not download and cycle in previous_sources:
            prior = previous_sources[cycle]
            if prior["sha256"] == source["sha256"]:
                source = prior
        sources.append(source)
        if path.exists():
            records.extend(parse_congressional_workbook(path, cycle, url))
        else:
            gaps.append({
                "cycle": cycle, "state_code": "", "chamber": "House and Senate",
                "district": "", "issue": "missing_complete_primary_results_workbook",
                "source_url": url, "detail": error or "Run with --download",
            })
    for row in records:
        senate_class = ((row["workbook_cycle"] - 2018) // 2) % 3 + 1
        if (
            row["chamber"] == "Senate" and senate_class in senate_classes
            and row["state_code"] not in senate_classes[senate_class]
        ):
            row["term"] = "special"
    by_seat = defaultdict(list)
    state_primary_path = output_dir / "state_primary_results.csv"
    state_primary_rows = read_csv(state_primary_path) if state_primary_path.exists() else []
    state_by_seat = defaultdict(list)
    for row in state_primary_rows:
        if row["stage"] not in PRIMARY_RESULT_STAGES:
            continue
        state_by_seat[(
            int(row["cycle"]), row["chamber"], row["state_code"], row["district"], row["term"],
        )].append(row)
    for row in records:
        by_seat[(row["workbook_cycle"], row["chamber"], row["state_code"],
                 row["district"], row["term"])].append(row)
        if not row["district"]:
            gaps.append({
                "cycle": row["workbook_cycle"], "state_code": row["state_code"],
                "chamber": row["chamber"], "district": "",
                "issue": "missing_district_in_source_row", "source_url": row["source_url"],
                "detail": f'{row["source_sheet"]}:{row["source_row"]} {row["candidate_name"]}',
            })

    inventory = []
    for cycle in cycles:
        for chamber in ("House", "Senate"):
            seats = _regular_seats(records, cycle, chamber)
            basis = cycle
            senate_class = ((cycle - 2018) // 2) % 3 + 1
            if chamber == "Senate" and senate_class in senate_classes:
                seats = {(state, "S") for state in senate_classes[senate_class]}
                basis = f"Senate Class {senate_class}"
            elif not seats:
                # A prior map supplies a checklist, never an invented ballot roster.
                basis = 2022 if chamber == "House" else cycle - 6
                seats = _regular_seats(records, basis, chamber)
                if chamber == "House" and cycle % 4 == 0:
                    seats |= {seat for seat in _regular_seats(records, 2020, chamber)
                              if seat[0] == "PR"}
            state_seats = {seat for seat in seats if seat[0] not in TERRITORIES}
            expected_count = 435 if chamber == "House" else (34 if cycle % 6 == 0 else 33)
            if len(state_seats) != expected_count:
                gaps.append({
                    "cycle": cycle, "state_code": "", "chamber": chamber, "district": "",
                    "issue": "regular_seat_count_mismatch", "source_url": FEC_INDEX,
                    "detail": f"Expected {expected_count}, found {len(state_seats)}; "
                              "special contests must be reconciled separately",
                })
            for state, district in sorted(seats):
                if state == "PR" and cycle % 4:
                    continue
                seat_records = by_seat.get((cycle, chamber, state, district, "regular"), [])
                state_records = state_by_seat.get((cycle, chamber, state, district, "regular"), [])
                inventory.append({
                    "cycle": cycle, "chamber": chamber, "state_code": state,
                    "district": district, "seat_basis_cycle": basis,
                    "status": (
                        "source_records_present" if seat_records else
                        "state_primary_returns_present" if state_records else "missing_results"
                    ),
                    "source_rows": len(seat_records),
                    "state_primary_candidate_rows": len(state_records),
                    "state_primary_source_ids": " | ".join(sorted({
                        row["source_id"] for row in state_records
                    })),
                    "candidate_rows": sum(row["row_kind"] == "candidate" for row in seat_records),
                    "primary_reported_rows": sum(bool(row["primary_votes_raw"]) for row in seat_records),
                    "primary_numeric_rows": sum(row["primary_votes"] != "" for row in seat_records),
                    "democratic_primary_rows": sum(
                        row["party_raw"] in {"D", "DEM", "DFL", "W(D)", "W(DEM)", "W(DFL)"}
                        and bool(row["primary_votes_raw"]) for row in seat_records
                    ),
                    "primary_type": (
                        "nonpartisan_top_two_or_top_four"
                        if state in {"CA", "WA"} or (state == "AK" and cycle >= 2022)
                        else "louisiana_open_primary" if state == "LA" and (chamber == "House" or cycle < 2026)
                        else "requires_election_specific_verification"
                    ),
                    "candidate_policy_text_status": "not_collected_for_nationwide_universe",
                })
                if not seat_records and not state_records:
                    gaps.append({
                        "cycle": cycle, "state_code": state, "chamber": chamber,
                        "district": district, "issue": "missing_seat_results",
                        "source_url": FEC_INDEX,
                        "detail": "Checklist entry only; not a verified candidate roster",
                    })

    special_path = raw_dir / "federalspecialelectionslist.pdf"
    if download:
        special_path.write_bytes(_download_workbook(SPECIAL_URL))
    specials = []
    if special_path.exists():
        _, text = _extract_pdf_text(special_path.read_bytes())
        specials = parse_special_elections(
            text, int(config.get("source_start", config["study_start"])[:4]),
            config["research_cutoff"],
        )
        write_csv(output_dir / "special_election_inventory.csv", specials, list(specials[0]))
        for row in specials:
            gaps.append({
                "cycle": row["reported_election_date"][:4], "state_code": row["state_code"],
                "chamber": row["chamber"], "district": row["district"],
                "issue": "special_primary_roster_not_reconciled", "source_url": SPECIAL_URL,
                "detail": row["special_election_id"],
            })
    # A dated special-election list is not proof of current exhaustive coverage.
    for year in range(int(config.get("source_start", config["study_start"])[:4]), end + 1):
        for state in sorted(STATE_CODES):
            gaps.append({
                "cycle": year, "state_code": state, "chamber": "House and Senate",
                "district": "", "issue": "special_election_census_not_reconciled",
                "source_url": FEC_INDEX,
                "detail": "Verify all special primaries, conventions, repeat elections and "
                          "runoffs; overlapping FEC appendices are not a unique-race census",
            })
    fields = ["cycle", "state_code", "chamber", "district", "issue", "source_url", "detail"]
    identity_issues = source_identity_issues(records + ballot_rows)
    write_csv(
        output_dir / "source_identity_review.csv", identity_issues,
        ["issue", "fec_candidate_id", "source_url", "names", "locators"],
    )
    write_csv(output_dir / "candidate_results.csv", records, RESULT_FIELDS)
    write_csv(output_dir / "regular_seat_inventory.csv", inventory, list(inventory[0]) if inventory else [
        "cycle", "chamber", "state_code", "district", "status",
    ])
    state_cycles = []
    for cycle, state in sorted({(row["cycle"], row["state_code"]) for row in inventory}):
        seats = [row for row in inventory if row["cycle"] == cycle and row["state_code"] == state]
        house = [row for row in seats if row["chamber"] == "House"]
        senate = [row for row in seats if row["chamber"] == "Senate"]
        state_cycles.append({
            "cycle": cycle,
            "state_code": state,
            "expected_house_districts": " | ".join(row["district"] for row in house),
            "house_districts_with_returns": " | ".join(
                row["district"] for row in house if row["status"] != "missing_results"
            ),
            "house_districts_missing_returns": " | ".join(
                row["district"] for row in house if row["status"] == "missing_results"
            ),
            "regular_senate_seat": (
                "not_scheduled" if not senate else
                "missing_returns" if senate[0]["status"] == "missing_results" else "returns_present"
            ),
            "regular_seats": len(seats),
            "regular_seats_missing_returns": sum(row["status"] == "missing_results" for row in seats),
            "primary_roster_reconciliation": "pending_election_specific_review",
            "special_election_reconciliation": "pending",
            "scope": "nonvoting_delegation" if state in TERRITORIES else "state",
        })
    write_csv(
        output_dir / "state_cycle_coverage.csv", state_cycles,
        list(state_cycles[0]) if state_cycles else ["cycle", "state_code"],
    )
    write_csv(output_dir / "coverage_gaps.csv", gaps, fields)
    write_csv(manifest_path, sources, list(sources[0]))
    party_accounting = export_party_accounting(output_dir, inventory, state_primary_rows)
    longitudinal_rows = export_longitudinal_primary_records(output_dir, records, state_primary_rows)
    reconciliation = reconcile_dsa_congressional(output_dir)
    summary = {
        **party_accounting,
        **reconciliation,
        "longitudinal_primary_observations": longitudinal_rows,
        "as_of": config["research_cutoff"],
        "scope": "All parties, U.S. House and Senate; independent of DSA endorsements",
        "source_start": config.get("source_start", config["study_start"]),
        "election_cycles": cycles,
        "source_result_rows": len(records),
        "state_primary_candidate_rows": len(state_primary_rows),
        "state_primary_contests": len({row["contest_id"] for row in state_primary_rows}),
        "source_rows_by_cycle": dict(sorted(Counter(
            row["workbook_cycle"] for row in records
        ).items())),
        "regular_seat_inventory_rows": len(inventory),
        "missing_regular_seat_results": sum(row["status"] == "missing_results" for row in inventory),
        "special_election_calendar_rows": len(specials),
        "special_election_calendar_updated_on": specials[0]["source_updated_on"] if specials else None,
        "general_ballot_2024_rows": len(ballot_rows),
        "source_identity_review_groups": len(identity_issues),
        "auxiliary_source_sha256": {
            path.name: hashlib.sha256(path.read_bytes()).hexdigest()
            for path in (
                raw_dir / "senate_class_i.html", raw_dir / "senate_class_ii.html",
                raw_dir / "senate_class_iii.html", special_path, ballot_path,
            )
            if path.exists()
        },
        "gap_records": len(gaps),
        "complete": False,
        "limitations": [
            "Source rows are not unique candidates or unique elections; fusion ballot lines "
            "and overlapping special-election appendices are retained with provenance.",
            "Blank votes are missing, not zero. Unopposed and asterisk flags are preserved, "
            "not invented votes or proof that a primary was held.",
            "Named write-ins and aggregate/scattered write-ins remain distinct.",
            "All stages and raw source columns are retained, including ranked-choice rounds.",
            "Primary dates, certified rosters, and special-election identities still require "
            "state-by-state reconciliation; source rows do not prove exhaustive coverage.",
            "Candidate filings and general-election nominees cannot replace primary losers.",
            "Nationwide campaign-policy text is not yet collected; existing DSA text remains "
            "a separate, narrower corpus.",
            "Missing regular-seat return counts describe numeric-source presence, not a census "
            "of unresolved primary nominations. Party accounting separately identifies official "
            "no-primary decisions and future-scheduled elections.",
        ],
    }
    (output_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8",
    )
    return summary
